"""The project bill (final account) file: figures from the tracker, a spreadsheet whose formulas give
the same figures, read back before any link is handed out (31 Oct 2026). Made-up figures."""
import io
import shutil
import subprocess

import pytest

from vula.commerce import project_bill as pb


def _sheet_position():
    """A project with the business's reconciled sheet on file (sections carry actual and paid)."""
    return {
        "project": "Seaview", "as_of": "2026-10-31", "fee_pct": 0, "fee_cents": 0,
        "sheet": {"actual_cents": 3_400_000},
        "sections": [
            {"section": "Building labour", "budget_cents": 1_000_000, "actual_cents": 1_200_000, "paid_cents": 1_000_000},
            {"section": "Plumbing", "budget_cents": 500_000, "actual_cents": 650_000, "paid_cents": 650_000},
            {"section": "Tiles", "budget_cents": 1_500_000, "actual_cents": 1_550_000, "paid_cents": 1_200_000},
        ],
        "paid_directly_cents": 900_000, "to_pay_directly_cents": 250_000, "received_cents": 1_800_000,
        "direct_payments": [],
    }


def _books_position():
    """No reconciled sheet: budget sections, the business's cost per trade, direct payments."""
    return {
        "project": "Oak Street", "as_of": "2026-10-31", "fee_pct": 10, "fee_cents": 120_000,
        "sheet": None,
        "sections": [{"section": "Electrical", "budget_cents": 400_000},
                     {"section": "Painting", "budget_cents": 300_000}],
        "trades": [{"trade": "Electrical", "cents": 450_000}, {"trade": "Scaffolding", "cents": 50_000}],
        "business_cost_cents": 1_200_000,          # 700,000 not yet on a trade
        "paid_directly_suppliers_cents": 80_000, "paid_directly_slips_cents": 20_000,
        "direct_payments": [{"supplier": "Sparks & Co", "amount_cents": 150_000, "status": "paid"},
                            {"supplier": "Glass <b>Ltd</b>", "amount_cents": 60_000, "status": "to_pay"}],
        "paid_directly_cents": 250_000, "to_pay_directly_cents": 60_000, "received_cents": 900_000,
        "spent_cents": 1_450_000,
    }


def test_balance_follows_the_final_account_method():
    f = pb.figures(_sheet_position())
    assert f["actual"] == 3_400_000 and f["budget"] == 3_000_000 and f["paid"] == 2_850_000
    assert f["over_budget"] == 400_000
    # total − paid directly − still to pay directly − paid to the business
    assert f["balance"] == 3_400_000 - 900_000 - 250_000 - 1_800_000 == 450_000


def test_without_a_sheet_the_lines_come_from_the_books_and_add_up_to_the_spend():
    pos = _books_position()
    f = pb.figures(pos)
    names = [ln["section"] for ln in f["lines"]]
    assert names[:3] == ["Electrical", "Painting", "Scaffolding"]
    elec = f["lines"][0]
    assert elec["budget"] == 400_000 and elec["actual"] == 450_000 == elec["paid"]
    assert "Other project costs (not yet put on a trade)" in names
    # every rand spent, plus what the client still has to pay suppliers directly
    assert f["actual"] == pos["spent_cents"] + pos["to_pay_directly_cents"]
    to_pay = next(ln for ln in f["lines"] if "Glass" in ln["section"])
    assert to_pay["paid"] == 0 and to_pay["actual"] == 60_000
    assert f["fee"] == 120_000 and f["total"] == f["actual"] + 120_000
    assert f["balance"] == f["total"] - 250_000 - 60_000 - 900_000


@pytest.mark.parametrize("make", [_sheet_position, _books_position])
def test_the_spreadsheet_reads_back_to_the_books(make):
    f = pb.figures(make())
    data, cells = pb.workbook("Made Up Builders", "Seaview", f, [], "2026-10-31")
    assert pb.read_back(data, cells, f) == []


def test_a_spreadsheet_that_does_not_add_up_is_caught():
    from openpyxl import load_workbook
    f = pb.figures(_sheet_position())
    data, cells = pb.workbook("Made Up Builders", "Seaview", f, [], "2026-10-31")
    wb = load_workbook(io.BytesIO(data))
    wb.active["C7"] = 7000           # Plumbing actual changed in the file
    buf = io.BytesIO()
    wb.save(buf)
    wrong = pb.read_back(buf.getvalue(), cells, f)
    assert any(w.startswith("actual") for w in wrong) and any(w.startswith("balance") for w in wrong)


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_libreoffice_gives_the_same_figures(tmp_path):
    """The formulas are real spreadsheet formulas: LibreOffice computes the same balance."""
    from openpyxl import load_workbook
    f = pb.figures(_books_position())
    data, cells = pb.workbook("Made Up Builders", "Oak Street", f, [], "2026-10-31")
    src = tmp_path / "bill.xlsx"
    src.write_bytes(data)
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(src)],
                   check=True, capture_output=True, timeout=120)
    ws = load_workbook(out / "bill.xlsx", data_only=True).active
    assert round(ws[cells["balance"]].value * 100) == f["balance"]
    assert round(ws[cells["actual"]].value * 100) == f["actual"]


def test_decisions_become_notes_on_their_line():
    f = pb.figures(_sheet_position())
    decs = [{"decision": "Plumbing change paid by the business, not the client", "created_at": "2026-10-30T10:00:00"},
            {"decision": "Further painting is billed separately", "created_at": "2026-10-30T11:00:00"}]
    data, _ = pb.workbook("Made Up Builders", "Seaview", f, decs, "2026-10-31")
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(data)).active
    assert "Plumbing change" in (ws["G7"].value or "")
    texts = [c.value for c in ws["B"] if isinstance(c.value, str)]
    assert "Further painting is billed separately" in texts      # listed under the decisions


def test_the_pdf_escapes_what_came_from_the_books():
    f = pb.figures(_books_position())
    page = pb.render_html("Made <i>Up</i>", "Oak Street", f, [], "2026-10-31")
    assert "<b>Ltd</b>" not in page and "&lt;b&gt;Ltd&lt;/b&gt;" in page
    assert "Made &lt;i&gt;Up&lt;/i&gt;" in page


def test_build_refuses_when_the_read_back_disagrees(monkeypatch):
    from vula.commerce import decisions, project_tracker
    monkeypatch.setattr(project_tracker, "position", lambda t, p: _sheet_position())
    monkeypatch.setattr(decisions, "for_project", lambda *a, **k: [])
    monkeypatch.setattr(pb, "read_back", lambda *a: ["balance: file 1 vs books 2"])
    uploads = []
    monkeypatch.setattr("vula.api.whatsapp._upload_to_storage", lambda *a: uploads.append(a) or "u")
    with pytest.raises(RuntimeError):
        pb.build("tenant-a", "Seaview", "Made Up Builders")
    assert uploads == []                               # nothing stored, nothing linked


def test_build_stores_both_files_privately_and_links_them(monkeypatch):
    from vula import storage_links
    from vula.commerce import decisions, project_tracker
    monkeypatch.setattr(project_tracker, "position", lambda t, p: _sheet_position())
    monkeypatch.setattr(decisions, "for_project", lambda *a, **k: [])
    uploads = []
    monkeypatch.setattr("vula.api.whatsapp._upload_to_storage",
                        lambda bucket, path, data, ct: uploads.append((bucket, path, ct)) or f"stored:{path}")
    monkeypatch.setattr(storage_links, "signed", lambda url, ttl: f"signed:{url}")
    out = pb.build("tenant-a", "Seaview", "Made Up Builders")
    assert {u[0] for u in uploads} == {"documents"}
    assert all(u[1].startswith("tenant-a/projects/seaview/") for u in uploads)
    assert out["links"]["xlsx"].startswith("signed:") and out["links"]["pdf"].startswith("signed:")
    assert out["figures"]["balance"] == 450_000


@pytest.mark.asyncio
async def test_the_tool_asks_which_project_and_reports_a_failed_check(monkeypatch):
    from core.skills.commerce_admin import CommerceAdminSkill
    from vula.commerce import service
    monkeypatch.setattr(service, "canonical_project", lambda t, raw: None)
    monkeypatch.setattr(service, "registered_projects", lambda t: [{"name": "Seaview"}])
    monkeypatch.setattr("vula.api.tenants.display_name", lambda t: "Made Up Builders")
    skill = CommerceAdminSkill()
    out = await skill._project_bill_file("tenant-a", {"project": "Mars base"})
    assert out["status"] == "need_info"

    def bad(*a):
        raise RuntimeError("the spreadsheet didn't add up to the books")
    monkeypatch.setattr(pb, "build", bad)
    out = await skill._project_bill_file("tenant-a", {"project": "Seaview"})
    assert "haven't sent it" in out["error"]


@pytest.mark.asyncio
async def test_the_tool_replies_with_the_summary_and_both_links(monkeypatch):
    from core.skills.commerce_admin import CommerceAdminSkill
    from vula.commerce import service
    monkeypatch.setattr(service, "canonical_project", lambda t, raw: "Seaview")
    monkeypatch.setattr(service, "registered_projects", lambda t: [{"name": "Seaview"}])
    monkeypatch.setattr("vula.api.tenants.display_name", lambda t: "Made Up Builders")
    f = pb.figures(_sheet_position())
    monkeypatch.setattr(pb, "build", lambda *a: {"project": "Seaview", "figures": f, "decisions": 0,
                                                  "links": {"xlsx": "https://x/bill.xlsx", "pdf": "https://x/bill.pdf"}})
    out = await CommerceAdminSkill()._project_bill_file("tenant-a", {"project": "seaview"})
    text = out["reply_verbatim"]
    assert "Balance due: R4,500.00" in text and "https://x/bill.xlsx" in text and "https://x/bill.pdf" in text

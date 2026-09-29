"""The price book: DIGG's QS rates learned from its own documents (2026-09-28).

Ian: "the QS rates are not adapting… so many costs, even the BOQ on HPC, and Vula hasn't picked
up any cost from invoices or labour." digg-demo had ~1,000 priced line items on filed invoices,
quotes and BOQs and 2 rows in vula_qs_rates. Every filed money document's priced lines are now
recorded (migration 184) and rolled up per item; the QS rate API, the calculations skill and
Takeoff pricing read them after the tenant's own rates, which are never changed by this. BOQ
spreadsheets are read row by row, and one project's spellings are merged. Real DIGG line shapes.
"""
import pytest

from vula.api import qs
from vula.commerce import price_book, service

TID = "digg-demo"


class _Res:
    def __init__(self, data):
        self.data = data


class _Q:
    def __init__(self, db, table):
        self.db, self.table, self.op, self.payload = db, table, "select", None
        self.filters, self.likes, self.rng = [], [], None

    def select(self, *_a, **_k):
        return self

    def insert(self, payload):
        self.op, self.payload = "insert", payload
        return self

    def update(self, payload):
        self.op, self.payload = "update", payload
        return self

    def upsert(self, payload, **_k):
        self.op, self.payload = "insert", payload
        return self

    def delete(self):
        self.op = "delete"
        return self

    def eq(self, col, val):
        self.filters.append((col, val))
        return self

    def in_(self, col, vals):
        self.filters.append((col, set(vals)))
        return self

    def ilike(self, col, pat):
        self.likes.append((col, pat.strip("%").lower()))
        return self

    def order(self, *_a, **_k):
        return self

    def limit(self, *_a, **_k):
        return self

    def range(self, a, b):
        self.rng = (a, b)
        return self

    def _hit(self, r):
        for c, v in self.filters:
            if isinstance(v, set):
                if r.get(c) not in v:
                    return False
            elif r.get(c) != v:
                return False
        return all(p in str(r.get(c) or "").lower() for c, p in self.likes)

    def execute(self):
        rows = self.db.tables.setdefault(self.table, [])
        if self.op == "insert":
            new = self.payload if isinstance(self.payload, list) else [self.payload]
            out = []
            for p in new:
                row = {"id": f"{self.table}-{len(rows) + 1}", **p}
                rows.append(row)
                out.append(dict(row))
            return _Res(out)
        hit = [r for r in rows if self._hit(r)]
        if self.op == "delete":
            self.db.tables[self.table] = [r for r in rows if r not in hit]
            return _Res(hit)
        if self.op == "update":
            for r in hit:
                r.update(self.payload)
            return _Res([dict(r) for r in hit])
        if self.rng:
            hit = hit[self.rng[0]:self.rng[1] + 1]
        return _Res([dict(r) for r in hit])


class FakeDB:
    def __init__(self, tables=None):
        self.tables = tables or {}

    def table(self, name):
        return _Q(self, name)


@pytest.fixture()
def db(monkeypatch):
    fake = FakeDB()
    monkeypatch.setattr(service, "_client", lambda: fake)
    return fake


def _inv(i, date, supplier, lines, category="Invoice", project="HPC Bokaap", filename=None):
    return {"id": f"d{i}", "tenant_id": TID, "category": category, "project": project,
            "filename": filename or f"inv{i}.pdf", "summary": "",
            "fields": {"supplier": supplier, "date": date, "line_items": lines}}


BOARD = "SOLID 12mm BOARD 3.6 X 1.2 TE"


def _board(qty, unit_c):
    return {"description": BOARD, "quantity": qty, "unit_price_cents": unit_c, "total_cents": qty * unit_c}


# ── recording ─────────────────────────────────────────────────────────────────

def test_only_real_priced_lines_are_recorded_and_classified():
    row = _inv(1, "2026-08-02", "SOLID CAPE (PTY) LTD", [
        _board(30, 26314),
        {"description": "Tiling Labour", "quantity": 1, "unit_price_cents": None, "total_cents": 150000},
        {"description": "Deliver", "quantity": 1, "unit_price_cents": 21000, "total_cents": 21000},
        {"description": "Jacks", "quantity": 4, "unit_price_cents": 8000, "total_cents": 32000},
        {"description": "Contingency — 15% of measured work", "quantity": 1, "unit_price_cents": 20993600},
        {"description": "VAT", "quantity": 1, "total_cents": 12000},
        {"description": "Commissioning", "quantity": 1.0, "unit_price_cents": None, "total_cents": None},
        {"description": "Selflevelling ITE F30 Screed 2mm thick (405 sqm x R168)", "quantity": 405,
         "unit_price_cents": 16800},
    ])
    obs = {o["description"]: o for o in price_book.observations_for(row)}
    assert set(obs) == {BOARD, "Tiling Labour", "Deliver", "Jacks",
                        "Selflevelling ITE F30 Screed 2mm thick (405 sqm x R168)"}
    assert obs[BOARD]["unit_price_cents"] == 26314 and obs[BOARD]["kind"] == "material"
    assert obs["Tiling Labour"]["kind"] == "labour" and obs["Tiling Labour"]["unit_price_cents"] == 150000
    assert obs["Deliver"]["kind"] == "delivery" and obs["Jacks"]["kind"] == "plant"
    assert obs["Selflevelling ITE F30 Screed 2mm thick (405 sqm x R168)"]["unit"] == "m2"
    assert obs[BOARD]["supplier"] == "SOLID CAPE (PTY) LTD" and obs[BOARD]["project"] == "HPC Bokaap"
    assert obs[BOARD]["observed_on"] == "2026-08-02" and obs[BOARD]["source_kind"] == "invoice"


def test_refunds_and_non_money_documents_add_nothing():
    refund = _inv(1, "2026-08-02", "Gardens Handiman", [_board(1, 26314)],
                  filename="POS Account Refund 21-366230.pdf")
    drawing = _inv(2, "2026-08-02", None, [_board(1, 26314)], category="Drawing / Plan")
    assert price_book.observations_for(refund) == [] and price_book.observations_for(drawing) == []


def test_refiling_a_document_replaces_its_lines(db):
    row = _inv(1, "2026-08-02", "Solid Cape", [_board(30, 26314)])
    assert price_book.record_from_document(TID, row) == 1
    row["fields"]["line_items"] = [_board(30, 27000), _board(5, 27000) | {"description": "SOLID CORNERBEAD 3m"}]
    assert price_book.record_from_document(TID, row) == 2
    assert len(db.tables["vula_price_observations"]) == 2          # not 3


# ── learned rates ─────────────────────────────────────────────────────────────

def test_a_rate_is_the_median_paid_with_latest_and_range(db):
    for i, (date, price, supplier) in enumerate([("2026-06-01", 25000, "Solid Cape"),
                                                  ("2026-07-10", 26314, "Solid Cape"),
                                                  ("2026-08-20", 27500, "City Build It")]):
        price_book.record_from_document(TID, _inv(i, date, supplier, [_board(10, price)]))
    # a quote at a silly price doesn't move a rate the business has actually paid
    price_book.record_from_document(TID, _inv(9, "2026-09-01", "X", [_board(10, 90000)],
                                              category="Quote / Estimate"))
    (r,) = price_book.rates(TID, "12mm board")
    assert r["rate_cents"] == 26314 and r["basis"] == "paid"
    assert (r["latest_cents"], r["latest_on"], r["latest_supplier"]) == (27500, "2026-08-20", "City Build It")
    assert (r["low_cents"], r["high_cents"]) == (25000, 27500)
    assert r["observations"] == 4 and r["suppliers"] == 3
    assert r["source"].startswith("Learned from 3 invoices, 1 quote")


def test_only_quoted_items_say_so(db):
    price_book.record_from_document(TID, _inv(1, "2026-07-01", "Flush Bathrooms", [
        {"description": "Diverter Mixer Concealed Part", "quantity": 2, "unit_price_cents": 295000}],
        category="Quote / Estimate"))
    (r,) = price_book.rates(TID, "diverter")
    assert r["basis"] == "quoted" and r["rate"] == 2950.0


def test_worker_day_rates_become_labour_rates(db):
    db.tables["commerce_workers"] = [
        {"id": "w1", "tenant_id": TID, "name": "Sipho", "type": "casual", "rate_cents": 35000,
         "rate_period": "daily", "active": True},
        {"id": "w2", "tenant_id": TID, "name": "Gone", "type": "casual", "rate_cents": 30000,
         "rate_period": "daily", "active": False}]
    assert price_book.record_worker_rates(TID) == 1
    (r,) = price_book.rates(TID, kind="labour")
    assert (r["rate_cents"], r["unit"], r["basis"]) == (35000, "day", "paid")


# ── the QS rate API: own rates first, never changed ──────────────────────────

@pytest.fixture()
def rates_db(db):
    db.tables["vula_qs_rates"] = [
        {"id": "r1", "tenant_id": TID, "code": "WALL-BB", "unit": "m2", "rate": 685.0,
         "description": "Breeze block walling, 140mm, plastered both sides", "source": "DIGG 2024"},
        {"id": "r2", "tenant_id": TID, "unit": "each", "rate": 220.0, "description": BOARD},
        {"id": "rx", "tenant_id": "other", "unit": "m2", "rate": 1.0, "description": "theirs"}]
    for i, price in enumerate([26314, 26314, 27500]):
        price_book.record_from_document(TID, _inv(i, f"2026-08-0{i + 1}", "Solid Cape", [_board(10, price)]))
    price_book.record_from_document(TID, _inv(7, "2026-08-09", "Solid Cape", [
        {"description": "SOLID CORNERBEAD 3m", "quantity": 40, "unit_price_cents": 15000}]))
    return db


@pytest.mark.asyncio
async def test_list_shows_own_rates_first_then_learned_and_flags_drift(rates_db):
    out = await qs.list_rates(TID)
    assert {r["id"] for r in out["own"]} == {"r1", "r2"}            # own tenant only
    assert [r["description"] for r in out["learned"]] == ["SOLID CORNERBEAD 3m"]   # board is own
    assert out["rates"][:2] == out["own"]
    board = next(r for r in out["own"] if r["id"] == "r2")
    assert board["drift"]["learned_rate"] == 263.14 and board["drift"]["change_pct"] == 19.6
    assert rates_db.tables["vula_qs_rates"][1]["rate"] == 220.0        # nothing written


def test_skill_lookup_finds_by_word_and_includes_learned(rates_db):
    assert [r["code"] for r in qs.search_rates(TID, "brick wall")] == []   # not a word match
    assert [r["code"] for r in qs.search_rates(TID, "block wall")] == ["WALL-BB"]
    learned = qs.search_rates(TID, "cornerbead")
    assert learned[0]["learned"] and learned[0]["rate"] == 150.0
    assert qs.search_rates(TID, "cornerbead", include_learned=False) == []


@pytest.mark.asyncio
async def test_saving_the_same_rate_updates_it_tenant_scoped(rates_db):
    out = await qs.upsert_rate(TID, qs.RateIn(description=BOARD.lower(), unit="each", rate=263.14))
    assert out["id"] == "r2" and out["rate"] == 263.14 and "T" in out["updated_at"]
    assert len([r for r in rates_db.tables["vula_qs_rates"] if r["tenant_id"] == TID]) == 2
    other = await qs.upsert_rate(TID, qs.RateIn(id="rx", description="theirs", rate=5))
    assert other == {"error": "Rate not found."}
    await qs.delete_rate(TID, "rx")
    assert any(r["id"] == "rx" for r in rates_db.tables["vula_qs_rates"])


def test_takeoff_prices_from_the_tenants_rate_first(rates_db):
    from vula.takeoff.boq_generator import BOQGenerator
    gen = BOQGenerator.__new__(BOQGenerator)
    gen.tenant_id, gen._own = TID, None
    low, high, mid, source = gen._own_rate("walling_block", "m²")
    assert (mid, source) == (685.0, "Your rate · DIGG 2024")
    assert gen._own_rate("paint_interior", "m²") is None                # no match → market rate
    gen.tenant_id, gen._own = "default", None
    assert gen._own_rate("walling_block", "m²") is None


# ── BOQ spreadsheets ──────────────────────────────────────────────────────────

def _hpc_sheet(path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Extra air"
    ws.append(["HPC Cape Town — Extra air"])
    ws.append([])
    ws.append(["Item", "Description", "Qty", "Unit", "Rate", "Amount"])
    ws.append([None, "Airconditioning", None, None, None, None])
    for n in range(1, 21):
        ws.append([f"1.{n}", f"Cassette unit {n} kW supply & install", 2, "no", 18500, 37000])
    ws.append([None, "Toilet Extract", None, None, None, None])
    for n in range(1, 21):
        ws.append([f"2.{n}", f"Extract fan type {n}", 1, "no", None, 4500])
    ws.append([None, "Sub-total", None, None, None, 830000])
    wb.save(path)


def test_a_boq_sheet_is_read_every_line_with_its_section(tmp_path):
    from vula.ingestion import boq_sheet
    p = tmp_path / "HPC Cape Town Extra air.xlsx"
    _hpc_sheet(p)
    lines = boq_sheet.parse(p)
    assert len(lines) == 40
    assert lines[0] == {"description": "Cassette unit 1 kW supply & install", "quantity": 2, "unit": "no",
                        "unit_price_cents": 1850000, "total_cents": 3700000, "section": "Airconditioning",
                        "code": "1.1"}
    assert lines[-1]["unit_price_cents"] == 450000 and lines[-1]["section"] == "Toilet Extract"
    fields = boq_sheet.apply_to_fields(
        {"total_cents": 135471992, "line_items": [{"description": "Toilet Extract", "quantity": 1}] * 4,
         "_unverified_figures": [{"field": "line_items[2].total_cents"}]}, lines)
    assert len(fields["line_items"]) == 40 and fields["total_cents"] == 135471992   # read total kept
    assert fields["sections"] == [{"section": "Airconditioning", "budget_cents": 74000000},
                                  {"section": "Toilet Extract", "budget_cents": 9000000}]
    assert "_unverified_figures" not in fields


@pytest.mark.asyncio
async def test_analyze_document_completes_a_boq_sheet(tmp_path, monkeypatch):
    from vula.api import whatsapp as wa
    p = tmp_path / "HPC_CapeTown_Interior_BOQ_1.xlsx"
    _hpc_sheet(p)
    out = await wa._complete_boq_lines({"category": "Bill of Quantities (BOQ)", "fields": {}}, p, "x", p.name)
    assert len(out["fields"]["line_items"]) == 40 and out["fields"]["total_cents"] == 83000000


# ── one name per project ──────────────────────────────────────────────────────

def test_project_spellings_merge_onto_one(db):
    db.tables["vula_projects"] = [{"tenant_id": TID, "name": "HPC Bokaap"}]
    db.tables["vula_filed_documents"] = [{"tenant_id": TID, "project": "PORTERFIELD"}] * 18 + \
        [{"tenant_id": TID, "project": "Porterfield"}] * 4
    assert service.canonical_project(TID, "HPC_Bokaap") == "HPC Bokaap"
    assert service.canonical_project(TID, "hpc  bokaap ") == "HPC Bokaap"
    assert service.canonical_project(TID, "porterfield") == "PORTERFIELD"     # the spelling used most
    assert service.canonical_project(TID, "Sporty – Phase 2") == "Sporty – Phase 2"


# ── learn from history ────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_learn_from_history_builds_the_book_and_is_idempotent(db, monkeypatch):
    from vula.commerce import reread
    db.tables["vula_filed_documents"] = [
        _inv(1, "2026-08-02", "Solid Cape", [_board(30, 26314)]),
        _inv(2, "2026-08-03", "Solid Cape", [_board(10, 27000)]),
        _inv(3, "2026-08-03", None, [], category="Drawing / Plan"),
    ]
    monkeypatch.setattr(reread, "candidates", lambda tid, limit=200: [])
    first = await reread.learn_from_history(TID)
    again = await reread.learn_from_history(TID)
    assert first["step"] == again["step"] == "done"
    assert (first["documents"], first["priced_lines"]) == (2, 2)
    assert len(db.tables["vula_price_observations"]) == 2
    assert again["summary"]["items"] == 1 and again["summary"]["suppliers"] == 1
    assert reread.learn_status(TID)["running"] is False


@pytest.mark.asyncio
async def test_assigning_a_project_in_the_dashboard_carries_the_bill_boq_and_prices(db, monkeypatch):
    from vula.api import documents
    monkeypatch.setattr(documents, "_client", lambda: db)
    monkeypatch.setattr(documents, "authorized_tenant", lambda req: TID)
    monkeypatch.setattr("vula.integrations.doc_filing.learn_filing_rule", lambda *a: None)
    monkeypatch.setattr("vula.integrations.finances.post_finance_from_doc", lambda *a: None)
    db.tables["vula_projects"] = [{"tenant_id": TID, "name": "HPC Bokaap"}]
    boq = _inv(1, "2026-08-20", None, [_board(10, 26314)], category="Bill of Quantities (BOQ)", project=None)
    boq.update(commerce_invoice_id="ci1", fields={**boq["fields"], "total_cents": 263140,
                                                   "sections": [{"section": "Boards", "budget_cents": 263140}]})
    db.tables["vula_filed_documents"] = [boq]
    db.tables["commerce_invoices"] = [{"id": "ci1", "tenant_id": TID, "project": None}]
    price_book.record_from_document(TID, boq)

    out = await documents.assign_project("d1", documents.AssignIn(project="hpc_bokaap"), object())
    assert out["project"] == "HPC Bokaap"
    assert db.tables["commerce_invoices"][0]["project"] == "HPC Bokaap"
    assert db.tables["vula_price_observations"][0]["project"] == "HPC Bokaap"
    (pb,) = db.tables["vula_project_boq"]
    assert (pb["project"], pb["total_cents"], pb["sections"][0]["section"]) == ("HPC Bokaap", 263140, "Boards")


# ── supplier price lists and stock sheets (2026-09-29, Gerflor) ──────────────
# Gerflor's documents were mostly product knowledge (fire reports, spec sheets, brochures), but
# the SPM wall-protection and vinyl-sheeting price lists were filed with a summary and no lines,
# and "DT SOH and Planning" was filed as a Programme before stock sheets were stored as rows.

_SPM_TEXT = """SPM WALL PROTECTION & HANDRAILS — PRICE LIST 2026 (excl VAT)
BR200 Bumper rail 200mm  per m  R 485.00
HR50 Handrail 50mm round  per m  R 612.50
CG75 Corner guard 75x75  per 3m length  R 1 290.00"""

_SOH_TEXT = """DT SOH and Planning — SOH m² – 07.09.26
VIRTUO 30 COL: SUNNY WHITE 2.00MM 532
VIRTUO 30 COL: BAITA MEDIUM 2.00MM 581.4 534 Est. Mid Nov - TBC
MAC TILES-COL: 612 ST/GREY 1.60MM 6504.3
AMBIANCE ULTRA TERRA COL: 0203 STEAM GREY 2.00W | 540 | 2000 | Est. Mid Oct - TBC
MAC TILES-COL: 656 BASIL 2.00MM 378 164m² Reserved TVET College
EL7 SD ROBUST COL: GREY 2.00MM 220"""


def _spm_lines():
    return [{"description": "BR200 Bumper rail 200mm", "unit": "m", "unit_price_cents": 48500},
            {"description": "HR50 Handrail 50mm round", "unit": "m", "unit_price_cents": 61250},
            {"description": "CG75 Corner guard 75x75", "unit": "3m length", "unit_price_cents": 129000}]


@pytest.mark.asyncio
async def test_a_price_list_is_itemised_but_a_catalogue_is_left_alone(monkeypatch):
    from vula.api import whatsapp as wa
    calls = []

    async def fake_lines(text, filename, doc_kind="Bill of Quantities"):
        calls.append(doc_kind)
        return _spm_lines()
    monkeypatch.setattr(wa, "_boq_lines_from_text", fake_lines)
    out = await wa._complete_price_list({"category": "Menu / Price List", "fields": {"supplier": "SPM"}},
                                        "Menu - SPM.pdf", _SPM_TEXT, "Menu - SPM.pdf")
    assert calls == ["supplier price list"] and len(out["fields"]["line_items"]) == 3
    assert out["fields"]["supplier"] == "SPM"
    catalogue = await wa._complete_price_list(
        {"category": "Menu / Price List", "fields": {}}, "Menu.pdf",
        "Mipolam Troplan — available colours, 2.0mm, 2m wide rolls, EN 13501-1 Bfl-s1", "Menu.pdf")
    assert catalogue["fields"] == {} and len(calls) == 1              # no prices → no LLM read


@pytest.mark.asyncio
async def test_learn_from_history_itemises_price_lists_and_stores_stock_sheets(db, monkeypatch):
    from vula.api import whatsapp as wa
    from vula.commerce import reread
    tid = "gerflor"
    db.tables["vula_filed_documents"] = [
        {"id": "p1", "tenant_id": tid, "category": "Menu / Price List", "filename": "Menu - SPM 20260828-0659.pdf",
         "file_url": "https://x/p1.pdf", "summary": "This document is a price list for SPM Wall Protection",
         "fields": {"supplier": "SPM"}, "created_at": "2026-08-28T07:00:00Z"},
        {"id": "s1", "tenant_id": tid, "category": "Programme / Schedule", "filename": "Programme 20260907-1232.pdf",
         "file_url": "https://x/s1.pdf", "doc_id": "d-s1", "fields": {},
         "summary": "This document is a new order proposal detailing various vinyl flooring products, their current stock on hand (SOH) in sqm"},
    ]
    texts = {"p1": _SPM_TEXT, "s1": _SOH_TEXT}

    async def fake_download(url):
        return b"%PDF"

    async def fake_text(tenant_id, row, path):
        return texts[row["id"]]

    async def fake_lines(text, filename, doc_kind="Bill of Quantities"):
        return _spm_lines()
    monkeypatch.setattr(reread, "_download", fake_download)
    monkeypatch.setattr(reread, "_file_text", fake_text)
    monkeypatch.setattr(wa, "_boq_lines_from_text", fake_lines)
    monkeypatch.setattr(reread, "candidates", lambda t, limit=200: [])

    st = await reread.learn_from_history(tid)
    assert st["step"] == "done" and st["failed"] == 0
    assert (st["price_lists_read"], st["stock_sheets"], st["priced_lines"]) == (1, 1, 3)
    sheet = db.tables["vula_stock_sheets"][0]
    assert sheet["doc_id"] == "d-s1" and sheet["as_at"] and len(sheet["rows"]) >= 5
    rates = {r["norm_key"]: r for r in db.tables["vula_price_observations"]}
    assert any(r["unit_price_cents"] == 61250 and r["supplier"] == "SPM" for r in rates.values())
    # a second run reads nothing again
    db.tables["vula_stock_sheets"][0]["tenant_id"] = tid
    again = await reread.learn_from_history(tid)
    assert (again["price_lists_read"], again["stock_sheets"]) == (0, 0)

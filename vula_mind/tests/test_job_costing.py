"""Job costing from the bank: is each project making its cost-plus fee? (2026-09-28, DIGG)

Ian: "Judy usually asks 10% on top of cost… Vula should see the operational costs and labour
costs based on transactions." Her statement breakdown (10 Jul – 12 Sep) showed HPC001 receiving
R1,514,438.13 in payment certificates against R1,521,336.97 paid out on it — a cost-plus-10%
job running negative. A statement sheet already categorised by project/trade is imported as it
is, allocations are learned for the next PDF statement, and job_costing turns it into received /
cost by trade / fee target / overhead share / profit; the owner agent quotes it.
"""
from datetime import datetime

import pytest

from tests.test_price_book import FakeDB as _BaseDB, _Q as _BaseQ
from vula.commerce import allocation, job_costing, service, statement_sheet

TID = "digg-demo"


class _Q(_BaseQ):
    def upsert(self, payload, on_conflict=None, **_k):
        self.op, self.payload, self.conflict = "upsert", payload, on_conflict
        return self

    def gte(self, col, val):
        self.filters.append((col, _Gte(val)))
        return self

    def lte(self, col, val):
        self.filters.append((col, _Lte(val)))
        return self

    def is_(self, col, val):
        self.filters.append((col, None))
        return self

    def execute(self):
        if self.op == "upsert":
            rows = self.db.tables.setdefault(self.table, [])
            keys = [k for k in (self.conflict or "").split(",") if k]
            for r in rows:
                if keys and all(r.get(k) == self.payload.get(k) for k in keys):
                    r.update(self.payload)
                    return type("R", (), {"data": [dict(r)]})()
            self.op = "insert"
        return super().execute()


class _Gte:
    def __init__(self, v):
        self.v = v

    def __eq__(self, other):
        if other is None:
            return False
        if isinstance(self.v, (int, float)):
            return float(other) >= self.v
        return str(other) >= self.v


class _Lte:
    def __init__(self, v):
        self.v = v

    def __eq__(self, other):
        if other is None:
            return False
        if isinstance(self.v, (int, float)):
            return float(other) <= self.v
        return str(other) <= self.v


class FakeDB(_BaseDB):
    def table(self, name):
        return _Q(self, name)


@pytest.fixture()
def db(monkeypatch):
    fake = FakeDB({"vula_projects": [{"tenant_id": TID, "name": "HPC Bokaap", "number": None}],
                   "vula_filed_documents": [{"tenant_id": TID, "project": "Sporty – Phase 2"}]})
    monkeypatch.setattr(service, "_client", lambda: fake)
    monkeypatch.setattr("vula.commerce.accounting.ensure_chart", lambda tid: [
        {"code": c, "vat_treatment": "none"} for c in
        ("sales", "other_income", "cost_of_sales", "casual_labour", "owner_drawings", "fuel",
         "bank_charges", "insurance", "utilities", "other_expense", "professional_fees")])
    monkeypatch.setattr("vula.commerce.accounting.is_vat_registered", lambda tid: False)
    return fake


HEAD = ["Transaction date", "Bank posted date", "Description", "Payee", "Money In", "Money Out",
        "Balance (as posted)", "Category", "Sub-category", "Action flag", "Adjustment note"]
ROWS = [  # real lines from DIGG's breakdown (amounts as on the sheet)
    (datetime(2026, 7, 13), "SEND Andries Sigauque", None, 2500, "HPC001 project", "Labour, wages & subcontractors"),
    (datetime(2026, 7, 13), "Hpc- Screen Blocks", None, 44252, "HPC001 project", "Materials & specialist items"),
    (datetime(2026, 7, 14), "HPC", None, 13250.62, "HPC001 project", "HPC - unspecified"),
    (datetime(2026, 7, 20), "HPC DOORS", None, 27292.48, "HPC001 project", "Doors, joinery & ironmongery"),
    (datetime(2026, 7, 22), "NELITHO WAGES", None, 12500, "HPC001 project", "Labour, wages & subcontractors"),
    (datetime(2026, 7, 29), "NELITHO WAGES", None, 13000, "HPC001 project", "Labour, wages & subcontractors"),
    (datetime(2026, 7, 13), "SPORTY CEILINGS", None, 50000, "Sporty.TV project", "Sporty.TV Phase 2 trades"),
    (datetime(2026, 7, 18), "HPC-PC4", 175916.13, None, "Income", "HPC001 payment certificate"),
    (datetime(2026, 7, 25), "Reversed payment", 7000, None, "Income", "Reversed payment (returned)"),
    (datetime(2026, 7, 13), "Home", None, 5000, "Owner drawings", "Drawings & household transfers"),
    (datetime(2026, 7, 14), "insurance payment.", None, 3000, "Overheads", "Insurance"),
    (datetime(2026, 7, 16), "FUJI EXPRESS", None, 288, "Travel & vehicle", "Fuel, parking & travel"),
    (datetime(2026, 7, 11), "BWH", None, 3274.2, "Materials (card)", "Building merchants & hardware on card"),
]


def _sheet(path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.title = "Summary"
    wb.active.append(["DIGG — Account 63214254607"])
    ws = wb.create_sheet("Transactions")
    ws.append(HEAD)
    for d, desc, cin, cout, cat, sub in ROWS:
        ws.append([d, d, desc, desc, cin, cout, 0, cat, sub, None, None])
    wb.save(path)
    return path


def test_the_sheet_is_read_with_its_categories(tmp_path):
    lines = statement_sheet.parse(_sheet(tmp_path / "DIGG_Statement_Breakdown.xlsx"))
    assert len(lines) == len(ROWS)
    labels = {statement_sheet.project_label(li) for li in lines}
    assert labels == {"HPC001", "Sporty.TV", None}
    cert = next(li for li in lines if li["description"] == "HPC-PC4")
    assert (cert["direction"], cert["amount_cents"]) == ("in", 17591613)
    assert statement_sheet.account_for(cert, "HPC Bokaap") == "sales"
    rev = next(li for li in lines if li["description"] == "Reversed payment")
    assert statement_sheet.account_for(rev, None) == "other_income"


def test_preview_suggests_the_tenants_project_names(db, tmp_path):
    out = statement_sheet.preview(TID, _sheet(tmp_path / "s.xlsx"))
    got = {p["label"]: p["suggested"] for p in out["projects"]}
    assert got == {"HPC001": "HPC Bokaap", "Sporty.TV": "Sporty – Phase 2"}


def _import(db, tmp_path):
    path = _sheet(tmp_path / "s.xlsx")
    return statement_sheet.import_sheet(TID, path, {"HPC001": "HPC Bokaap", "Sporty.TV": "Sporty – Phase 2"})


def test_import_keeps_allocations_dedupes_and_learns(db, tmp_path):
    first = _import(db, tmp_path)
    again = _import(db, tmp_path)
    assert first["saved"] == again["saved"] == len(ROWS)
    rows = db.tables["commerce_bank_transactions"]
    assert len(rows) == len(ROWS)                                   # re-import dedupes
    doors = next(r for r in rows if r["description"] == "HPC DOORS")
    assert (doors["project"], doors["trade"], doors["account_code"], doors["categorized_by"]) == \
        ("HPC Bokaap", "Doors, joinery & ironmongery", "cost_of_sales", "owner")
    wages = next(r for r in rows if r["description"] == "NELITHO WAGES")
    assert wages["account_code"] == "casual_labour"
    # learned: the next statement's "NELITHO WAGES" and any "HPC …" line land on HPC
    rules = allocation.load_rules(TID)
    assert allocation.suggest(rules, "NELITHO WAGES") == ("HPC Bokaap", "Labour, wages & subcontractors")
    assert allocation.suggest(rules, "HPC PLAMMING")[0] == "HPC Bokaap"
    assert allocation.suggest(rules, "BWH") == (None, None)          # never allocated → not guessed



def test_a_single_allocation_is_not_yet_a_rule(db):
    allocation.learn(TID, "SEND Andries Sigauque", "HPC Bokaap", "Labour")
    assert allocation.suggest(allocation.load_rules(TID), "SEND Andries Sigauque") == (None, None)
    allocation.learn(TID, "SEND Andries Sigauque", "HPC Bokaap", "Labour")
    assert allocation.suggest(allocation.load_rules(TID), "SEND Andries Sigauque") == ("HPC Bokaap", "Labour")


def test_job_costing_hpc_is_below_its_fee(db, tmp_path):
    _import(db, tmp_path)
    res = job_costing.costing(TID)
    hpc = next(p for p in res["projects"] if p["project"] == "HPC Bokaap")
    cost = 250000 + 4425200 + 1325062 + 2729248 + 1250000 + 1300000
    assert (hpc["received_cents"], hpc["cost_cents"]) == (17591613, cost)
    assert hpc["fee_pct"] == 10.0 and hpc["target_received_cents"] == cost + round(cost * 0.10)
    assert hpc["fee_earned_cents"] == 17591613 - cost
    assert hpc["unallocated_trade_cents"] == 1325062              # "HPC - unspecified"
    labour = next(t for t in hpc["trades"] if t["trade"] == "Labour, wages & subcontractors")
    assert labour["cents"] == 250000 + 1250000 + 1300000
    # overheads (drawings + insurance + fuel = R8,288) split by July project spend
    assert res["overheads_cents"] == 828800
    sporty = next(p for p in res["projects"] if p["project"] == "Sporty – Phase 2")
    assert hpc["overhead_share_cents"] + sporty["overhead_share_cents"] == 828800
    assert hpc["overhead_share_cents"] == round(828800 * cost / (cost + 5000000))
    assert sporty["status"] == "loss"                              # R50k out, nothing in
    assert res["unallocated_project_spend_cents"] == 327420       # BWH card materials


def test_fee_per_project_changes_the_target(db, tmp_path):
    _import(db, tmp_path)
    job_costing.set_fee(TID, "HPC Bokaap", 15)
    hpc = next(p for p in job_costing.costing(TID)["projects"] if p["project"] == "HPC Bokaap")
    assert hpc["fee_pct"] == 15.0 and hpc["fee_target_cents"] == round(hpc["cost_cents"] * 0.15)


def test_the_agent_gets_quotable_text(db, tmp_path):
    _import(db, tmp_path)
    out = job_costing.project_profit(TID, "hpc")
    assert out["project"]["project"] == "HPC Bokaap"
    assert "received R175,916.13" in out["text"] and "cost + 10%" in out["text"]
    assert "not allocated to a trade" in out["text"].lower() or "isn't allocated to a trade" in out["text"]
    everything = job_costing.project_profit(TID)
    assert "Sporty – Phase 2" in everything["text"] and "Overheads R8,288.00" in everything["text"]


def test_price_advice_starts_from_what_was_paid(db, tmp_path, monkeypatch):
    from vula.commerce import price_book
    for i, (kind, price) in enumerate([("Quote / Estimate", 14000), ("Invoice", 16800), ("Invoice", 16800)]):
        price_book.record_from_document(TID, {"id": f"d{i}", "category": kind, "fields": {
            "supplier": "Teck Flooring", "date": f"2026-08-0{i + 1}",
            "line_items": [{"description": "Selflevelling screed 2mm", "quantity": 100,
                            "unit": "m2", "unit_price_cents": price}]}})
    monkeypatch.setattr(job_costing, "overhead_rate", lambda tid, days=90: 7.5)
    adv = job_costing.price_advice(TID, "screed", quantity=405, unit="m2")
    assert adv["cost_cents"] == 16800 and adv["price_cost_plus_cents"] == 18480
    assert adv["price_covering_overheads_cents"] == round(16800 * 1.075 * 1.10)
    assert adv["under_quoted_pct"] == 20.0 and adv["total_cost_plus_cents"] == 18480 * 405
    assert "leaves about 2.5% as profit" in adv["text"]
    assert "No price on file" in job_costing.price_advice(TID, "unobtainium")["message"]


@pytest.mark.asyncio
async def test_weekly_alert_names_the_losing_project_once(db, tmp_path, monkeypatch):
    _import(db, tmp_path)
    sent = []

    async def notify_team(tid, event, text, idem_key=None):
        sent.append((event, text, idem_key))
        return 1
    monkeypatch.setattr("vula.integrations.notify.notify_team", notify_team)
    monkeypatch.setattr("vula.integrations.notify._members", lambda tid: [
        {"name": "Judy", "role": "owner", "whatsapp": "27827077080", "notify": ["project_margin"]}])
    job_costing._last_alert.pop(TID, None)
    text = await job_costing.weekly_alert(TID)
    assert "Sporty – Phase 2: loss" in text and sent[0][0] == "project_margin"
    assert sent[0][2].startswith("project-check:")                     # stored send key
    assert await job_costing.weekly_alert(TID) is None                # once a week


@pytest.mark.asyncio
async def test_weekly_alert_is_not_resent_after_a_restart(db, tmp_path, monkeypatch):
    """2 Oct, Ian: "Judy is getting WhatsApps about a costing recoup" — the same 'R611,903.69 of
    materials/labour isn't allocated' check went to her 13 times in 4 days, once per deploy,
    because the once-a-week marker lived only in memory."""
    _import(db, tmp_path)
    claimed, delivered = set(), []

    async def send(to, text, tenant_id="", idem_key=None):
        if idem_key and idem_key in claimed:     # what _claim_outbound's DB key does
            return False
        if idem_key:
            claimed.add(idem_key)
        delivered.append(to)
        return True
    monkeypatch.setattr("vula.api.whatsapp._send_reply", send)
    monkeypatch.setattr("vula.integrations.notify._members", lambda tid: [
        {"name": "Judy", "role": "owner", "whatsapp": "27827077080", "notify": []}])
    for _restart in range(3):
        job_costing._last_alert.clear()          # a deploy wipes the process
        await job_costing.weekly_alert(TID)
    assert delivered == ["27827077080"]
    job_costing._last_alert.clear()
    await job_costing.weekly_alert(TID, force=True)                   # the manual trigger still sends
    assert len(delivered) == 2


@pytest.mark.asyncio
async def test_notify_team_keys_each_recipient(monkeypatch):
    from vula.integrations import notify
    keys = []

    async def send(to, text, tenant_id="", idem_key=None):
        keys.append(idem_key)
        return True
    monkeypatch.setattr("vula.api.whatsapp._send_reply", send)
    monkeypatch.setattr(notify, "_members", lambda tid: [
        {"whatsapp": "27820000001", "notify": ["low_stock"]},
        {"whatsapp": "27820000002", "notify": ["low_stock"]}])
    await notify.notify_team(TID, "low_stock", "x", idem_key="low-stock:2026-10-02")
    assert keys == ["low-stock:2026-10-02:27820000001", "low-stock:2026-10-02:27820000002"]
    keys.clear()
    await notify.notify_team(TID, "low_stock", "x")                    # ad-hoc alerts: no key
    assert keys == [None, None]


def test_identical_lines_on_one_day_are_both_kept(db, tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(HEAD)
    for _ in range(2):
        ws.append([datetime(2026, 7, 11), None, "PAY", "PAY", None, 500, 0, "Unallocated", "Unnamed payment", None, None])
    path = tmp_path / "dup.xlsx"
    wb.save(path)
    statement_sheet.import_sheet(TID, path)
    statement_sheet.import_sheet(TID, path)
    assert sorted(r["description"] for r in db.tables["commerce_bank_transactions"]) == ["PAY", "PAY (2)"]


# ── a shop has no projects (Ian, 2026-09-28: "Off the Hook is pure commerce… not projects") ──

def test_only_a_project_business_uses_projects(monkeypatch):
    from vula.api import tenants
    cfgs = {"digg-demo": {"modules": ["invoices", "projects", "finances"]},
            "off-the-hook": {"modules": ["products", "orders", "invoices"]}}
    monkeypatch.setattr(tenants, "get_config", lambda tid: cfgs.get(tid))
    assert tenants.uses_projects("digg-demo") is True
    assert tenants.uses_projects("off-the-hook") is False
    assert tenants.uses_projects("unknown") is True          # no config → old behaviour

    import core.skills.commerce_admin as ca
    monkeypatch.setattr(tenants, "enabled_modules", lambda tid: (cfgs.get(tid) or {}).get("modules", []))
    names = lambda tid: {t["function"]["name"] for t in ca._tools_for(tid)}   # noqa: E731
    assert {"project_profit", "price_advice"} <= names("digg-demo")
    assert "project_profit" not in names("off-the-hook") and "price_advice" in names("off-the-hook")


@pytest.mark.asyncio
async def test_no_weekly_project_check_for_a_shop(db, tmp_path, monkeypatch):
    _import(db, tmp_path)
    monkeypatch.setattr("vula.api.tenants.uses_projects", lambda tid: False)
    job_costing._last_alert.pop(TID, None)
    assert await job_costing.weekly_alert(TID) is None


# ── only the business's own bank statement becomes bank lines (2026-09-28) ─────

def test_supplier_statements_and_invoices_are_not_bank_statements():
    from vula.commerce.bank_rec import looks_like_own_bank_statement
    capitec = ("Capitec Bank  Account Statement  Opening Balance R156,069.73  "
               "13/07/2026 HPC DOORS -27,292.48  Closing Balance R1,968.69")
    fnb = "FIRST NATIONAL BANK  Cheque Account  Balance brought forward  NELITHO WAGES 12,500.00"
    supplier = ("SOLID CAPE (PTY) LTD  Statement of Account  Account: DIG003  Current 30 Days 60 Days "
                "90 Days  Amount Due R77,513.00")
    ar = "Accounts Receivable Statements  Customer: DIGG  Invoice No 113651  Balance R1,062.34"
    receipt = "Receipt No 2493-1665-2943  Total R23.00  Thank you"
    assert looks_like_own_bank_statement(capitec) and looks_like_own_bank_statement(fnb)
    assert not looks_like_own_bank_statement(supplier)
    assert not looks_like_own_bank_statement(ar)
    assert not looks_like_own_bank_statement(receipt)
    # 2026-10-01: Solid Cape's statement prints its own banking details and a balance — still
    # a supplier's statement, not DIGG's bank.
    solid_cape = ("Accounts Receivable Statements  SOLID CAPE (PTY) LTD  Customer: AWEH BELEKKER T/A DIGG  "
                  "Balance brought forward R0.00  INV06069 Sales Order 1,142.07  Balance due R369.54  "
                  "Banking details: First National Bank  Acc 62012345678  Branch 250655")
    assert not looks_like_own_bank_statement(solid_cape)
    from tests.test_fnb_statement import TEXT as real_fnb_layout
    assert looks_like_own_bank_statement(real_fnb_layout)     # the bank's own layout still passes


def test_the_sheet_can_replace_pdf_lines_for_its_dates_and_allocates_the_rest(db, tmp_path):
    rows = db.tables.setdefault("commerce_bank_transactions", [])
    rows += [
        {"id": "p1", "tenant_id": TID, "txn_date": "2026-07-20", "description": "HPC DOORS", "amount_cents": 2729248,
         "direction": "out", "source_file": "20 Jul 2026 - (Free).pdf", "match_status": "unmatched", "project": None},
        {"id": "p4", "tenant_id": TID, "txn_date": "2026-07-21", "description": "BWH x3 (separate lines)",
         "amount_cents": 109140, "direction": "out", "source_file": "20 Jul 2026 - (Free).pdf",
         "match_status": "unmatched", "project": None},
        {"id": "p2", "tenant_id": TID, "txn_date": "2026-07-18", "description": "matched one", "amount_cents": 100,
         "direction": "in", "source_file": "x.pdf", "match_status": "matched", "project": None},
        {"id": "p3", "tenant_id": TID, "txn_date": "2026-06-30", "description": "HPC SKIPS", "amount_cents": 560000,
         "direction": "out", "source_file": "30 Jun 2026 - (Free).pdf", "match_status": "unmatched", "project": None},
    ]
    path = _sheet(tmp_path / "s.xlsx")
    assert statement_sheet.preview(TID, path)["existing_lines_in_period"] == 2     # p1, p4
    out = statement_sheet.import_sheet(TID, path, {"HPC001": "HPC Bokaap"}, replace_existing=True)
    byid = {r["id"]: r for r in db.tables["commerce_bank_transactions"]}
    assert out["set_aside"] == 2 and byid["p4"]["match_status"] == "ignored"      # PDF-only line set aside
    # the identical line (same date, amount, description) simply becomes the sheet's line
    assert (byid["p1"]["match_status"], byid["p1"]["project"]) == ("unmatched", "HPC Bokaap")
    assert byid["p2"]["match_status"] == "matched"                     # matched work untouched
    assert byid["p3"]["project"] == "HPC Bokaap"                      # June line allocated by "hpc"
    assert out["also_allocated"] >= 1


def test_variations_are_counted_per_project(db, tmp_path, monkeypatch):
    _import(db, tmp_path)
    monkeypatch.setattr("vula.api.tenants.get_config", lambda tid: {"display_name": "DIGG"})
    db.tables.setdefault("vula_filed_documents", []).extend([
        {"id": "v1", "tenant_id": TID, "category": "Quote / Estimate", "project": "HPC Bokaap",
         "fields": {"supplier": None, "total_cents": 3433533, "labels": ["Variation — over BOQ"]}},
        {"id": "v2", "tenant_id": TID, "category": "Invoice", "project": "HPC Bokaap",
         "fields": {"supplier": "Storeplay", "total_cents": 580000, "labels": ["Variation — over BOQ", "Back-charge"]}},
        {"id": "v3", "tenant_id": TID, "category": "Invoice", "project": "HPC Bokaap",
         "fields": {"supplier": "Solid Cape", "total_cents": 999}},
    ])
    hpc = next(p for p in job_costing.costing(TID)["projects"] if p["project"] == "HPC Bokaap")
    assert hpc["variations"] == {"documents": 2, "claimed_cents": 3433533, "extra_cost_cents": 580000}
    assert "Variations over the BOQ: 2 document(s)" in job_costing.project_profit(TID, "hpc")["text"]


# ── DIGG's own wages are a shared running cost (Ian's default, 2026-09-28) ──────

def test_own_wages_are_overhead_and_never_a_project_clue(db, monkeypatch):
    monkeypatch.setattr("vula.api.tenants.get_config", lambda tid: {"display_name": "DIGG"})
    assert allocation.is_own_wages(TID, "DIGG WAGE")
    assert allocation.is_own_wages(TID, "DIGG SALARIES SEP")
    assert allocation.is_own_wages(TID, "SEND DIGG")
    assert not allocation.is_own_wages(TID, "NELITHO WAGES")          # a site worker
    # real FNB lines: "Digg Wage" is DIGG's reference on its site worker's pay → project labour
    assert not allocation.is_own_wages(TID, "FNB App Payment To Nelitho Wages Digg Wage")
    assert not allocation.is_own_wages(TID, "FNB App Payment To Hpc Cleaner Digg")
    assert allocation.is_own_wages(TID, "FNB App Payment To Digg Wage")
    assert not allocation.is_own_wages(TID, "DIGG Reimbursement HPC")  # names more than DIGG
    allocation.learn(TID, "DIGG WAGE", "HPC Bokaap", "Labour")
    allocation.learn(TID, "DIGG WAGE", "HPC Bokaap", "Labour")
    assert allocation.suggest(allocation.load_rules(TID), "DIGG WAGE") == (None, None)

    db.tables["commerce_bank_transactions"] = [
        {"id": "a", "tenant_id": TID, "txn_date": "2026-08-05", "description": "HPC DOORS",
         "amount_cents": 900000, "direction": "out", "account_code": "cost_of_sales", "project": "HPC Bokaap"},
        {"id": "b", "tenant_id": TID, "txn_date": "2026-08-06", "description": "DIGG WAGE",
         "amount_cents": 100000, "direction": "out", "account_code": "casual_labour", "project": None},
    ]
    res = job_costing.costing(TID)
    assert res["overheads_cents"] == 100000 and res["unallocated_project_spend_cents"] == 0
    assert res["projects"][0]["overhead_share_cents"] == 100000


# ── Costs paid from a private account count until they're paid back (Ian, 2026-10-02) ──

def test_unreimbursed_personal_claims_count_as_job_cost(db, monkeypatch):
    monkeypatch.setattr("vula.api.tenants.get_config", lambda tid: {"display_name": "DIGG"})
    db.tables["commerce_bank_transactions"] = [
        {"id": "a", "tenant_id": TID, "txn_date": "2026-08-05", "description": "HPC DOORS",
         "amount_cents": 900000, "direction": "out", "account_code": "cost_of_sales", "project": "HPC Bokaap"},
    ]
    claim = dict(tenant_id=TID, reimbursable=True, status="submitted", reimbursed_at=None,
                 paid_by_name="Judy Downing", account_code="cost_of_sales",
                 paid_with="personal", channel="statement")
    db.tables["commerce_expenses"] = [
        # real lines from Judy's ABSA statement
        {**claim, "id": "c1", "date": "2026-08-01", "description": "Porterfield payment (Judy ABSA)",
         "amount_cents": 45000, "project": "HPC Bokaap"},
        {**claim, "id": "c2", "date": "2026-08-08", "description": "Atlantic Electrical (Judy ABSA card)",
         "amount_cents": 377244, "project": None},
        # already paid back → the bank carries it, never counted twice
        {**claim, "id": "c3", "date": "2026-08-09", "description": "BWH Tableview (Judy ABSA card)",
         "amount_cents": 27250, "project": "HPC Bokaap", "status": "reimbursed",
         "reimbursed_at": "2026-09-01T00:00:00Z"},
        # business card → not owed to anyone
        {**claim, "id": "c4", "date": "2026-08-09", "description": "Build It", "amount_cents": 1000,
         "project": "HPC Bokaap", "reimbursable": False},
        # a WhatsApp receipt flagged reimbursable but really paid on DIGG's card (already in the
        # bank: FNB "POS Purchase Italtile Cape Town" R884, 29 Jul) — not proven personal
        {**claim, "id": "c5", "date": "2026-07-29", "channel": "whatsapp",
         "description": "Customer copy of a card transaction from Italtile Cape Town",
         "amount_cents": 88400, "project": "HPC Bokaap"},
    ]
    res = job_costing.costing(TID)
    hpc = res["projects"][0]
    assert hpc["cost_cents"] == 900000 + 45000
    assert res["unallocated_project_spend_cents"] == 377244
    assert "not yet paid back (R4,222.44)" in job_costing.cost_basis(TID)

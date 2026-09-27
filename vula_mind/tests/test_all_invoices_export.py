"""'All invoices' requests that name no supplier (Ian, 2026-09-27).

2026-09-26, digg-demo: "Please supply me all invoice and material with costs In a excel sheet"
named no supplier. The model searched its own words, fuzzy-matched a Teck Flooring invoice and an
unnamed WhatsApp photo, and the owner got those 2 documents in a workbook titled "invoice and
material with costs" — out of ~200 invoices from ~40 suppliers. Minutes later "all invoice for
gardens handyman" (the supplier is GARDENS HANDIMAN CENTRE) also missed, fell to the model, and
it replied with a fabricated email tool result.
"""
import io
from unittest.mock import AsyncMock, MagicMock, patch

import openpyxl
import pytest

from vula.commerce import service

TID = "digg-demo"


def _invoice(fn, supplier, cents, date, summary="", lines=None):
    fields = {"date": date}
    if supplier:
        fields["supplier"] = supplier
    if cents is not None:
        fields["total_cents"] = cents
    if lines:
        fields["line_items"] = lines
    return {"id": fn, "filename": fn, "category": "Invoice", "summary": summary,
            "created_at": f"{date}T09:00:00", "fields": fields}


ROWS = [
    _invoice("POS Account Sale 24-225537.pdf", "Gardens Handiman Centre", 94200, "2026-09-22",
             lines=[{"description": "SAND PER BAG ACC", "quantity": 15, "unit_price_cents": 3100,
                     "total_cents": 46500}]),
    _invoice("POS Account Refund 21-366230.pdf", "GARDENS HANDIMAN CENTRE", 95400, "2026-09-22"),
    _invoice("00090668.pdf", "SOLID CAPE (PTY) LTD", 1000000, "2026-09-19"),
    _invoice("00090614.pdf", None, None, "2026-07-20",
             summary="This document is a delivery note and tax invoice from SOLID CAPE (PTY) LTD to "
                     "AWEH BELEKKER T/A DIGG"),
    _invoice("Inv_51934_from_Teck_Flooring_Pty_Ltd_3792.pdf", "Teck Flooring (Pty) Ltd", 27843245, "2026-07-16"),
    _invoice("CITY BUILD IT 1.pdf", "CITY BI (PTY) LTD T/A (CITY) BUILD IT", 100000, "2026-07-24"),
    _invoice("CITY BUILD IT 2.pdf", "CITY BUILD IT", 104460, "2026-07-24"),
]


@pytest.mark.parametrize("q,expected", [
    ("Please supply me all invoice and material with costs In a excel sheet", True),
    ("all invoce and matirial costs", True),                 # typos of request words
    ("Please show all and do a full.break down in excel", True),
    ("all invoices from Teck", False),
    ("Please give me all invoice for gardens handyman with and Matirial Costs on Excel", False),
    ("Need all jack hammer invoices", False),
])
def test_names_no_party(q, expected):
    assert service.names_no_party(q) is expected


@pytest.mark.asyncio
async def test_all_invoices_answer_covers_every_supplier_and_sends_the_workbook():
    send = AsyncMock(return_value=True)
    with patch.object(service, "_all_invoice_rows", return_value=ROWS), \
         patch("vula.api.whatsapp._send_invoice_document", new=send):
        reply = await service.answer_all_invoices(
            TID, "Please supply me all invoice and material with costs In a excel sheet", phone="+2782")
    # 7 invoices, 4 suppliers after merging spelling variants
    assert reply.startswith("*All suppliers*: 7 invoices from 4 suppliers")
    total = 94200 - 95400 + 1000000 + 27843245 + 100000 + 104460
    assert f"R{total / 100:,.2f}" in reply and "after 1 refund" in reply
    assert "1 invoice has no amount on file" in reply
    assert "📎 Sent every invoice" in reply
    assert reply.index("TECK FLOORING") < reply.index("SOLID CAPE")        # biggest first
    (phone, data, filename, *_), _ = send.call_args
    assert filename.startswith("All_Suppliers_invoices_")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Invoices", "Line items", "Materials"]
    names = {r[0] for r in wb["Summary"].iter_rows(values_only=True)}
    assert {"TECK FLOORING", "SOLID CAPE", "GARDENS HANDIMAN CENTRE", "CITY BUILD IT"} <= names
    assert len([r for r in wb["Invoices"].iter_rows(values_only=True)][1:-1]) == 7


@pytest.mark.asyncio
async def test_all_invoices_without_excel_wording_replies_in_text_only():
    send = AsyncMock()
    with patch.object(service, "_all_invoice_rows", return_value=ROWS), \
         patch("vula.api.whatsapp._send_invoice_document", new=send):
        reply = await service.answer_all_invoices(TID, "show me all invoices", phone="+2782")
    assert "*All suppliers*" in reply and "📎" not in reply
    send.assert_not_awaited()


@pytest.mark.asyncio
@pytest.mark.parametrize("q", ["all invoices from Teck", "what did we spend with gardens",
                               "Please show all and do a full.break down in excel"])
async def test_all_invoices_leaves_other_questions_alone(q):
    fetch = MagicMock(return_value=ROWS)
    with patch.object(service, "_all_invoice_rows", fetch):
        assert await service.answer_all_invoices(TID, q) is None
    fetch.assert_not_called()


@pytest.mark.asyncio
async def test_answer_supplier_history_hands_an_unnamed_all_invoices_ask_to_the_all_suppliers_path():
    with patch.object(service, "list_suppliers", AsyncMock(return_value=[])), \
         patch.object(service, "_all_invoice_rows", return_value=ROWS), \
         patch.object(service, "find_filed_document", AsyncMock()) as find:
        reply = await service.answer_supplier_history(
            TID, "Please supply me all invoice and material with costs In a excel sheet")
    assert reply.startswith("*All suppliers*")
    find.assert_not_awaited()


# ── supplier names ───────────────────────────────────────────────────────────

_SUPPLIERS = [{"name": "GARDENS HANDIMAN CENTRE", "aliases": ["Jack Hammer"]},
              {"name": "SAND MASTERS", "aliases": []}]


@pytest.mark.asyncio
@pytest.mark.parametrize("q", [
    "Please give me all invoice for gardens handyman with and Matirial Costs on Excel",
    "all invoices from gardens handiman",
    "need all jack hammer invoices",
])
async def test_misspelt_or_aliased_supplier_resolves(q):
    with patch.object(service, "list_suppliers", AsyncMock(return_value=_SUPPLIERS)):
        names = await service._resolve_supplier_names(TID, q)
    assert names[0] == "GARDENS HANDIMAN CENTRE"
    assert service.mentions_supplier_name(q, names)


@pytest.mark.parametrize("q", ["how many bags of sand did we buy", "all tile invoices",
                               "gardens area invoices"])
def test_fuzzy_mentions_need_every_distinctive_word(q):
    assert not service.mentions_supplier_name(q, ["GARDENS HANDIMAN CENTRE", "SAND MASTERS"])


def test_supplier_label_never_falls_back_to_the_search_words():
    mixed = {"matches": [{"party": "Teck Flooring (Pty) Ltd", "amount": 278432.45},
                         {"party": None, "amount": None}]}
    assert service.supplier_label(mixed) == "TECK FLOORING"
    two = {"matches": [{"party": "Teck Flooring", "amount": 1.0}, {"party": "SOLID CAPE", "amount": 2.0}]}
    assert service.supplier_label(two) is None
    assert service.supplier_label({"resolved_supplier": "X", "matches": []}) == "X"


@pytest.mark.asyncio
async def test_direct_answer_on_a_generic_search_sends_every_supplier_not_two_fuzzy_hits():
    """The exact 2026-09-26 17:15 path: the model called find_document("invoice and material
    with costs") and got two unrelated hits."""
    from core.skills.email_admin import _direct_supplier_answer
    fuzzy = {"status": "found", "match_type": "resolved_via_knowledge_base", "total_matches": 2,
             "total_amount": "R278,432.45", "total_amount_cents": 27843245, "matches_with_amount": 1,
             "matches": [{"filename": "Inv_51934_from_Teck_Flooring_Pty_Ltd_3792.pdf",
                          "party": "Teck Flooring", "amount": 278432.45},
                         {"filename": "image-wamid.HBgLMjc4", "party": "Solid Cape", "amount": None}]}
    fuzzy["matches"][1]["amount"] = 5.0   # two different parties -> no single supplier
    with patch.object(service, "_all_invoice_rows", return_value=ROWS), \
         patch("vula.api.whatsapp._send_invoice_document", new=AsyncMock(return_value=True)):
        reply = await _direct_supplier_answer(
            "Please supply me all invoice and material with costs In a excel sheet", "find_document",
            {"query": "invoice and material with costs"}, fuzzy, tenant_id=TID, phone="+2782")
    assert reply.startswith("*All suppliers*: 7 invoices")
    assert "invoice and material with costs" not in reply

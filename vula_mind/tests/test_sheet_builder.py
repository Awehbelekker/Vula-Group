"""The sheet builder (31 Oct 2026): a spreadsheet the owner describes in words, built from the books,
previewed, checked to the cent, and kept as the business's own template. Made-up figures."""
import io
import shutil
import subprocess
import uuid
from datetime import date

import pytest

from vula.commerce import sheet_builder as sb

TODAY = date(2026, 10, 31)


class _Q:
    def __init__(self, db, table):
        self.db, self.table, self.filters, self._range, self._limit = db, table, [], None, None
        self._op, self._payload = "select", None

    def select(self, *_a, **_k):
        return self

    def eq(self, k, v):
        self.filters.append(lambda r: r.get(k) == v)
        return self

    def gte(self, k, v):
        self.filters.append(lambda r: str(r.get(k) or "") >= v)
        return self

    def lte(self, k, v):
        self.filters.append(lambda r: str(r.get(k) or "") <= v)
        return self

    def in_(self, k, vs):
        self.filters.append(lambda r: r.get(k) in vs)
        return self

    def is_(self, k, v):
        self.filters.append(lambda r: r.get(k) is None)
        return self

    def order(self, k, desc=False):
        self._order = (k, desc)
        return self

    def range(self, a, b):
        self._range = (a, b)
        return self

    def limit(self, n):
        self._limit = n
        return self

    def insert(self, row):
        self._op, self._payload = "insert", row
        return self

    def update(self, patch):
        self._op, self._payload = "update", patch
        return self

    def execute(self):
        rows = self.db.setdefault(self.table, [])
        if self._op == "insert":
            row = {"id": str(uuid.uuid4()), **self._payload}
            rows.append(row)
            return type("R", (), {"data": [row]})
        hit = [r for r in rows if all(f(r) for f in self.filters)]
        if self._op == "update":
            for r in hit:
                r.update(self._payload)
            return type("R", (), {"data": hit})
        if getattr(self, "_order", None):
            k, desc = self._order
            hit = sorted(hit, key=lambda r: str(r.get(k) or ""), reverse=desc)
        if self._range:
            hit = hit[self._range[0]:self._range[1] + 1]
        if self._limit:
            hit = hit[: self._limit]
        return type("R", (), {"data": [dict(r) for r in hit]})


class _DB:
    def __init__(self):
        self.tables = {}

    def table(self, name):
        return _Q(self.tables, name)


def _doc(tenant, supplier, number, day, total, vat, project="", lines=None, refund=False, superseded=False):
    return {"id": str(uuid.uuid4()), "tenant_id": tenant, "category": "Invoice", "project": project,
            "filename": ("refund " if refund else "") + f"{number}.pdf", "summary": "", "status": "filed",
            "doc_status": "superseded" if superseded else "current", "created_at": day + "T10:00:00",
            "fields": {"supplier": supplier, "document_number": number, "date": day, "total_cents": total,
                       "vat_cents": vat, "line_items": lines or []}}


@pytest.fixture
def books(monkeypatch):
    db = _DB()
    t = db.tables
    t["vula_filed_documents"] = [
        _doc("tenant-a", "Tile Co (Pty) Ltd", "T-1", "2026-09-03", 115_000, 15_000, "Seaview",
             [{"description": "Grout", "quantity": 4, "unit_price_cents": 25_000, "total_cents": 100_000}]),
        _doc("tenant-a", "Tile Co (Pty) Ltd", "T-2", "2026-10-10", 230_000, 30_000, "Oak Street"),
        _doc("tenant-a", "Tile Co (Pty) Ltd", "T-3", "2026-10-12", 23_000, 3_000, "Seaview", refund=True),
        _doc("tenant-a", "Tile Co (Pty) Ltd", "T-2b", "2026-10-10", 230_000, 30_000, "Oak Street", superseded=True),
        _doc("tenant-a", "Brick Ltd", "B-9", "2026-10-20", 57_500, 7_500, "Seaview"),
        _doc("tenant-a", "Tile Co (Pty) Ltd", "T-0", "2025-12-01", 11_500, 1_500, "Seaview"),
        _doc("tenant-b", "Tile Co (Pty) Ltd", "X-1", "2026-10-01", 999_900, 0, "Seaview"),
    ]
    t["commerce_bank_transactions"] = [
        {"tenant_id": "tenant-a", "txn_date": "2026-10-02", "description": "EFT Tile Co Seaview", "payee": "",
         "amount_cents": 115_000, "direction": "out", "account_code": "cost_of_sales", "project": "Seaview",
         "match_status": "matched"},
        {"tenant_id": "tenant-a", "txn_date": "2026-10-05", "description": "Client deposit", "payee": "Sam",
         "amount_cents": 500_000, "direction": "in", "account_code": "sales", "project": "Seaview",
         "match_status": "unmatched"},
        {"tenant_id": "tenant-a", "txn_date": "2026-10-06", "description": "Duplicate line", "payee": "",
         "amount_cents": 77_700, "direction": "out", "account_code": "", "project": "Seaview",
         "match_status": "ignored"},
    ]
    t["commerce_expenses"] = [
        {"tenant_id": "tenant-a", "date": "2026-10-03", "supplier": "Fuel Stop", "description": "Diesel",
         "amount_cents": 45_000, "vat_cents": 0, "account_code": "fuel", "project": "Seaview",
         "paid_with": "personal", "paid_by_name": "Thandi", "reimbursable": True, "status": "unpaid"},
        {"tenant_id": "tenant-a", "date": "2026-10-04", "supplier": "Fuel Stop", "description": "Duplicate",
         "amount_cents": 45_000, "vat_cents": 0, "account_code": "fuel", "project": "Seaview",
         "paid_with": "personal", "status": "cancelled"},
    ]
    t["commerce_invoices"] = [
        {"tenant_id": "tenant-a", "invoice_number": "INV-1", "doc_type": "invoice", "direction": "outbound",
         "customer_name": "Sam", "issue_date": "2026-10-01", "project": "Seaview", "subtotal_cents": 100_000,
         "vat_cents": 15_000, "total_cents": 115_000, "total_paid_cents": 50_000, "status": "partially_paid"},
        {"tenant_id": "tenant-a", "invoice_number": "CN-1", "doc_type": "credit_note", "direction": "outbound",
         "customer_name": "Sam", "issue_date": "2026-10-15", "project": "Seaview", "subtotal_cents": 10_000,
         "vat_cents": 1_500, "total_cents": 11_500, "total_paid_cents": 0, "status": "issued"},
        {"tenant_id": "tenant-a", "invoice_number": "BILL-1", "doc_type": "invoice", "direction": "inbound",
         "customer_name": "", "issue_date": "2026-10-01", "total_cents": 999, "status": "unpaid"},
        {"tenant_id": "tenant-a", "invoice_number": "INV-2", "doc_type": "invoice", "direction": "outbound",
         "customer_name": "Lee", "issue_date": "2026-10-20", "total_cents": 50_000, "status": "draft"},
    ]
    t["commerce_orders"] = [
        {"tenant_id": "tenant-a", "display_id": "OTH-1", "created_at": "2026-10-09T09:00:00", "customer_name": "Sam",
         "channel": "whatsapp", "status": "paid", "subtotal_cents": 20_000, "delivery_cents": 5_000, "total_cents": 25_000},
        {"tenant_id": "tenant-a", "display_id": "OTH-2", "created_at": "2026-10-09T10:00:00", "customer_name": "Lee",
         "channel": "whatsapp", "status": "cancelled", "subtotal_cents": 9_000, "delivery_cents": 0, "total_cents": 9_000},
    ]
    from vula.commerce import accounting, service
    monkeypatch.setattr(service, "_client", lambda: db)
    monkeypatch.setattr(accounting, "ensure_chart", lambda tid: [{"code": "fuel", "name": "Fuel"},
                                                                 {"code": "sales", "name": "Sales"},
                                                                 {"code": "cost_of_sales", "name": "Cost of sales"}])

    async def names(tid, q):
        return ["TILE CO (PTY) LTD", "Tile Co"] if "tile" in q.lower() else []
    monkeypatch.setattr(service, "_resolve_supplier_names", names)
    monkeypatch.setattr(service, "canonical_project", lambda tid, n: "Seaview" if "sea" in (n or "").lower() else n)
    return db


async def _rows(spec):
    spec = sb.validate(spec)
    resolved = await sb.resolve("tenant-a", spec)
    return spec, resolved, sb.rows_for("tenant-a", spec, resolved, today=TODAY)


# ── the spec ─────────────────────────────────────────────────────────────────

def test_the_spec_is_checked_and_says_what_is_allowed():
    with pytest.raises(sb.SpecError, match="Which data"):
        sb.validate({"source": "payroll"})
    with pytest.raises(sb.SpecError, match="can't have the column"):
        sb.validate({"source": "bank", "columns": ["date", "vat"]})
    with pytest.raises(sb.SpecError, match="grouped"):
        sb.validate({"source": "orders", "group_by": "project"})
    with pytest.raises(sb.SpecError, match="filtered by"):
        sb.validate({"source": "orders", "filters": {"supplier": "Tile Co"}})
    with pytest.raises(sb.SpecError, match="date"):
        sb.validate({"source": "bank", "filters": {"since": "last Tuesday"}})
    spec = sb.validate({"source": "supplier_invoices"})
    assert spec["columns"] == list(sb.SOURCES["supplier_invoices"].default_columns)
    assert spec["totals"] == ["excl_vat", "vat", "total"] and spec["group_by"] == "none"


def test_periods_are_worked_out_by_code():
    assert sb.period_range({"period": "this_year"}, TODAY) == ("2026-01-01", "2026-10-31")
    assert sb.period_range({"period": "last_year"}, TODAY) == ("2025-01-01", "2025-12-31")
    assert sb.period_range({"since": "2026-10-01", "until": "2026-10-31"}, TODAY) == ("2026-10-01", "2026-10-31")
    assert sb.period_range({"period": "all"}, TODAY) == (None, None)


# ── rows from the books ──────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_supplier_invoices_this_year_refunds_negative_superseded_and_other_tenants_left_out(books):
    spec, _, rows = await _rows({"source": "supplier_invoices", "filters": {"supplier": "tile", "period": "this_year"}})
    assert sorted(r["number"] for r in rows) == ["T-1", "T-2", "T-3"]          # not T-0 (2025), T-2b, X-1, B-9
    assert sb.totals(spec, rows) == {"excl_vat": 100_000 + 200_000 - 20_000, "vat": 15_000 + 30_000 - 3_000,
                                     "total": 115_000 + 230_000 - 23_000}


@pytest.mark.asyncio
async def test_a_project_named_loosely_is_the_projects_own_spelling(books):
    spec, resolved, rows = await _rows({"source": "supplier_invoices", "filters": {"project": "seaview ", "period": "all"}})
    assert resolved["project"] == "Seaview"
    assert sorted(r["number"] for r in rows) == ["B-9", "T-0", "T-1", "T-3"]


@pytest.mark.asyncio
async def test_lines_bank_expenses_invoices_and_orders_carry_the_right_signs(books):
    _, _, lines = await _rows({"source": "supplier_lines", "filters": {"supplier": "tile"}})
    assert [(r["description"], r["line_total"]) for r in lines] == [("Grout", 100_000)]

    spec, _, bank = await _rows({"source": "bank", "filters": {"project": "Seaview"}})
    assert {r["description"] for r in bank} == {"EFT Tile Co Seaview", "Client deposit"}     # ignored line out
    assert sb.totals(sb.validate({"source": "bank", "columns": ["date", "amount", "money_in", "money_out"]}), bank) \
        == {"amount": 500_000 - 115_000, "money_in": 500_000, "money_out": 115_000}
    assert {r["account"] for r in bank} == {"Cost of sales", "Sales"}

    _, _, exp = await _rows({"source": "expenses"})
    assert [r["description"] for r in exp] == ["Diesel"]                                      # cancelled left out
    _, _, cancelled = await _rows({"source": "expenses", "filters": {"status": "cancelled"}})
    assert [r["description"] for r in cancelled] == ["Duplicate"]

    spec, _, inv = await _rows({"source": "customer_invoices"})
    assert [r["number"] for r in inv] == ["INV-1", "CN-1"]                       # supplier bill and draft out
    assert sb.totals(spec, inv) == {"total": 115_000 - 11_500, "paid": 50_000, "owed": 65_000 - 11_500}

    spec, _, orders = await _rows({"source": "orders"})
    assert [r["number"] for r in orders] == ["OTH-1"]


# ── the preview and the file ─────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_the_preview_says_what_was_understood_and_the_totals(books):
    spec, resolved, rows = await _rows({"source": "supplier_invoices", "group_by": "month", "title": "Tile Co this year",
                                        "filters": {"supplier": "tile", "period": "this_year"}})
    text = sb.preview_text(spec, resolved, rows, TODAY)
    assert "Supplier: TILE CO (PTY) LTD" in text and "Period: 1 Jan 2026 – 31 Oct 2026" in text
    assert "3 row(s)" in text and "Total R3,220.00" in text
    assert "2026-09 R1,150.00" in text and "2026-10 R2,070.00" in text
    assert "build it" in text


@pytest.mark.asyncio
async def test_nothing_found_says_so(books):
    spec, resolved, rows = await _rows({"source": "supplier_invoices", "filters": {"supplier": "nobody at all"}})
    assert rows == [] and "Nothing in the books matches" in sb.preview_text(spec, resolved, rows, TODAY)


@pytest.mark.asyncio
@pytest.mark.parametrize("group", ["none", "month", "project"])
async def test_the_file_reads_back_to_the_books(books, group):
    from vula.commerce import xlsx_check
    spec, resolved, rows = await _rows({"source": "supplier_invoices", "group_by": group,
                                        "filters": {"supplier": "tile", "period": "this_year"}})
    data, want, _exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    assert set(want.values()) == {280_000, 42_000, 322_000}
    assert xlsx_check.mismatches(data, want) == []


@pytest.mark.asyncio
async def test_a_changed_file_is_caught(books):
    from openpyxl import load_workbook

    from vula.commerce import xlsx_check
    spec, resolved, rows = await _rows({"source": "supplier_invoices", "filters": {"supplier": "tile"}})
    data, want, _exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    wb = load_workbook(io.BytesIO(data))
    wb.active["G7"] = 1.0
    buf = io.BytesIO()
    wb.save(buf)
    assert xlsx_check.mismatches(buf.getvalue(), want)


@pytest.mark.asyncio
async def test_text_from_the_books_never_becomes_a_formula(books):
    from openpyxl import load_workbook
    books.tables["commerce_expenses"][0]["description"] = '=HYPERLINK("http://example.invalid","x")'
    spec, resolved, rows = await _rows({"source": "expenses", "columns": ["date", "description", "amount"]})
    data, _, _x = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    cell = load_workbook(io.BytesIO(data)).active["B7"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
@pytest.mark.asyncio
async def test_libreoffice_gives_the_same_totals(books, tmp_path):
    from openpyxl import load_workbook
    spec, resolved, rows = await _rows({"source": "supplier_invoices", "group_by": "month",
                                        "filters": {"supplier": "tile", "period": "this_year"}})
    data, want, _exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    src = tmp_path / "s.xlsx"
    src.write_bytes(data)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "o"), str(src)],
                   check=True, capture_output=True, timeout=120)
    ws = load_workbook(tmp_path / "o" / "s.xlsx", data_only=True).active
    assert {ref: round(ws[ref].value * 100) for ref in want} == want


# ── the conversation (the custom_sheet tool) ─────────────────────────────────

@pytest.fixture
def tool(books, monkeypatch):
    from core.skills import commerce_admin as ca
    from vula.commerce import service
    sent, stored = [], []
    monkeypatch.setattr(ca, "_member_access", lambda tid, phone: None)

    async def accent(tid):
        return None
    monkeypatch.setattr(service, "_brand_accent", accent)
    monkeypatch.setattr("vula.api.tenants.display_name", lambda tid: "Made Up Builders")

    async def send(to, data, name, caption, tid, content_type=None):
        sent.append((to, name, data))
        return True
    monkeypatch.setattr("vula.api.whatsapp._send_invoice_document", send)
    monkeypatch.setattr("vula.api.whatsapp._upload_to_storage", lambda b, p, d, c: stored.append(p) or f"u:{p}")
    monkeypatch.setattr("vula.storage_links.signed", lambda u, ttl: "signed:" + u)
    skill = ca.CommerceAdminSkill()
    ctx = {"tenant_id": "tenant-a", "phone": "27820000001", "message": "tile co this year in excel"}

    async def call(**args):
        return await skill._custom_sheet("tenant-a", args, ctx)
    call.sent, call.stored, call.ctx = sent, stored, ctx
    return call


SPEC = {"source": "supplier_invoices", "group_by": "month", "title": "Tile Co this year",
        "filters": {"supplier": "tile", "period": "this_year"}}


@pytest.mark.asyncio
async def test_design_then_build_sends_the_checked_file(tool):
    out = await tool(action="design", spec=SPEC)
    assert "here's what I understood" in out["reply_verbatim"]
    out = await tool(action="build")
    assert "checked against the books" in out["reply_verbatim"] and "Total R3,220.00" in out["reply_verbatim"]
    (to, name, data), = tool.sent
    assert to == "27820000001" and name.startswith("Tile_Co_this_year_") and data[:2] == b"PK"
    assert tool.stored[0].startswith("tenant-a/sheets/")


@pytest.mark.asyncio
async def test_a_change_is_a_new_preview_and_the_build_uses_the_latest(tool):
    await tool(action="design", spec=SPEC)
    await tool(action="design", spec={**SPEC, "filters": {**SPEC["filters"], "project": "Seaview"}})
    out = await tool(action="build")
    assert "2 row(s)" in out["reply_verbatim"]                     # T-1 and the T-3 refund, Seaview only


@pytest.mark.asyncio
async def test_if_the_books_change_after_the_preview_it_shows_it_again(tool, books):
    await tool(action="design", spec=SPEC)
    books.tables["vula_filed_documents"].append(
        _doc("tenant-a", "Tile Co (Pty) Ltd", "T-9", "2026-10-30", 11_500, 1_500, "Seaview"))
    out = await tool(action="build")
    assert out["reply_verbatim"].startswith("The books changed since the preview") and not tool.sent
    out = await tool(action="build")
    assert "4 row(s)" in out["reply_verbatim"] and tool.sent


@pytest.mark.asyncio
async def test_build_without_a_preview_asks_first(tool):
    out = await tool(action="build")
    assert out["status"] == "need_info" and not tool.sent


@pytest.mark.asyncio
async def test_save_use_and_list_a_template(tool):
    await tool(action="design", spec=SPEC)
    await tool(action="build")
    out = await tool(action="save", name="Tile Co monthly")
    assert "Saved as *Tile Co monthly*" in out["reply_verbatim"]
    out = await tool(action="list")
    assert "Tile Co monthly" in out["reply_verbatim"]
    out = await tool(action="use", name="tile co monthly", spec={"filters": {"since": "2026-10-01", "until": "2026-10-31"}})
    assert "Period: 1 Oct 2026 – 31 Oct 2026" in out["reply_verbatim"] and "2 row(s)" in out["reply_verbatim"]
    again = await tool(action="save", name="Tile Co monthly")
    assert "replacing the old one" in again["reply_verbatim"]


@pytest.mark.asyncio
async def test_a_staff_member_only_gets_data_their_access_allows(tool, monkeypatch):
    from core.skills import commerce_admin as ca
    monkeypatch.setattr(ca, "_member_access", lambda tid, phone: ["finances"])
    assert "Your access doesn't include orders" in (await tool(action="design", spec={"source": "orders"}))["message"]
    assert "reply_verbatim" in await tool(action="design", spec={"source": "bank"})
    monkeypatch.setattr(ca, "_member_access", lambda tid, phone: ["orders"])
    assert "Your access doesn't include bank lines" in (await tool(action="design", spec={"source": "bank"}))["message"]


@pytest.mark.asyncio
async def test_a_bad_column_comes_back_as_what_is_possible(tool):
    out = await tool(action="design", spec={"source": "orders", "columns": ["number", "vat"]})
    assert out["status"] == "need_info" and "Columns it can have" in out["message"]


def test_the_tool_is_for_finance_access_and_not_for_reps():
    from core.skills import commerce_admin as ca
    assert ca._TOOL_SCOPE["custom_sheet"] == "finances"
    assert "custom_sheet" not in {t["function"]["name"] for t in ca._tools_for("t", role="sales_rep")}


# ── the fixed export vs the builder ──────────────────────────────────────────

def test_a_described_layout_goes_to_the_builder_a_plain_export_does_not():
    from core.skills.base import wants_custom_layout
    assert wants_custom_layout("every Tile Co invoice this year by project, VAT split, per month in excel")
    assert wants_custom_layout("bank lines for Seaview in October in excel")
    assert not wants_custom_layout("Tile Co invoices in excel")
    assert not wants_custom_layout("maybe send the tile co invoices in excel")
    assert not wants_custom_layout("invoices per month")                       # no file asked for


@pytest.mark.asyncio
async def test_find_document_does_not_answer_a_described_layout_itself():
    from core.skills.commerce_admin import _direct_supplier_answer
    out = await _direct_supplier_answer("every Tile Co invoice this year per month in excel", "find_document",
                                        {"query": "tile co"}, {"matches": []}, tenant_id="tenant-a")
    assert out is None


# ── item 1: calculations, counts and averages, summaries, charts ─────────────

def test_a_calculation_only_uses_the_sheets_number_columns_and_plain_maths():
    for bad, why in [("total; drop table", "only use"), ("total ** 2", "couldn't read"), ("sum(total)", "isn't a number column"),
                     ("=1+1", "only use"), ("vat * total", "Money times money"), ("profit - vat", "isn't a number column"),
                     ("(total - vat", "isn't closed")]:
        with pytest.raises(sb.SpecError, match=why):
            sb.validate({"source": "supplier_invoices", "computed": [{"label": "X", "formula": bad}]})


def test_kinds_are_worked_out_and_inputs_are_kept_in_the_sheet():
    spec = sb.validate({"source": "supplier_invoices", "columns": ["date", "supplier", "total"],
                        "computed": [{"label": "Excl", "formula": "total - vat"},
                                     {"label": "Margin %", "formula": "(total - vat) / total"},
                                     {"label": "Per unit", "formula": "total / 2"}]})
    assert [c["kind"] for c in spec["computed"]] == ["money", "percent", "money"]
    assert "vat" in spec["columns"] and spec["columns"][-3:] == ["calc1", "calc2", "calc3"]
    assert spec["totals"] == ["total", "vat", "calc1", "calc3"]
    assert sb.validate(spec) == spec                                   # a saved spec checks out the same


@pytest.mark.parametrize("bad,why", [
    ({"chart": "bar"}, "needs a grouping"), ({"detail": False}, "needs a grouping"),
    ({"group_by": "month", "detail": False, "summary": ["average"]}, "need the rows"),
    ({"running": "supplier"}, "money column"), ({"group_by": "month", "chart": "radar"}, "bar, line or pie"),
    ({"summary": ["median"]}, "count, average"),
])
def test_extras_are_checked(bad, why):
    with pytest.raises(sb.SpecError, match=why):
        sb.validate({"source": "supplier_invoices", **bad})


MARGIN = {"source": "supplier_invoices", "group_by": "month", "filters": {"supplier": "tile", "period": "this_year"},
          "computed": [{"label": "Excl VAT", "formula": "total - vat"},
                       {"label": "Margin %", "formula": "(total - vat) / total"}]}


@pytest.mark.asyncio
async def test_calculated_columns_per_row_and_the_percentage_on_the_total(books):
    spec, resolved, rows = await _rows({**MARGIN, "columns": ["date", "number", "total"]})
    by = {r["number"]: r for r in rows}
    assert by["T-1"]["calc1"] == 100_000 and by["T-3"]["calc1"] == -20_000
    assert by["T-1"]["calc2"] == 0.8696                          # 1000 / 1150, to 4 places
    tot = sb.totals(spec, rows)
    assert tot["calc1"] == 280_000
    assert sb.ratio_on(spec, spec["computed"][1], tot) == 0.8696   # on the totals, not a sum of %
    text = sb.preview_text(spec, resolved, rows, TODAY)
    assert "Margin % on the total: 87.0%" in text and "Margin % = (total - vat) / total (as a %)" in text


@pytest.mark.asyncio
async def test_a_division_by_zero_is_a_blank_cell_not_an_error(books):
    from openpyxl import load_workbook

    from vula.commerce import xlsx_check
    books.tables["vula_filed_documents"].append(_doc("tenant-a", "Tile Co (Pty) Ltd", "T-Z", "2026-10-25", 0, 0, "Seaview"))
    spec, resolved, rows = await _rows({**MARGIN, "group_by": "none", "columns": ["date", "number", "total"]})
    zero = next(r for r in rows if r["number"] == "T-Z")
    assert zero["calc2"] is None
    data, want, exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    assert xlsx_check.mismatches(data, want, exact=exact) == []
    ws = load_workbook(io.BytesIO(data)).active
    row = next(c.row for c in ws["B"] if c.value == "T-Z")
    col = spec["columns"].index("calc2") + 1
    assert xlsx_check.value(ws, ws.cell(row=row, column=col).coordinate) == ""


@pytest.mark.asyncio
async def test_quantity_times_price_rounds_half_up_like_excel(books):
    books.tables["vula_filed_documents"][0]["fields"]["line_items"].append(
        {"description": "Sealer", "quantity": 2.5, "unit_price_cents": 12_345, "total_cents": 30_863})
    spec, resolved, rows = await _rows({"source": "supplier_lines", "filters": {"supplier": "tile"},
                                        "computed": [{"label": "Check", "formula": "quantity * unit_price"}]})
    sealer = next(r for r in rows if r["description"] == "Sealer")
    assert sealer["calc1"] == 30_863                              # 308.625 → 308.63
    from vula.commerce import xlsx_check
    data, want, exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    assert xlsx_check.mismatches(data, want, exact=exact) == []


@pytest.mark.asyncio
@pytest.mark.parametrize("group", ["none", "project"])
async def test_running_balance_ends_at_the_total(books, group):
    from vula.commerce import xlsx_check
    spec, resolved, rows = await _rows({"source": "bank", "group_by": group, "running": "amount",
                                        "columns": ["date", "description", "amount"]})
    assert [r["running"] for r in rows] == [-115_000, 385_000]
    data, want, exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    assert 385_000 in want.values() and xlsx_check.mismatches(data, want, exact=exact) == []


@pytest.mark.asyncio
async def test_count_average_lowest_and_highest_leave_out_headers_and_subtotals(books):
    from vula.commerce import xlsx_check
    spec, resolved, rows = await _rows({"source": "supplier_invoices", "group_by": "month",
                                        "columns": ["date", "number", "total"], "totals": ["total"],
                                        "summary": ["count", "average", "min", "max"],
                                        "filters": {"supplier": "tile", "period": "this_year"}})
    data, want, exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    assert 3.0 in exact.values()                                   # 3 invoices, not 3 + headers + subtotals
    assert any(abs(v - (1150 + 2300 - 230) / 3) < 1e-9 for v in exact.values() if v)
    assert {-23_000, 230_000} <= set(want.values())
    assert xlsx_check.mismatches(data, want, exact=exact) == []


@pytest.mark.asyncio
async def test_totals_per_supplier_only(books):
    from openpyxl import load_workbook

    from vula.commerce import xlsx_check
    spec, resolved, rows = await _rows({"source": "supplier_invoices", "group_by": "supplier", "detail": False})
    data, want, exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    assert xlsx_check.mismatches(data, want, exact=exact) == []
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A6"].value == "Supplier" and [ws["A7"].value, ws["A8"].value] == ["BRICK", "TILE CO"]
    assert "One line per supplier: 2 lines from 5 row(s)" in sb.preview_text(spec, resolved, rows, TODAY)


@pytest.mark.asyncio
async def test_a_chart_is_in_the_file_and_its_figures_read_back(books):
    import zipfile

    from vula.commerce import xlsx_check
    spec, resolved, rows = await _rows({**MARGIN, "chart": "bar"})
    data, want, exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    assert any(n.startswith("xl/charts/chart") for n in names)
    assert {115_000, 207_000} <= set(want.values())                # Sep and Oct totals in the chart data
    assert xlsx_check.mismatches(data, want, exact=exact) == []
    text = sb.preview_text(spec, resolved, rows, TODAY)
    assert "▇" in text and "With a bar chart of Total per month" in text


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
@pytest.mark.asyncio
async def test_libreoffice_agrees_with_every_checked_cell(books, tmp_path):
    from openpyxl import load_workbook
    books.tables["vula_filed_documents"].append(_doc("tenant-a", "Tile Co (Pty) Ltd", "T-Z", "2026-10-25", 0, 0, "Seaview"))
    spec, resolved, rows = await _rows({**MARGIN, "chart": "line", "summary": ["count", "average", "min", "max"]})
    data, want, exact = sb.workbook("Made Up Builders", spec, resolved, rows, today=TODAY)
    src = tmp_path / "s.xlsx"
    src.write_bytes(data)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "o"), str(src)],
                   check=True, capture_output=True, timeout=120)
    ws = load_workbook(tmp_path / "o" / "s.xlsx", data_only=True).active
    assert {ref: round(ws[ref].value * 100) for ref in want} == want
    for ref, v in exact.items():
        got = ws[ref].value
        assert (got in (None, "") if v is None else abs(float(got) - v) < 1e-6), (ref, got, v)


@pytest.mark.asyncio
async def test_design_and_build_with_calculations_and_a_chart(tool):
    out = await tool(action="design", spec={**MARGIN, "chart": "bar", "summary": ["average"]})
    assert "Margin % on the total" in out["reply_verbatim"]
    out = await tool(action="build")
    assert "checked against the books" in out["reply_verbatim"] and tool.sent

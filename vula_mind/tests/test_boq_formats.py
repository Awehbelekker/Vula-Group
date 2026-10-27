"""Other QS layouts, 27 Oct.

Ian: "Vula must be able to import other BOQ templates if the tenant QS uses a different format."
The reader knew DIGG's and HPC's headings; a bill headed in Afrikaans, with a heading that only
contains the word, a JBCC "R | c" split, no heading row at all, decimal commas or an old .xls fell
through to the model's 6,000-character read. All figures here are made up.
"""
from pathlib import Path

from openpyxl import Workbook

from vula.ingestion import boq_sheet

FIXTURES = Path(__file__).parent / "fixtures"


def _sheet(path, rows, title="Bill"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def _summary(lines):
    return [(li["section"], li["description"], li["quantity"], li["unit"], li["unit_price_cents"],
             li["total_cents"]) for li in lines]


def test_an_afrikaans_bill_is_read(tmp_path):
    p = _sheet(tmp_path / "Hoeveelheidslys.xlsx", [
        ["Item Nr", "Beskrywing", "Eenheid", "Hoeveelheid", "Tarief", "Bedrag"],
        [None, "Messelwerk"],
        ["1", "Halwe steen mure", "m²", 12, 180, 2160],
        ["2", "Lateie", "m", 3, 95.5, 286.5],
        [None, "Teëlwerk"],
        ["3", "Vloerteëls", "m²", 20, 220, 4400],
    ])
    assert _summary(boq_sheet.parse(p)) == [
        ("Messelwerk", "Halwe steen mure", 12, "m²", 18000, 216000),
        ("Messelwerk", "Lateie", 3, "m", 9550, 28650),
        ("Teëlwerk", "Vloerteëls", 20, "m²", 22000, 440000)]


def test_headings_that_only_contain_the_word_are_read(tmp_path):
    p = _sheet(tmp_path / "Tender bill.xlsx", [
        ["Ref", "Description of Item / Work", "Unit of Msr", "Quantity (m²)", "Unit Rate Excl VAT",
         "Total Amount Excl VAT", "Previous Amount Certified"],
        ["A1", "Screed to floors", "m²", 40, 85, 3400, 1000],
        ["A2", "Skirting", "m", 30, 45, 1350, 0],
    ])
    lines = boq_sheet.parse(p)
    assert [(li["description"], li["quantity"], li["unit_price_cents"], li["total_cents"]) for li in lines] == [
        ("Screed to floors", 40, 8500, 340000), ("Skirting", 30, 4500, 135000)]


def test_a_jbcc_bill_with_rand_and_cents_columns_adds_the_cents(tmp_path):
    p = _sheet(tmp_path / "JBCC bill.xlsx", [
        ["BILL No. 3 – CONCRETE WORK"],
        ["Item", "Description", "Unit", "Qty", "Rate", None, "Amount", None],
        [None, None, None, None, "R", "c", "R", "c"],
        ["3.1", "Concrete in surface beds", "m³", 4, 1450, 75, 5803, 0],
        ["3.2", "Formwork to edges", "m", 12, 62, 50, 750, 0],
    ])
    lines = boq_sheet.parse(p)
    assert [(li["description"], li["unit_price_cents"], li["total_cents"]) for li in lines] == [
        ("Concrete in surface beds", 145075, 580300), ("Formwork to edges", 6250, 75000)]


def test_a_two_row_heading_is_read(tmp_path):
    p = _sheet(tmp_path / "Two row heading.xlsx", [
        ["Item", "Work", "Unit", "Quantity", "Rate", "Amount"],
        [None, "description", None, None, "R", "R"],
        ["1", "Paint walls", "m²", 100, 35, 3500],
    ])
    assert [(li["description"], li["total_cents"]) for li in boq_sheet.parse(p)] == [("Paint walls", 350000)]


def test_a_bill_with_no_heading_row_is_read_from_its_figures(tmp_path):
    p = _sheet(tmp_path / "QS export.xlsx", [
        ["Project: Sample Clinic", None, None, None, None, None],
        [None, "Electrical", None, None, None, None],
        ["E1", "Supply and fit double plug socket", "No", 8, 650, 5200],
        ["E2", "LED downlight complete", "No", 12, 420, 5040],
        ["E3", "Distribution board, 24-way", "Item", 1, 7800, 7800],
        [None, "Plumbing", None, None, None, None],
        ["P1", "uPVC soil pipe, 110 mm", "m", 15, 310.5, 4657.5],
        [None, "Total carried to summary", None, None, None, 22697.5],
    ])
    assert _summary(boq_sheet.parse(p)) == [
        ("Electrical", "Supply and fit double plug socket", 8, "No", 65000, 520000),
        ("Electrical", "LED downlight complete", 12, "No", 42000, 504000),
        ("Electrical", "Distribution board, 24-way", 1, "Item", 780000, 780000),
        ("Plumbing", "uPVC soil pipe, 110 mm", 15, "m", 31050, 465750)]


def test_a_sheet_of_figures_that_are_not_a_bill_is_left_alone(tmp_path):
    p = _sheet(tmp_path / "Cash flow.xlsx", [
        ["Week", 1, 2, 3],
        ["Income", 5000, 7000, 6500],
        ["Costs", 3000, 4100, 3900],
        ["Net", 2000, 2900, 2600],
    ])
    assert boq_sheet.parse(p) == []


def test_decimal_commas_and_bracketed_negatives_read_as_numbers():
    assert boq_sheet._num("1 234,56") == 1234.56
    assert boq_sheet._num("R 12 500,00") == 12500.0
    assert boq_sheet._num("1.234,56") == 1234.56
    assert boq_sheet._num("1,250") == 1250.0
    assert boq_sheet._num("(1 500.00)") == -1500.0
    assert boq_sheet._num("ZAR 90") == 90.0


def test_a_csv_with_decimal_commas_is_read(tmp_path):
    p = tmp_path / "bill.csv"
    p.write_text('Description,Unit,Qty,Rate,Amount\n"Ceiling board",m²,"12,5","95,00","1 187,50"\n',
                 encoding="utf-8")
    lines = boq_sheet.parse(p)
    assert [(li["quantity"], li["unit_price_cents"], li["total_cents"]) for li in lines] == [(12.5, 9500, 118750)]


def test_an_old_xls_workbook_is_read():
    lines = boq_sheet.parse(FIXTURES / "qs_bill_sample.xls")
    assert _summary(lines) == [
        ("Masonry", "Half brick walls", 10, "m²", 18000, 180000),
        ("Masonry", "Precast lintels", 4, "m", 9500, 38000),
        ("Tiling", "Floor tiles", 25, "m²", 22000, 550000)]


def test_the_known_layouts_still_read_the_same(tmp_path):
    p = _sheet(tmp_path / "Standard.xlsx", [
        ["Item", "Description", "Unit", "Qty", "Rate", "Amount"],
        ["1", "Excavate trenches", "m³", 6, 250, 1500],
    ])
    assert _summary(boq_sheet.parse(p)) == [("Bill", "Excavate trenches", 6, "m³", 25000, 150000)]


def test_a_quote_tab_with_no_headings_is_not_added_to_a_bill_already_read(tmp_path):
    """HPC's PC No. 10 carries the joiner's proforma on its own tab ("QTY | Rate | Amount", no
    description heading). The bill holds joinery as one item; reading the quote too counts it twice."""
    wb = Workbook()
    bill = wb.active
    bill.title = "Bill"
    for r in [["Item", "Description", "Unit", "Qty", "Rate", "Amount"],
              ["1", "Joinery as per quote", "Item", 1, 90000, 90000]]:
        bill.append(r)
    quote = wb.create_sheet("Joinery quote")
    for r in [["JOINER CO", "QTY", "Rate", "Amount"],
              ["Reception desk", 1, 40000, 40000], ["Vanity units", 3, 10000, 30000],
              ["Kitchen", 1, 20000, 20000]]:
        quote.append(r)
    p = tmp_path / "Valuation.xlsx"
    wb.save(p)
    assert [li["description"] for li in boq_sheet.parse(p)] == ["Joinery as per quote"]

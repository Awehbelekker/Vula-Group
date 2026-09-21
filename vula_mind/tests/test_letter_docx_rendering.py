"""Tests for render_letter_docx and draft_letter's output_format opt-in (2026-09-18). No .docx
generation capability existed anywhere in this codebase before this — python-docx was a
dependency used only for READING inbound .docx files.
"""
import io
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from vula.commerce.pdf import render_letter_docx, _markdown_to_docx, _add_inline_runs


def _open(docx_bytes: bytes):
    from docx import Document
    return Document(io.BytesIO(docx_bytes))


# ── render_letter_docx: basic shape ────────────────────────────────────────────────

def test_returns_valid_docx_bytes():
    docx_bytes = render_letter_docx(tenant_id="digg-demo", body="Hello world.", doc_label="Letter")
    assert isinstance(docx_bytes, bytes)
    assert docx_bytes[:2] == b"PK"  # .docx is a zip archive
    doc = _open(docx_bytes)
    assert any("Hello world" in p.text for p in doc.paragraphs)


def test_doc_label_and_date_present():
    docx_bytes = render_letter_docx(tenant_id="digg-demo", body="Body.", doc_label="Fee Proposal",
                                    issue_date="2026-09-18")
    text = "\n".join(p.text for p in _open(docx_bytes).paragraphs)
    assert "Fee Proposal" in text
    assert "2026-09-18" in text


def test_recipient_and_subject_rendered():
    docx_bytes = render_letter_docx(
        tenant_id="digg-demo", body="Body.", recipient="Mr R Downing\nCape Town",
        subject="Bokaap Renovation")
    text = "\n".join(p.text for p in _open(docx_bytes).paragraphs)
    assert "Mr R Downing" in text and "Cape Town" in text
    assert "Re: Bokaap Renovation" in text


def test_no_logo_falls_back_to_tenant_name_heading():
    docx_bytes = render_letter_docx(tenant_id="digg-demo", body="Body.",
                                    tenant_profile={"name": "DIGG Architects"})
    doc = _open(docx_bytes)
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "DIGG Architects" in headings


# ── sign-off + signature (with branding.get fallback and explicit-param override) ─

def test_sign_off_rendered_when_given():
    docx_bytes = render_letter_docx(tenant_id="digg-demo", body="Body.",
                                    sign_off="Kind regards,\nJudy Downing")
    text = "\n".join(p.text for p in _open(docx_bytes).paragraphs)
    assert "Kind regards," in text and "Judy Downing" in text


def test_signature_name_rendered_bold():
    docx_bytes = render_letter_docx(tenant_id="digg-demo", body="Body.",
                                    signature_name="Judy Downing")
    doc = _open(docx_bytes)
    match = next(p for p in doc.paragraphs if p.text == "Judy Downing")
    assert match.runs[0].bold is True


def _tiny_png_bytes() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (10, 10), color="white").save(buf, format="PNG")
    return buf.getvalue()


def test_signature_url_download_attempted_and_embedded_on_success():
    fake_png = _tiny_png_bytes()
    with patch("vula.commerce.pdf._download_image_bytes", return_value=fake_png) as mock_dl:
        docx_bytes = render_letter_docx(tenant_id="digg-demo", body="Body.",
                                        signature_url="https://storage.example/sig.png")
    mock_dl.assert_any_call("https://storage.example/sig.png")
    doc = _open(docx_bytes)
    assert len(doc.inline_shapes) >= 1


def test_signature_download_failure_never_breaks_the_document():
    with patch("vula.commerce.pdf._download_image_bytes", return_value=None):
        docx_bytes = render_letter_docx(tenant_id="digg-demo", body="Body.",
                                        signature_url="https://storage.example/broken.png",
                                        sign_off="Kind regards,")
    assert isinstance(docx_bytes, bytes)
    assert docx_bytes[:2] == b"PK"


# ── markdown walker: headings, bullets, numbered lists, inline formatting ────────

def test_markdown_headings_and_bullets():
    from docx import Document
    doc = Document()
    _markdown_to_docx(doc, "# Title\n\n## Subheading\n\n- First\n- Second\n\n1. One\n2. Two\n\nPlain paragraph.")
    styles = [(p.text, p.style.name) for p in doc.paragraphs]
    assert ("Title", "Heading 2") in styles
    assert ("Subheading", "Heading 3") in styles
    assert ("First", "List Bullet") in styles
    assert ("Second", "List Bullet") in styles
    assert ("One", "List Number") in styles
    assert ("Two", "List Number") in styles
    assert ("Plain paragraph.", "Normal") in styles


def test_markdown_multiline_paragraph_joined_on_blank_line_boundary():
    from docx import Document
    doc = Document()
    _markdown_to_docx(doc, "Line one\nLine two continues.\n\nSeparate paragraph.")
    texts = [p.text for p in doc.paragraphs]
    assert "Line one Line two continues." in texts
    assert "Separate paragraph." in texts


def test_inline_bold_and_italic_produce_separate_runs():
    from docx import Document
    doc = Document()
    p = doc.add_paragraph()
    _add_inline_runs(p, "**Fee:** R85,000, *payable monthly*, plain text.")
    assert p.text == "Fee: R85,000, payable monthly, plain text."
    bold_runs = [r for r in p.runs if r.bold]
    italic_runs = [r for r in p.runs if r.italic]
    assert any(r.text == "Fee:" for r in bold_runs)
    assert any(r.text == "payable monthly" for r in italic_runs)


def test_plain_text_with_no_markdown_markers_is_one_run():
    from docx import Document
    doc = Document()
    p = doc.add_paragraph()
    _add_inline_runs(p, "Nothing special here.")
    assert p.text == "Nothing special here."


# ── draft_letter: output_format branching ──────────────────────────────────────────

async def _run_draft_letter(output_format=None):
    from core.skills.draft_admin import draft_letter
    args = {"document_type": "appointment_letter", "brief": "Appoint the contractor."}
    if output_format is not None:
        args["output_format"] = output_format
    with (
        patch("core.skills.draft_admin._fee_proposal_gaps", return_value=[]),
        patch("vula.api.draft._retrieve_context", new=AsyncMock(return_value=("", 0))),
        patch("vula.api.draft._generate_document", new=AsyncMock(return_value=("Content.", "test-model"))),
        patch("core.skills.draft_admin._resolve_sign_off", new=AsyncMock(return_value="Kind regards,")),
        patch("vula.commerce.service.get_invoice_settings", new=AsyncMock(return_value={})),
        patch("vula.commerce.pdf.render_letter_pdf", return_value=b"%PDF-fake") as mock_pdf,
        patch("vula.commerce.pdf.render_letter_docx", return_value=b"PK-fake") as mock_docx,
        patch("vula.api.whatsapp._send_invoice_document", new=AsyncMock(return_value=True)) as mock_send,
    ):
        result = await draft_letter(args, "digg-demo", "27827077080")
    return result, mock_pdf, mock_docx, mock_send


@pytest.mark.asyncio
async def test_default_output_is_pdf():
    result, mock_pdf, mock_docx, mock_send = await _run_draft_letter(output_format=None)
    mock_pdf.assert_called_once()
    mock_docx.assert_not_called()
    assert result["output_format"] == "pdf"
    assert mock_send.call_args.kwargs["content_type"] == "application/pdf"
    assert mock_send.call_args.args[2].endswith(".pdf")


@pytest.mark.asyncio
async def test_explicit_docx_request_uses_docx_renderer():
    result, mock_pdf, mock_docx, mock_send = await _run_draft_letter(output_format="docx")
    mock_docx.assert_called_once()
    mock_pdf.assert_not_called()
    assert result["output_format"] == "docx"
    assert mock_send.call_args.kwargs["content_type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert mock_send.call_args.args[2].endswith(".docx")


@pytest.mark.asyncio
async def test_unrecognised_output_format_value_defaults_to_pdf():
    """Anything other than the literal string 'docx' (a typo, an unexpected value) fails safe
    to the existing, proven PDF path rather than erroring."""
    result, mock_pdf, mock_docx, mock_send = await _run_draft_letter(output_format="word")
    mock_pdf.assert_called_once()
    mock_docx.assert_not_called()
    assert result["output_format"] == "pdf"


# ── _send_invoice_document: content_type plumbing ──────────────────────────────────

@pytest.mark.asyncio
async def test_send_invoice_document_defaults_to_pdf_content_type():
    from vula.api.whatsapp import _send_invoice_document

    mock_response = MagicMock()
    mock_response.raise_for_status = MagicMock()
    mock_response.json.return_value = {"id": "media123"}
    mock_client = AsyncMock()
    mock_client.__aenter__ = AsyncMock(return_value=mock_client)
    mock_client.__aexit__ = AsyncMock(return_value=None)
    mock_client.post = AsyncMock(return_value=mock_response)

    with (
        patch("vula.api.whatsapp._get_tenant_wa_creds",
              new=AsyncMock(return_value={"token": "tok", "phone_id": "123"})),
        patch("vula.api.whatsapp.httpx.AsyncClient", return_value=mock_client),
        patch("vula.api.whatsapp._record_outbound"),
    ):
        ok = await _send_invoice_document("27821234567", b"bytes", "file.pdf", tenant_id="digg-demo")

    assert ok is True
    upload_call = mock_client.post.call_args_list[0]
    assert upload_call.kwargs["data"]["type"] == "application/pdf"
    assert upload_call.kwargs["files"]["file"][2] == "application/pdf"


@pytest.mark.asyncio
async def test_send_invoice_document_uses_explicit_docx_content_type():
    from vula.api.whatsapp import _send_invoice_document

    mock_response = MagicMock()
    mock_response.raise_for_status = MagicMock()
    mock_response.json.return_value = {"id": "media123"}
    mock_client = AsyncMock()
    mock_client.__aenter__ = AsyncMock(return_value=mock_client)
    mock_client.__aexit__ = AsyncMock(return_value=None)
    mock_client.post = AsyncMock(return_value=mock_response)
    docx_ct = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    with (
        patch("vula.api.whatsapp._get_tenant_wa_creds",
              new=AsyncMock(return_value={"token": "tok", "phone_id": "123"})),
        patch("vula.api.whatsapp.httpx.AsyncClient", return_value=mock_client),
        patch("vula.api.whatsapp._record_outbound"),
    ):
        ok = await _send_invoice_document("27821234567", b"bytes", "file.docx",
                                          tenant_id="digg-demo", content_type=docx_ct)

    assert ok is True
    upload_call = mock_client.post.call_args_list[0]
    assert upload_call.kwargs["data"]["type"] == docx_ct
    assert upload_call.kwargs["files"]["file"][2] == docx_ct

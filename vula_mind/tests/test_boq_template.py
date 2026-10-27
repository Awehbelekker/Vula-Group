"""The BOQ template as a Vula feature, 27 Oct.

Ian asked for a branded Bill of Quantities template "based on the BOQ given by the QS at HPC…
that uses updated costs", then for it to follow the ASAQS standard, then for Vula to send it.
The workbook is built from the tenant's own brand settings and price book. All figures here are
made up.
"""
import asyncio
import io
from datetime import date
from unittest.mock import patch

from openpyxl import load_workbook

from vula.commerce import boq_template as bt
from vula.ingestion import boq_sheet

TID = "test-tenant"
TODAY = date(2026, 10, 27)


def _obs(desc, cents, on="2026-10-01", kind="material", source="invoice", section=None, supplier="Supplier A",
         unit="each"):
    return {"description": desc, "norm_key": desc.lower(), "unit": unit, "kind": kind, "section": section,
            "supplier": supplier, "project": None, "source_kind": source, "quantity": 1,
            "unit_price_cents": cents, "observed_on": on, "doc_id": "d1"}


def test_items_are_filed_under_their_asaqs_trade():
    assert bt.trade_of("Shower door 900mm") == "Glazing"
    assert bt.trade_of("Lever door handle, satin chrome") == "Ironmongery"
    assert bt.trade_of("Interior door 813 x 2032") == "Carpentry and Joinery"
    assert bt.trade_of("Cement plaster to walls") == "Plastering"
    assert bt.trade_of("CEMENT 50KG 42.5N") == "Concrete, Formwork and Reinforcement"
    assert bt.trade_of("SOLID 12mm FIRE BOARD 3.0 X 1.2") == "Ceilings, Partitions and Access Flooring"
    assert bt.trade_of("Grout 5kg grey") == "Tiling"
    assert bt.trade_of("Basin mixer, chrome") == "Plumbing and Drainage"
    assert bt.trade_of("Scaffold standard 3m") == "Preliminaries"


def test_the_bill_section_names_the_trade_before_the_description():
    assert bt.trade_of("Supply and fix 600 x 600", "CEILINGS, PARTITIONS AND ACCESS FLOORING") == \
        "Ceilings, Partitions and Access Flooring"
    # A room heading names no trade: the description decides.
    assert bt.trade_of("Floor tiles 600 x 600", "Guest Bathroom") == "Tiling"


def test_lines_that_are_not_building_work_have_no_trade():
    for line in ("PORSCHE X SMEG 2-SLICE TOASTER", "SEWERAGE (Period 11/08/2026 to 08/09/2026) (Actual reading)",
                 "October Rent", "Time and distance for 13 days rental of car group A",
                 "FNB App Payment To Hpc Skipp Rubble", "OBI CHAIR, 5-Star black nylon", "Gross amount certified"):
        assert bt.trade_of(line) is None, line


def test_the_rate_library_comes_from_the_tenants_own_price_book():
    obs = [
        _obs("Floor tiles 600x600", 18000, on="2026-09-01"),                    # paid
        _obs("Floor tiles 600x600", 22000, on="2026-10-01"),                    # paid
        _obs("Floor tiles 600x600", 30000, on="2026-10-02", source="quote"),    # quoted — not the rate
        _obs("Smeg toaster", 250000),                                          # not building work
        _obs("Grout 5kg", 9000, on="2024-01-10"),                               # older than 18 months
        _obs("Delivery to site", 50000, kind="delivery"),                       # not work
        _obs("Site establishment", 1500000, source="boq", section="PRELIMINARIES"),  # a job's lump sum
        _obs("Provisional sum for electrical", 2000000, source="boq"),          # not a rate
        _obs("Half brick walls", 18500, source="boq", section="MASONRY", supplier=None, unit="m2"),
        _obs("Tile cutter 1.5m", 35000, kind="plant", unit="day"),
    ]
    with patch("vula.commerce.price_book._fetch", return_value=obs) as fetch:
        rows = bt.learned_rates(TID, TODAY)
    fetch.assert_called_once_with(TID)
    got = {(r[0], r[1]): r for r in rows}
    assert set(got) == {("Tiling", "Floor tiles 600x600"), ("Masonry", "Half brick walls"),
                        ("Preliminaries", "Tile cutter 1.5m")}
    tiles = got[("Tiling", "Floor tiles 600x600")]
    assert (tiles[2], tiles[3], tiles[4], tiles[6]) == ("No", 200.0, "Material", "1 Oct 2026")   # median paid
    bricks = got[("Masonry", "Half brick walls")]
    assert (bricks[2], bricks[3], bricks[4], bricks[5]) == ("m²", 185.0, "Bill rate", "Priced bill")
    assert got[("Preliminaries", "Tile cutter 1.5m")][4] == "Plant hire"


def test_own_rates_come_first_and_every_dropdown_key_is_unique():
    own = [["Tiling", "Floor tiles 600x600", "No", 210.0, "Your rate", "Your rate list", "1 Oct 2026"]]
    learned = [["Tiling", "Floor tiles 600x600", "No", 200.0, "Material", "Your rate list", "1 Oct 2026"],
               ["Masonry", "Half brick walls", "m²", 185.0, "Bill rate", "Priced bill", "26 Oct 2026"]]
    with patch.object(bt, "_own_rates", return_value=own), patch.object(bt, "learned_rates", return_value=learned):
        lib = bt.library(TID, TODAY)
    assert [r[1] for r in lib] == ["Masonry", "Tiling", "Tiling"]            # trade order
    assert lib[1][5] == "Your rate"                                          # own rate first in its trade
    assert len({r[0] for r in lib}) == 3 and lib[2][0].endswith("(2)")


BRAND = {"display_name": "Example Builders", "legal_name": "Example Holdings (Pty) Ltd t/a Example Builders",
         "reg_no": "2020/000000/07", "address": "1 Example Street, Cape Town", "vat_registered": False,
         "vat_number": "", "accent": "2C5545", "ink": "1E1E1E", "logo_url": None}
LIB = [["Masonry: Half brick walls [m²] — Priced bill", "Masonry", "Half brick walls", "m²", 185.0, "Bill rate",
        "Priced bill", "26 Oct 2026"],
       ["Tiling: Floor tiles 600x600 [No] — Supplier A", "Tiling", "Floor tiles 600x600", "No", 200.0, "Material",
        "Supplier A", "1 Oct 2026"]]


def _workbook(data):
    return load_workbook(io.BytesIO(data))


def test_the_workbook_carries_the_tenants_brand_and_the_asaqs_layout():
    wb = _workbook(bt.render(BRAND, LIB, None, TODAY))
    assert wb.sheetnames == ["Cover", "Read me", "BOQ", "Priced Bill", "Summary", "Rate Library", "Lists"]
    cover = wb["Cover"]
    assert cover["E2"].value == "Example Builders"
    assert "Example Holdings (Pty) Ltd t/a Example Builders" in cover.oddFooter.left.text
    assert cover["C37"].value == bt.BASES[0]                   # basis of measurement, contractor's estimate
    lists = wb["Lists"]
    assert [lists.cell(i, 1).value for i in range(2, 2 + len(bt.TRADES))] == bt.TRADES
    assert wb["Rate Library"]["C3"].value == "Floor tiles 600x600"
    assert wb["BOQ"]["C9"].value == LIB[1][0]                   # the example line uses a tiling item
    assert wb.calculation.fullCalcOnLoad


def test_the_bill_prints_only_as_far_as_its_last_line():
    import zipfile
    with zipfile.ZipFile(io.BytesIO(bt.render(BRAND, LIB, None, TODAY))) as z:
        names = z.read("xl/workbook.xml").decode()
    assert names.count('name="_xlnm.Print_Area"') == 2 and "OFFSET('Priced Bill'!$A$1" in names


def test_vula_reads_the_templates_boq_sheet_back():
    """A filled-in copy sent to Vula sets the project budget: the BOQ sheet's headings are the
    ones the sheet reader knows, and its other sheets are not read as bills."""
    wb = _workbook(bt.render(BRAND, LIB, None, TODAY))
    rows = [list(r) for r in wb["BOQ"].iter_rows(min_row=8, max_row=8, values_only=True)]
    cols = boq_sheet._header_map(rows[0])
    assert (cols["desc"], cols["qty"], cols["rate"], cols["amount"]) == (3, 5, 8, 9)
    for name in ("Rate Library", "Summary", "Cover"):
        sheet = [list(r) for r in wb[name].iter_rows(values_only=True)]
        assert boq_sheet.parse_rows(sheet, name, infer=False) == [], name


def test_whatsapp_tool_sends_a_private_link():
    from core.skills.commerce_admin import CommerceAdminSkill
    with patch.object(bt, "link", return_value=("https://example.test/signed", {"rates": 12, "trades": 3,
                                                                                "own_rates": 0, "logo": True})):
        out = asyncio.run(CommerceAdminSkill()._boq_template(TID))
    assert "https://example.test/signed" in out["reply_verbatim"] and "12 of your own prices" in out["reply_verbatim"]
    assert "](" not in out["reply_verbatim"]                     # WhatsApp: raw URL, no markdown link


def test_only_project_businesses_get_the_tool():
    from core.skills.commerce_admin import _tools_for
    for projects, expected in ((True, True), (False, False)):
        with (patch("vula.api.tenants.enabled_modules", return_value=[]),
              patch("vula.api.tenants.uses_projects", return_value=projects)):
            names = {t["function"]["name"] for t in _tools_for(TID, role=None)}
        assert ("boq_template" in names) is expected

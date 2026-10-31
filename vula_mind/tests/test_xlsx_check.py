"""The read-back's formula evaluator understands exactly what Vula writes, and nothing else (31 Oct 2026)."""
import pytest
from openpyxl import Workbook

from vula.commerce import xlsx_check as x


@pytest.fixture
def ws():
    ws = Workbook().active
    ws["A1"], ws["A2"], ws["A3"] = 10, 2.5, "text"
    return ws


@pytest.mark.parametrize("formula,expected", [
    ("=A1-A2", 7.5), ("=(A1-A2)/A1", 0.75), ("=-A2*2", -5.0), ("=SUM(A1:A3)", 12.5),
    ("=AVERAGE(A1:A2,A1:A1)", 7.5), ("=COUNT(A1:A3)", 2.0), ("=MAX(A1:A2)-MIN(A1:A2)", 7.5),
    ("=ROUND(2.675,2)", 2.68), ("=ROUND(-2.675,2)", -2.68), ("=ROUND(308.625,2)", 308.63),
    ('=IFERROR(ROUND(A1/0,2),"")', ""), ('=IFERROR(ROUND(A1/A2,4),"")', 4.0),
    ('=IFERROR(A3*2,"")', ""), ("=A9+1", 1.0),
])
def test_what_vula_writes_is_evaluated(ws, formula, expected):
    ws["B1"] = formula
    assert x.value(ws, "B1") == expected


def test_a_formula_that_uses_another_formula_cell(ws):
    ws["B1"], ws["B2"], ws["B3"] = "=A1*2", "=B1+A2", "=B2-B1-A2"
    assert x.value(ws, "B3") == 0.0


@pytest.mark.parametrize("formula", ["=VLOOKUP(A1,A1:A2,1)", "=A1&A2", "=SUM(A1:A2", "=A1 A2"])
def test_anything_else_is_never_passed_as_correct(ws, formula):
    ws["B1"] = formula
    with pytest.raises(ValueError):
        x.value(ws, "B1")


def test_an_error_value_outside_iferror_fails_the_check(ws):
    ws["B1"] = "=A1/0"
    with pytest.raises(ValueError):
        x.value(ws, "B1")

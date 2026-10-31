"""Forms the team fills in (31 Oct 2026, item 2): made from the books, read back when returned,
booked only on the owner's Confirm, once. Made-up businesses and figures."""
import io
import shutil
import subprocess
import uuid
from datetime import date

import pytest
from openpyxl import load_workbook

from vula.commerce import forms

TODAY = date(2026, 10, 31)


class _Q:
    def __init__(self, db, table):
        self.db, self.table, self.filters, self._op, self._payload, self._limit = db, table, [], "select", None, None

    def select(self, *_a, **_k):
        return self

    def eq(self, k, v):
        self.filters.append(lambda r: r.get(k) == v)
        return self

    def in_(self, k, vs):
        self.filters.append(lambda r: r.get(k) in vs)
        return self

    def gte(self, k, v):
        self.filters.append(lambda r: str(r.get(k) or "") >= v)
        return self

    def lte(self, k, v):
        self.filters.append(lambda r: str(r.get(k) or "") <= v)
        return self

    def order(self, k, desc=False):
        self._order = (k, desc)
        return self

    def limit(self, n):
        self._limit = n
        return self

    def insert(self, row):
        self._op, self._payload = "insert", row
        return self

    def update(self, patch):
        self._op, self._payload = "update", patch
        return self

    def execute(self):
        rows = self.db.setdefault(self.table, [])
        if self._op == "insert":
            new = [{"id": str(uuid.uuid4()), **r} for r in (self._payload if isinstance(self._payload, list) else [self._payload])]
            for r in new:
                r.setdefault("id", str(uuid.uuid4()))
            rows.extend(new)
            return type("R", (), {"data": new})
        hit = [r for r in rows if all(f(r) for f in self.filters)]
        if self._op == "update":
            for r in hit:
                r.update(self._payload)
            return type("R", (), {"data": hit})
        if getattr(self, "_order", None):
            k, desc = self._order
            hit = sorted(hit, key=lambda r: str(r.get(k) or ""), reverse=desc)
        if self._limit:
            hit = hit[: self._limit]
        return type("R", (), {"data": [dict(r) for r in hit]})


class _DB:
    def __init__(self):
        self.tables = {}

    def table(self, name):
        return _Q(self.tables, name)


PRODUCTS = [{"id": "p1", "name": "Round basin", "sku": "RB-1", "price_cents": 145_000, "stock_quantity": 12},
            {"id": "p2", "name": "Oval basin", "sku": "OB-1", "price_cents": 165_000, "stock_quantity": 4},
            {"id": "p3", "name": "Sealer 1L", "sku": "", "price_cents": 12_345, "stock_quantity": None}]


@pytest.fixture
def biz(monkeypatch):
    db = _DB()
    db.tables["commerce_workers"] = [
        {"id": "w1", "tenant_id": "tenant-a", "name": "Sipho", "rate_cents": 45_000, "rate_period": "daily", "active": True},
        {"id": "w2", "tenant_id": "tenant-a", "name": "Lindiwe", "rate_cents": 250_000, "rate_period": "weekly", "active": True},
        {"id": "w9", "tenant_id": "tenant-b", "name": "Other Co worker", "rate_cents": 1, "rate_period": "daily", "active": True}]
    from vula.commerce import boq_template, service
    monkeypatch.setattr(service, "_client", lambda: db)

    async def products(tid, in_stock_only=False, include_archived=False):
        return [dict(p) for p in PRODUCTS] if tid == "tenant-a" else []
    monkeypatch.setattr(service, "list_products", products)
    monkeypatch.setattr(service, "registered_projects", lambda tid: [{"name": "Seaview"}, {"name": "Oak Street"}])
    monkeypatch.setattr(service, "canonical_project", lambda tid, n: "Seaview" if "sea" in (n or "").lower() else n)
    monkeypatch.setattr(boq_template, "brand", lambda tid: {"display_name": "Made Up Builders", "legal_name": "Made Up Builders",
                                                             "accent": "2C5545", "vat_registered": True, "logo_url": None})
    monkeypatch.setattr(boq_template, "_logo_bytes", lambda url: None)
    return db


def _fill(data: bytes, cells: dict, tmp_path, name="filled.xlsx"):
    wb = load_workbook(io.BytesIO(data))
    ws = wb.worksheets[0]
    for ref, v in cells.items():
        ws[ref] = v
    path = tmp_path / name
    wb.save(path)
    return path


# ── making ───────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_a_quote_form_has_the_dropdowns_locked_sums_and_its_hidden_id(biz):
    out = await forms.make("tenant-a", "quote", by="27820000001", today=TODAY)
    wb = load_workbook(io.BytesIO(out["bytes"]))
    ws = wb.worksheets[0]
    assert wb["_vula"]["A2"].value == out["form_id"] and wb["_vula"].sheet_state == "hidden"
    assert wb["Lists"].sheet_state == "hidden" and wb["Lists"]["A2"].value == "Oval basin"     # A–Z
    assert ws.protection.sheet is True
    assert ws["A9"].protection.locked is False and ws["B9"].protection.locked is False
    assert ws["C9"].protection.locked is True and ws["C9"].value.startswith("=IF(A9")
    assert any("A9" in str(dv.sqref) for dv in ws.data_validations.dataValidation)
    row = biz.tables["vula_forms"][0]
    assert row["tenant_id"] == "tenant-a" and row["status"] == "sent" and row["spec"]["layout"]["priced"]


@pytest.mark.asyncio
async def test_the_logo_goes_on_the_form(biz, monkeypatch):
    from PIL import Image

    from vula.commerce import boq_template
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "green").save(buf, "PNG")
    monkeypatch.setattr(boq_template, "_logo_bytes", lambda url: buf.getvalue())
    out = await forms.make("tenant-a", "timesheet", today=TODAY)
    assert len(load_workbook(io.BytesIO(out["bytes"])).worksheets[0]._images) == 1


@pytest.mark.asyncio
async def test_what_can_and_cant_be_made(biz):
    with pytest.raises(forms.FormError, match="Which project"):
        await forms.make("tenant-a", "site_checklist")
    with pytest.raises(forms.FormError, match="can't find a project"):
        await forms.make("tenant-a", "site_checklist", project="Mars Base")
    with pytest.raises(forms.FormError, match="stock count"):
        await forms.make("tenant-b", "stock_count")
    with pytest.raises(forms.FormError):
        await forms.make("tenant-a", "payroll")


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
@pytest.mark.asyncio
async def test_the_forms_own_sums_agree_with_what_code_reads_back(biz, tmp_path):
    out = await forms.make("tenant-a", "quote", today=TODAY)
    path = _fill(out["bytes"], {"B5": "Sam", "A9": "Round basin", "B9": 2, "A10": "Sealer 1L", "B10": 2.5}, tmp_path)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp_path / "o"), str(path)],
                   check=True, capture_output=True, timeout=120)
    ws = load_workbook(tmp_path / "o" / "filled.xlsx", data_only=True).worksheets[0]
    got = await forms.stage("tenant-a", path, "27820000001")
    lines = got["staged"]["lines"]
    assert [round(ws["D9"].value * 100), round(ws["D10"].value * 100)] == [ln["total_cents"] for ln in lines]
    assert lines[1]["total_cents"] == 30_863                                   # 2.5 × R123.45, half up


# ── reading back ─────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_a_returned_quote_is_checked_row_by_row(biz, tmp_path):
    out = await forms.make("tenant-a", "quote", today=TODAY)
    path = _fill(out["bytes"], {"B5": "Sam", "A9": "round basin", "B9": 3, "A10": "Gold tap", "B10": 1,
                                "A11": "Oval basin", "B11": "lots"}, tmp_path)
    got = await forms.stage("tenant-a", path, "27820000002")
    assert got["bookable"]
    lines, problems = got["staged"]["lines"], got["staged"]["problems"]
    assert [(ln["description"], ln["quantity"], ln["unit_price_cents"]) for ln in lines] == [("Round basin", 3.0, 145_000)]
    assert any("Gold tap" in p for p in problems) and any("Oval basin" in p for p in problems)
    assert "Draft quote for *Sam*" in got["text"] and "Left out" in got["text"]
    assert biz.tables["vula_forms"][0]["status"] == "returned"


@pytest.mark.asyncio
async def test_a_quote_without_a_customer_is_not_staged(biz, tmp_path):
    out = await forms.make("tenant-a", "quote", today=TODAY)
    got = await forms.stage("tenant-a", _fill(out["bytes"], {"A9": "Round basin", "B9": 1}, tmp_path), "27820000002")
    assert not got["bookable"] and "customer's name is missing" in got["text"]


@pytest.mark.asyncio
async def test_another_business_or_a_plain_spreadsheet_is_not_a_form(biz, tmp_path):
    from openpyxl import Workbook
    out = await forms.make("tenant-a", "quote", today=TODAY)
    path = _fill(out["bytes"], {"B5": "Sam", "A9": "Round basin", "B9": 1}, tmp_path)
    assert await forms.stage("tenant-b", path, "27820000002") is None
    plain = tmp_path / "plain.xlsx"
    Workbook().save(plain)
    assert await forms.stage("tenant-a", plain, "27820000002") is None


# ── booking ──────────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_a_quote_is_booked_once_as_a_draft(biz, tmp_path, monkeypatch):
    from vula.commerce import service
    made = []

    async def create_invoice(tid, data):
        row = {"id": str(uuid.uuid4()), "tenant_id": tid, "invoice_number": "MUB-QUO-00007", "status": "draft",
               "total_cents": 333_500, **{k: v for k, v in data.items() if k != "line_items"}}
        biz.tables.setdefault("commerce_invoices", []).append(row)
        made.append(data)
        return row
    monkeypatch.setattr(service, "create_invoice", create_invoice)
    out = await forms.make("tenant-a", "quote", today=TODAY)
    path = _fill(out["bytes"], {"B5": "Sam", "A9": "Round basin", "B9": 2}, tmp_path)
    await forms.stage("tenant-a", path, "27820000001")
    text = await forms.book("tenant-a", out["form_id"], by="27820000001")
    assert "Draft quote *MUB-QUO-00007*" in text and "Nothing has gone to the customer" in text
    assert made[0]["doc_type"] == "quote" and made[0]["line_items"] == [
        {"description": "Round basin", "quantity": 2.0, "unit_price_cents": 145_000}]
    with pytest.raises(forms.FormError, match="already booked"):
        await forms.book("tenant-a", out["form_id"])
    again = await forms.stage("tenant-a", path, "27820000001")
    assert not again["bookable"] and "already booked" in again["text"] and len(made) == 1


@pytest.mark.asyncio
async def test_a_failed_booking_books_nothing_and_can_be_tried_again(biz, tmp_path, monkeypatch):
    from vula.commerce import service

    async def broken(tid, data):
        raise RuntimeError("database down")
    monkeypatch.setattr(service, "create_invoice", broken)
    out = await forms.make("tenant-a", "quote", today=TODAY)
    await forms.stage("tenant-a", _fill(out["bytes"], {"B5": "Sam", "A9": "Round basin", "B9": 1}, tmp_path), "x")
    with pytest.raises(RuntimeError):
        await forms.book("tenant-a", out["form_id"])
    assert biz.tables["vula_forms"][0]["status"] == "returned"


@pytest.mark.asyncio
async def test_a_stock_count_shows_the_differences_then_sets_the_stock(biz, tmp_path, monkeypatch):
    from vula.commerce import stock_counts
    calls = {}

    async def create_count(tid, note=None, actor=None):
        calls["count"] = tid
        return {"id": "c1"}

    async def record_scans(tid, cid, scans, actor=None):
        calls["scans"] = scans
        return []

    async def apply(tid, cid, actor=None):
        return {"applied": len(calls["scans"]), "failed": []}
    monkeypatch.setattr(stock_counts, "create_count", create_count)
    monkeypatch.setattr(stock_counts, "record_scans", record_scans)
    monkeypatch.setattr(stock_counts, "apply", apply)
    out = await forms.make("tenant-a", "stock_count", today=TODAY)
    ws = load_workbook(io.BytesIO(out["bytes"])).worksheets[0]
    assert [ws["B9"].value, ws["B10"].value, ws["B11"].value] == ["Oval basin", "Round basin", None]  # sealer has no count
    got = await forms.stage("tenant-a", _fill(out["bytes"], {"D9": 4, "D10": 10}, tmp_path), "x")
    assert "Round basin: system 12.0, counted 10 (-2)" in got["text"]
    assert "1 different from the system" in got["text"] and "• Oval basin" not in got["text"]
    text = await forms.book("tenant-a", out["form_id"])
    assert "2 product(s) set to the count" in text
    assert {(s["product_id"], s["set"]) for s in calls["scans"]} == {("p2", 4.0), ("p1", 10.0)}


@pytest.mark.asyncio
async def test_a_timesheet_records_time_and_pay_due_not_a_cost(biz, tmp_path):
    out = await forms.make("tenant-a", "timesheet", today=TODAY)
    lists = load_workbook(io.BytesIO(out["bytes"]))["Lists"]
    assert [lists["C2"].value, lists["C3"].value] == ["Lindiwe", "Sipho"] and lists["C4"].value is None   # own workers only
    path = _fill(out["bytes"], {"A9": "2026-10-27", "B9": "Sipho", "C9": "seaview", "D9": 2,
                                "A10": "2026-10-28", "B10": "Sipo", "D10": 1,
                                "A11": "yesterday", "B11": "Sipho", "D11": 1,
                                "A12": date(2026, 10, 29), "B12": "Lindiwe", "C12": "Oak Street", "D12": 1}, tmp_path)
    got = await forms.stage("tenant-a", path, "x")
    assert [ln["total_cents"] for ln in got["staged"]["lines"]] == [90_000, 250_000]
    assert len(got["staged"]["problems"]) == 2 and "not booked as a cost" in got["text"]
    text = await forms.book("tenant-a", out["form_id"])
    assert "2 row(s), pay due R3,400.00" in text
    entries = biz.tables["vula_timesheet_entries"]
    assert {e["worker_name"] for e in entries} == {"Sipho", "Lindiwe"} and all(e["tenant_id"] == "tenant-a" for e in entries)
    due = forms.pay_due("tenant-a", "2026-10-26", "2026-11-01")
    assert due["total_cents"] == 340_000 and "Sipho: 2 day(s), R900.00 (Seaview)" in due["text"]
    assert forms.pay_due("tenant-b")["total_cents"] == 0


@pytest.mark.asyncio
async def test_a_site_checklist_goes_on_the_projects_board(biz, tmp_path):
    out = await forms.make("tenant-a", "site_checklist", project="seaview", today=TODAY)
    ws = load_workbook(io.BytesIO(out["bytes"])).worksheets[0]
    assert ws["B5"].value == "Seaview" and ws["A9"].value == forms.DEFAULT_CHECKLIST[0]
    assert ws["A9"].protection.locked is True and ws[f"A{9 + len(forms.DEFAULT_CHECKLIST)}"].protection.locked is False
    path = _fill(out["bytes"], {"B9": "Done", "B10": "Not done", "C10": "tiles short", "B11": "N/A", "B12": "maybe",
                                f"A{9 + len(forms.DEFAULT_CHECKLIST)}": "Skip bin collected",
                                f"B{9 + len(forms.DEFAULT_CHECKLIST)}": "Done"}, tmp_path)
    got = await forms.stage("tenant-a", path, "x")
    assert "Not done: Materials delivered and checked (tiles short)" in got["text"]
    assert any("maybe" in p for p in got["staged"]["problems"])
    text = await forms.book("tenant-a", out["form_id"])
    card = biz.tables["vula_project_cards"][0]
    assert card["project"] == "Seaview" and card["card_type"] == "checklist"
    done = {i["text"].split(" —")[0]: i["done"] for i in card["content"]["items"]}
    assert done["Site safe and tidy"] and not done["Materials delivered and checked"] and done["Skip bin collected"]
    assert "Still to do: Materials delivered and checked" in text
    second = await forms.make("tenant-a", "site_checklist", project="Seaview", today=TODAY)   # starts from its own list
    assert load_workbook(io.BytesIO(second["bytes"])).worksheets[0]["A9"].value == forms.DEFAULT_CHECKLIST[0]


# ── the tools and the WhatsApp path ──────────────────────────────────────────

@pytest.mark.asyncio
async def test_the_tool_respects_access_and_sends_the_form(biz, monkeypatch):
    from core.skills import commerce_admin as ca
    sent = []

    async def send(to, data, name, caption, tid, content_type=None):
        sent.append((to, name))
        return True
    monkeypatch.setattr("vula.api.whatsapp._send_invoice_document", send)
    monkeypatch.setattr("vula.api.whatsapp._upload_to_storage", lambda b, p, d, c: f"u:{p}")
    monkeypatch.setattr("vula.storage_links.signed", lambda u, ttl: "signed:" + u)
    skill, ctx = ca.CommerceAdminSkill(), {"tenant_id": "tenant-a", "phone": "27820000003"}
    monkeypatch.setattr(ca, "_member_access", lambda tid, phone: ["orders"])
    out = await skill._fill_in_form("tenant-a", {"form": "quote"}, ctx)
    assert "access doesn't include the quote" in out["message"] and not sent
    monkeypatch.setattr(ca, "_member_access", lambda tid, phone: None)
    out = await skill._fill_in_form("tenant-a", {"form": "quote"}, ctx)
    assert "is ready" in out["reply_verbatim"] and sent[0][0] == "27820000003"
    assert "waiting to come back" in (await skill._fill_in_form("tenant-a", {"action": "list"}, ctx))["reply_verbatim"]
    assert "fill_in_form" in {t["function"]["name"] for t in ca._tools_for("t", role="owner")}
    assert "fill_in_form" not in {t["function"]["name"] for t in ca._tools_for("t", role="sales_rep")}


@pytest.mark.asyncio
@pytest.mark.parametrize("role,who_confirms", [("owner", "sender"), ("staff", "owner")])
async def test_a_returned_form_waits_for_an_owner_or_manager(biz, tmp_path, monkeypatch, role, who_confirms):
    from vula import team_index
    from vula.api import whatsapp as wa
    asked, told = [], []
    monkeypatch.setattr(team_index, "member_for_phone", lambda tid, ph: {"role": role, "whatsapp": ph})
    monkeypatch.setattr(wa, "_form_approvers", lambda tid: ["27829990000"])

    async def ask(to, tid, tool, args, preview, caller_role=None, summary=None, labels=()):
        asked.append((to, tool, args))

    async def reply(to, text, tid=None, **k):
        told.append((to, text))
    monkeypatch.setattr(wa, "_ask_admin_confirm", ask)
    monkeypatch.setattr(wa, "_send_reply", reply)
    out = await forms.make("tenant-a", "quote", today=TODAY)
    path = _fill(out["bytes"], {"B5": "Sam", "A9": "Round basin", "B9": 1}, tmp_path)
    assert await wa._returned_form_first("tenant-a", "27820000004", path)
    to, tool, args = asked[0]
    assert tool == "form_book" and args == {"form_id": out["form_id"]}
    if who_confirms == "sender":
        assert to == "27820000004" and not told
    else:
        assert to == "27829990000" and "sent it to the owner to approve" in told[0][1]


@pytest.mark.asyncio
async def test_a_form_from_someone_not_on_the_team_is_filed_as_usual(biz, tmp_path, monkeypatch):
    from vula import team_index
    from vula.api import whatsapp as wa
    monkeypatch.setattr(team_index, "member_for_phone", lambda tid, ph: None)
    monkeypatch.setattr(wa, "_is_tenant_owner", lambda tid, ph: False)
    out = await forms.make("tenant-a", "quote", today=TODAY)
    path = _fill(out["bytes"], {"B5": "Sam", "A9": "Round basin", "B9": 1}, tmp_path)
    assert await wa._returned_form_first("tenant-a", "27830000000", path) is False
    assert biz.tables["vula_forms"][0]["status"] == "sent"


@pytest.mark.asyncio
async def test_confirm_books_through_the_tool_and_the_reply_is_sent_as_written(biz, tmp_path, monkeypatch):
    from core.skills import commerce_admin as ca
    from vula.commerce import service

    async def create_invoice(tid, data):
        row = {"id": "i1", "tenant_id": tid, "invoice_number": "MUB-QUO-00008", "status": "draft", "total_cents": 166_750}
        biz.tables.setdefault("commerce_invoices", []).append(row)
        return row
    monkeypatch.setattr(service, "create_invoice", create_invoice)
    out = await forms.make("tenant-a", "quote", today=TODAY)
    await forms.stage("tenant-a", _fill(out["bytes"], {"B5": "Sam", "A9": "Round basin", "B9": 1}, tmp_path), "x")
    skill = ca.CommerceAdminSkill()
    res = await skill._dispatch_tool("form_book", {"form_id": out["form_id"], "confirm": True},
                                     {"tenant_id": "tenant-a", "phone": "27820000001"})
    assert ca.checked_reply(res).startswith("✅ Draft quote *MUB-QUO-00008*")
    res = await skill._dispatch_tool("form_book", {"form_id": out["form_id"], "confirm": True},
                                     {"tenant_id": "tenant-a", "phone": "27820000001"})
    assert "already booked" in ca.checked_reply(res)

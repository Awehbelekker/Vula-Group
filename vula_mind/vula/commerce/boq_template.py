"""
vula/commerce/boq_template.py — a branded Bill of Quantities workbook for the tenant, priced from
its own price book.

2026-10-27 (Ian: "a template that a client can use that is based on the BOQ given by the QS at
HPC… dropdown selection… uses updated costs", then "branded by tenant brand kit and logo… so
they can add the client/project details", then "align the template with the ASAQS standard").
Built by hand for DIGG first; this makes it a feature any project business can ask for.

The workbook:
- Cover: the tenant's logo, company, registration and VAT details; client, project and bill
  details to fill in; the basis of measurement (contractor's estimate / ASAQS Standard System /
  SANS 1200), with the matching statement printed under it; acceptance lines.
- BOQ: one line per row — pick a trade section and a library item; description, unit and rate
  fill in; type the quantity. An own rate replaces the library rate on that line; an item not in
  the library is typed in. An optional Area / Room per line.
- Priced Bill: the client copy — the used lines only, in trade order, ending on the bill total.
- Summary: totals by trade, then contingency, site management, builder's fee and VAT; totals by
  kind (bill rate / material / labour / plant) and by room.
- Rate Library: the tenant's own rates (vula_qs_rates) first, then what its documents taught the
  price book (vula/commerce/price_book.py): the median of what was PAID, else of what was quoted
  or priced in a bill, over the last 18 months. Each item is filed under its ASAQS trade (the
  bill's own section first, then the words in its description); anything that is not building
  work (a quote's toaster, a municipal rates line) has no trade and stays out.
- Read me.

Trades follow the order of the ASAQS model bills (Standard System of Measuring Building Work);
units are the standard ones (m, m², m³, No, kg, t, Item, Sum, Prov Sum, PC Sum), a supplier's
pack size (bag, tin, sheet) kept as a materials price. Formulas are locked, inputs open; A4 page
setup with the tenant's footer; the BOQ and Priced Bill print only as far as their last line.

Every figure in the library is a price the tenant's own documents hold — nothing is invented,
and no other tenant's prices are ever read (every query is scoped by tenant_id). The workbook's
BOQ sheet reads back through vula/ingestion/boq_sheet.py, so a filled-in copy sent to Vula sets
the project budget. Money in the price book is integer cents; the sheet shows Rands.
"""
from __future__ import annotations

import io
import logging
import re
from collections import Counter
from datetime import date, timedelta
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

log = logging.getLogger(__name__)

# Trade order of the ASAQS model bills (Standard System of Measuring Building Work).
TRADES = ["Preliminaries", "Alterations", "Earthworks", "Concrete, Formwork and Reinforcement",
          "Precast Concrete", "Masonry", "Waterproofing", "Roof Coverings, etc.", "Carpentry and Joinery",
          "Ceilings, Partitions and Access Flooring", "Floor Coverings, Wall Linings, etc.", "Ironmongery",
          "Structural Steelwork", "Metalwork", "Plastering", "Tiling", "Plumbing and Drainage",
          "Electrical Installation", "Glazing", "Paintwork", "Paperhanging", "External Works",
          "Provisional Sums and Prime Cost Sums"]

# Which trade an item belongs to, from its words. Checked in this order: the specific before the
# general ("shower door" is glazing before "shower" is plumbing; "door handle" is ironmongery
# before "door" is joinery; "cement plaster" is plastering before "cement" is concrete).
_TRADE_WORDS: List[Tuple[str, re.Pattern]] = [(t, re.compile(p, re.IGNORECASE)) for t, p in [
    ("Provisional Sums and Prime Cost Sums", r"\bprovisional\b|\bp\.?\s?c\.? sums?\b|prime cost"),
    ("Preliminaries", r"preliminar|\bp ?& ?g\b|site establishment|site management|scaffold|"
                      r"\bhire\b|\bskip\b|drop ?sheet|\bdiscs?\b|\bblades?\b|\bsds\b|drill bits?|tile cutter|"
                      r"chipper"),
    ("Ceilings, Partitions and Access Flooring", r"acoustic tile|ceiling tile"),
    ("Floor Coverings, Wall Linings, etc.", r"carpet"),
    ("Tiling", r"\btil(e|es|ing)\b|grout|mosaic|tile adhesive|\btal\b|tile trim|\bspacers?\b|m2 per box|"
               r"\d+ ?boxes\b"),
    ("Glazing", r"glass|glazing|mirror|shower (door|screen|enclosure)|window panes?"),
    ("Electrical Installation", r"electric|\bcable|socket|\bplugs?\b|\blights\b|lighting|light fitting|"
                                r"downlight|\bled (strip|light|panel|bulb)|\blamps?\b|\bswitch|distribution board|"
                                r"conduit|isolator|extractor|\bcoc\b|wiring|data point"),
    ("Paperhanging", r"wallpaper"),
    ("Ironmongery", r"\block|handle|hinge|latch|door closer|cylinder|door ?stop|robe hook|padlock|"
                    r"push ?bar|panic bar"),
    ("Ceilings, Partitions and Access Flooring", r"ceiling|drywall|gypsum|rhino ?board|partition|cornice|"
                                                 r"bulkhead|\bstuds?\b|\btrack\b|cross tee|main tee|"
                                                 r"corner ?bead|trap ?door|insulation|isotherm|acoustic|wall channel|"
                                                 r"fibatape|fibre tape|jointing|tek screw|brandering|"
                                                 r"\b(9|12|15) ?mm (fire |moisture |sound )?board|access floor"),
    ("Plumbing and Drainage", r"plumb|\bpipe|toilet|\bwc\b|basin|\btaps?\b|mixer|shower|geyser|drain|"
                              r"\btrap\b|\bwaste\b|cistern|\bbath\b|bathtub|\bsink|valve|sanitary|bidet|"
                              r"flush|towel rail|gully|manhole|sewer|\belbow|coupling"),
    ("Paintwork", r"paint|primer|undercoat|enamel|varnish|polyfilla|\broller|duram|plascon|dulux|"
                  r"prominent|base pastel|non-?drip|\bpva\b"),
    ("Plastering", r"plaster|\bskim|rhinolite|screed|\brender|cretestone|bonding liquid|permabond|key ?it"),
    ("Waterproofing", r"waterproof|damp ?proof|\bdpc\b|membrane|torch[- ]?on"),
    ("Masonry", r"\bbricks?\b|brickwork|\bblocks?\b|lintel|brickforce|mortar|wall ties?|masonry|\bsand\b"),
    ("Concrete, Formwork and Reinforcement", r"concrete|ready ?mix|rebar|reinforc|\bmesh\b|formwork|"
                                             r"shutter|\bcement\b|binding wire|wire binding"),
    ("Roof Coverings, etc.", r"\broof|\bridge|\bibr\b|chromadek|gutter|flashing|fascia|barge ?board"),
    ("Floor Coverings, Wall Linings, etc.", r"vinyl|laminat|\bfloor(s|ing)?\b|\bspc\b|wall panel|cladding|"
                                            r"nosing"),
    ("Structural Steelwork", r"\bi-?beam|\bh-?section|\brsj\b|structural steel"),
    ("Metalwork", r"balustrade|handrail|burglar|\bgates?\b|\bsteel\b|alumin|\bangle\b|tubing|"
                  r"\banchors?\b|security door"),
    ("Carpentry and Joinery", r"\bdoors?\b|skirting|cupboard|kitchen unit|joinery|timber|shel(f|ves|ving)|"
                              r"vanity|worktop|counter ?top|melamine|architrave|\bnails?\b|plywood|\bmdf\b|"
                              r"\bbic\b|built[- ]in|edg(e|ing)\b|edge ?band|\bpine\b|\bskirt\b|\bpvc\b.*/m\b"),
    ("Earthworks", r"excavat|trench|backfill|compact|earthwork"),
    ("External Works", r"paving|\bkerb|fenc(e|ing)|landscap|irrigation|boundary wall|driveway"),
    ("Alterations", r"demolish|demolition|remove existing|\bremoval|hack off|break out|form opening|"
                    r"rubble|strip out|reduce opening"),
]]

UNITS = ["m", "m²", "m³", "No", "kg", "t", "l", "Item", "Sum", "Prov Sum", "PC Sum", "Week", "Month",
         "Day", "Hour"]
_UNIT = {"each": "No", "ea": "No", "no": "No", "nr": "No", "item": "Item", "m2": "m²", "sqm": "m²",
         "m3": "m³", "lm": "m", "sum": "Sum"}
BASES = ["Contractor's estimate (not measured to a standard)",
         "ASAQS Standard System (7th ed.) with Model Preambles for Trades",
         "SANS 1200 (civil engineering)", "Other (see scope)"]
KINDS = ["Your rate", "Bill rate", "Material", "Labour", "Plant hire", "Other"]
_KIND_LABEL = {"material": "Material", "labour": "Labour", "plant": "Plant hire", "other": "Other"}

LOOKBACK_DAYS = 548          # 18 months of prices
PER_TRADE = 40               # the most-used items per trade
UNLOCKED = Protection(locked=False)
_LINES = 300                 # BOQ lines
_OWN_RATES = 40              # blank rows for the tenant's own library items
_AREAS = 15                  # room / area names


# Lines a price book holds that are not building work, whatever words they share with it: a
# municipal account ("SEWERAGE (Period 11/08/2026 …) (Actual reading)"), rent, car hire, a bank
# payment line, furniture and appliances from a fit-out quote, a certificate's own totals.
_NOT_WORK = re.compile(
    r"\bperiod\b.*\d{4}|actual reading|municipal|rates and taxes|\brefuse\b|\brent\b|monthly rental|"
    r"rental of car|car group|time and distance|valet|waiver|\bfuel\b|\bfnb\b|app payment|payment to\b|"
    r"\bchairs?\b|\bsofa|\bstools?\b|refrigerator|fridge|microwave|kettle|toaster|\bsmeg\b|coffee|"
    r"cappuccino|\blocker\b|dining|\btables?\b|amount certified|value of work|\bprofit\b|retention",
    re.IGNORECASE)


def trade_of(description: str, section: Optional[str] = None) -> Optional[str]:
    """The ASAQS trade for an item: its bill section when that names a trade ("CEILINGS,
    PARTITIONS AND ACCESS FLOORING", "Electrical Installation"), else its description. None when
    neither is building work — a room heading ("Guest Bathroom") falls through to the
    description, and a toaster, a rates account or office rent has no trade at all."""
    if _NOT_WORK.search(description or ""):
        return None
    for text in (section, description):
        for trade, pat in _TRADE_WORDS:
            if text and pat.search(text):
                return trade
    return None


def unit_label(unit: Optional[str]) -> str:
    u = (unit or "").strip()
    return _UNIT.get(u.lower(), u or "Item")


# ── data ───────────────────────────────────────────────────────────────────

def _client():
    from vula.commerce import service
    return service._client()


def brand(tenant_id: str) -> Dict[str, Any]:
    """The tenant's brand kit from its invoice settings (the same details its invoices carry)."""
    try:
        rows = (_client().table("commerce_invoice_settings").select("*")
                .eq("tenant_id", tenant_id).limit(1).execute().data or [])
    except Exception as exc:
        log.debug("BOQ template: invoice settings read failed for %s: %s", tenant_id, exc)
        rows = []
    st = rows[0] if rows else {}
    try:
        from vula.api.tenants import display_name
        fallback = display_name(tenant_id)
    except Exception:
        fallback = tenant_id
    company, trading = (st.get("company_name") or "").strip(), (st.get("trading_as") or "").strip()
    shown = trading or company or fallback
    legal = f"{company} t/a {trading}" if company and trading and company.lower() != trading.lower() else (company or shown)
    colour = lambda v, d: (re.sub(r"[^0-9a-fA-F]", "", v or "") or d)[:6].upper().rjust(6, "0")
    return {"display_name": shown, "legal_name": legal, "reg_no": (st.get("company_reg") or "").strip(),
            "address": (st.get("registered_address") or "").strip(),
            "vat_registered": bool(st.get("vat_registered")), "vat_number": (st.get("vat_number") or "").strip(),
            "accent": colour(st.get("accent_color"), "2C5545"), "ink": colour(st.get("ink_color"), "1E1E1E"),
            "logo_url": st.get("logo_url")}


def _logo_bytes(url: Optional[str]) -> Optional[bytes]:
    if not url:
        return None
    try:
        from vula.storage_links import fetch_sync
        return fetch_sync(url, timeout=20.0)
    except Exception as exc:
        log.info("BOQ template: logo not fetched (%s) — the workbook goes without it", exc)
        return None


def _own_rates(tenant_id: str) -> List[List[Any]]:
    """The tenant's own rate list (vula_qs_rates) — always first, as entered."""
    try:
        rows = (_client().table("vula_qs_rates").select("description,unit,rate,category,updated_at")
                .eq("tenant_id", tenant_id).execute().data or [])
    except Exception as exc:
        log.debug("BOQ template: own rates read failed for %s: %s", tenant_id, exc)
        return []
    out = []
    for r in rows:
        try:
            rate = float(r.get("rate") or 0)
        except (TypeError, ValueError):
            continue
        desc = re.sub(r"\s+", " ", r.get("description") or "").strip()
        trade = trade_of(desc, r.get("category"))
        if rate > 0 and desc and trade:
            out.append([trade, desc, unit_label(r.get("unit")), round(rate, 2), "Your rate", "Your rate list",
                        _day(r.get("updated_at"))])
    return out


def _day(v: Any) -> str:
    try:
        d = date.fromisoformat(str(v)[:10])
        return f"{d.day} {d:%b %Y}"
    except ValueError:
        return ""


def learned_rates(tenant_id: str, today: Optional[date] = None) -> List[List[Any]]:
    """[trade, description, unit, rate (R), kind, source, price date] from the price book's
    observations of the last 18 months: one rate per item and unit (price_book.rollup — the
    median of what was paid, else of what was quoted), filed under its trade, the most-used
    PER_TRADE per trade. A bill's preliminaries and provisional sums are lump sums for that job,
    not rates, and deliveries are not work — none of them is a library item."""
    from vula.commerce import price_book
    today = today or date.today()
    since, until = (today - timedelta(days=LOOKBACK_DAYS)).isoformat(), (today + timedelta(days=31)).isoformat()
    groups: Dict[tuple, List[Dict[str, Any]]] = {}
    for o in price_book._fetch(tenant_id):
        on = str(o.get("observed_on") or "")[:10]
        if not on or not since <= on <= until or o.get("kind") == "delivery" or not o.get("unit_price_cents"):
            continue
        groups.setdefault((o.get("norm_key"), o.get("unit")), []).append(o)
    by_trade: Dict[str, List[Tuple[int, str, List[Any]]]] = {}
    for obs in groups.values():
        r = price_book.rollup(obs)[0]
        section = Counter(o.get("section") for o in obs if o.get("section")).most_common(1)
        trade = trade_of(r["description"] or "", section[0][0] if section else None)
        if trade and r.get("kind") == "plant":
            trade = "Preliminaries"        # plant hire is a preliminary, whatever it is used for
        bill = all(o.get("source_kind") == "boq" for o in obs)
        if not trade or trade == "Provisional Sums and Prime Cost Sums" or (bill and trade == "Preliminaries"):
            continue
        desc = re.sub(r"\s+", " ", r["description"] or "").strip()[:200]
        kind = "Bill rate" if bill else _KIND_LABEL.get(r.get("kind"), "Material")
        source = (r.get("latest_supplier") or ("Priced bill" if bill else "Your documents")).strip()[:60]
        if r["basis"] == "quoted" and not bill:
            source += " (quoted)"
        row = [trade, desc, unit_label(r.get("unit")), round(r["rate_cents"] / 100, 2), kind, source,
               _day(r.get("latest_on"))]
        by_trade.setdefault(trade, []).append((r["observations"], str(r.get("latest_on") or ""), row))
    out = []
    for items in by_trade.values():       # the most-used, then the most recent
        out += [row for *_, row in sorted(items, key=lambda x: (x[0], x[1]), reverse=True)[:PER_TRADE]]
    return out


def library(tenant_id: str, today: Optional[date] = None) -> List[List[Any]]:
    """The Rate Library rows, keyed for the BOQ dropdown: own rates, then learned rates, in trade
    order. Each key is unique (a repeat gets a number)."""
    rows = _own_rates(tenant_id) + learned_rates(tenant_id, today)
    rows.sort(key=lambda r: (TRADES.index(r[0]), r[4] != "Your rate", r[1].lower()))
    seen: Counter = Counter()
    out = []
    for r in rows:
        key = f"{r[0]}: {r[1]} [{r[2]}] — {r[5]}"
        seen[key] += 1
        if seen[key] > 1:
            key += f" ({seen[key]})"
        out.append([key] + r)
    return out


# ── workbook ───────────────────────────────────────────────────────────────

def build(tenant_id: str, today: Optional[date] = None) -> Tuple[bytes, Dict[str, Any]]:
    """The workbook's bytes and what went into it ({"rates", "trades", "own_rates", "logo"})."""
    today = today or date.today()
    b = brand(tenant_id)
    lib = library(tenant_id, today)
    logo = _logo_bytes(b.get("logo_url"))
    data = render(b, lib, logo, today)
    return data, {"rates": len(lib), "trades": len({r[1] for r in lib}),
                  "own_rates": sum(1 for r in lib if r[5] == "Your rate"), "logo": bool(logo)}


def render(b: Dict[str, Any], lib: List[List[Any]], logo: Optional[bytes], today: date) -> bytes:
    A, INK = b["accent"], b["ink"]
    F = lambda size=10, bold=False, color=INK: Font(name="Arial", size=size, bold=bold, color=color)
    fill = lambda c: PatternFill("solid", fgColor=c)
    F_HEAD, F_INPUT, F_GREY = F(bold=True, color="FFFFFF"), F(color="0000FF"), F(9, color="6B6B6B")
    FILL_HEAD, FILL_IN, FILL_AUTO, FILL_BAND, FILL_SOFT = (fill(A), fill("FFF2CC"), fill("F2F2F2"),
                                                          fill("F7F9F8"), fill("E8EFEC"))
    thin = Side(style="thin", color="C9C9C9")
    BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    RAND = '"R" #,##0.00;[Red]-"R" #,##0.00;"-"'
    footer = " · ".join(x for x in (b["legal_name"], f"Reg No {b['reg_no']}" if b["reg_no"] else "",
                                    b["address"]) if x)
    vat_line = f"VAT No {b['vat_number']}" if b["vat_registered"] and b["vat_number"] else (
        "VAT registered" if b["vat_registered"] else "Not VAT registered")

    def head(ws, r, c, text):
        x = ws.cell(r, c, text)
        x.font, x.fill, x.border = F_HEAD, FILL_HEAD, BOX
        x.alignment = Alignment(vertical="center", wrap_text=True)

    def inp(ws, ref, value=None, fmt=None):
        x = ws[ref]
        x.value, x.fill, x.font, x.border = value, FILL_IN, F_INPUT, BOX
        x.protection = UNLOCKED
        if fmt:
            x.number_format = fmt

    def page(ws, landscape, title_rows=None):
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top, ws.page_margins.bottom, ws.page_margins.header = 0.75, 0.6, 0.25
        if title_rows:
            ws.print_title_rows = title_rows
        ws.oddHeader.right.text, ws.oddHeader.right.size = f"{b['display_name']} — Bill of Quantities", 8
        ws.oddFooter.left.text, ws.oddFooter.left.size = footer, 7
        ws.oddFooter.right.text, ws.oddFooter.right.size = "Page &P of &N", 8
        ws.sheet_view.showGridLines = False

    def put_logo(ws, anchor, height):
        if not logo:
            return
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(io.BytesIO(logo))
            img.width, img.height = int(img.width * height / img.height), height
            ws.add_image(img, anchor)
        except Exception as exc:
            log.info("BOQ template: logo not placed (%s)", exc)

    wb = Workbook()
    # ── Cover ──
    cv = wb.active
    cv.title = "Cover"
    for col, w in zip("ABCDEF", (3, 26, 34, 4, 22, 30)):
        cv.column_dimensions[col].width = w
    put_logo(cv, "B2", 110)
    cv["E2"] = b["display_name"]
    cv["E2"].font = F(13, True, A)
    for i, t in enumerate([x for x in (b["legal_name"] if b["legal_name"] != b["display_name"] else "",
                                       f"Reg No {b['reg_no']}" if b["reg_no"] else "", b["address"], vat_line) if x], 3):
        cv.cell(i, 5, t).font = F(9, color="555555")
    cv["B9"] = "BILL OF QUANTITIES"
    cv["B9"].font = F(22, True, A)
    cv["B10"] = '=IF(C21="","Project name",C21)'
    cv["B10"].font = F(14, True)
    cv.merge_cells("B10:F10")

    def block(r, title, rows):
        for c in range(2, 7):
            cv.cell(r, c).fill = FILL_HEAD
        cv.cell(r, 2, title).font = F_HEAD
        for j, (label, value, fmt) in enumerate(rows, r + 1):
            cv.cell(j, 2, label).font = F(10, True)
            inp(cv, f"C{j}", value, fmt)
            cv[f"C{j}"].alignment = Alignment(horizontal="left")
            cv.merge_cells(f"C{j}:F{j}")

    block(12, "CLIENT", [(x, None, None) for x in ("Client name", "Company", "Contact person", "Phone", "Email",
                                                    "Billing address", "VAT number (if any)")])
    block(20, "PROJECT", [("Project name", None, None), ("Project number", None, None), ("Site address", None, None),
                          ("Project type", None, None), ("Start date", None, "d mmm yyyy"),
                          ("Duration (weeks)", None, None), ("Architect / QS", None, None), ("Site manager", None, None)])
    block(30, "THIS BILL", [("BOQ reference", None, None), ("Revision", "Rev 0", None), ("Date issued", None, "d mmm yyyy"),
                            ("Valid for (days)", 30, "0"), ("Prepared by", b["display_name"], None)])
    cv["B36"] = "Valid until"
    cv["B36"].font = F(10, True)
    cv["C36"] = '=IF(OR(C33="",C34=""),"",C33+C34)'
    cv["C36"].number_format, cv["C36"].fill = "d mmm yyyy", FILL_AUTO
    cv["C36"].alignment = Alignment(horizontal="left")
    cv.merge_cells("C36:F36")
    cv["B37"] = "Basis of measurement"
    cv["B37"].font = F(10, True)
    inp(cv, "C37", BASES[0])
    cv["C37"].alignment = Alignment(horizontal="left")
    cv.merge_cells("C37:F37")
    cv["B38"] = "Scope, assumptions and exclusions"
    cv["B38"].font = F(10, True)
    inp(cv, "B39")
    cv.merge_cells("B39:F43")
    cv["B39"].alignment = Alignment(wrap_text=True, vertical="top")
    cv["B45"] = "Acceptance"
    cv["B45"].font = F(10, True, A)
    for j, label in enumerate(["Accepted by (client)", "Signature", "Date"], 46):
        cv.cell(j, 2, label).font = F(10)
        cv.cell(j, 3).border = Border(bottom=Side(style="thin", color="7F7F7F"))
        cv.merge_cells(f"C{j}:F{j}")
        cv.row_dimensions[j].height = 24
    cv["B50"] = "Prices in South African Rand. " + (
        "VAT at 15% shown separately." if b["vat_registered"] else "No VAT is charged (not VAT registered).")
    cv["B50"].font = F_GREY
    cv["B51"] = ('=IF(LEFT(C37,5)="ASAQS","Measured in accordance with the ASAQS Standard System of Measuring '
                 'Building Work (7th edition) and the Model Preambles for Trades.",IF(LEFT(C37,4)="SANS",'
                 '"Measured in accordance with SANS 1200 as set out in the project specification.",'
                 'IF(LEFT(C37,10)="Contractor","Quantities are the contractor\'s estimate and are not measured '
                 'to a published standard.","")))')
    cv["B51"].font = F_GREY
    cv.merge_cells("B51:F51")
    page(cv, False)

    # ── Lists (hidden) ──
    ls = wb.create_sheet("Lists")
    for col, title, values in ((1, "Sections", TRADES), (2, "Kinds", KINDS), (3, "Yes/No", ["Yes", "No"]),
                               (4, "Units", UNITS), (5, "Basis of measurement", BASES)):
        ls.cell(1, col, title)
        for i, v in enumerate(values, 2):
            ls.cell(i, col, v)
    ls.sheet_state = "hidden"
    NS = len(TRADES)
    sections = f"Lists!$A$2:$A${NS + 1}"
    dv_type = DataValidation(type="list", formula1='"New build,Renovation,Alteration,Fit-out,Extension,Maintenance,Other"',
                             allow_blank=True)
    dv_rev = DataValidation(type="list", formula1='"Rev 0,Rev 1,Rev 2,Rev 3,Rev 4,Final"', allow_blank=True)
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                             error="Enter a date, e.g. 27/10/2026", showErrorMessage=True)
    dv_basis = DataValidation(type="list", formula1=f"=Lists!$E$2:$E${len(BASES) + 1}", allow_blank=False)
    for dv, ref in ((dv_type, "C24"), (dv_rev, "C32"), (dv_date, "C25"), (dv_date, "C33"), (dv_basis, "C37")):
        if dv not in cv.data_validations.dataValidation:
            cv.add_data_validation(dv)
        dv.add(ref)

    # ── Rate Library ──
    rl = wb.create_sheet("Rate Library")
    for c, h in enumerate(["Key (shown in the BOQ dropdown)", "Bill section", "Library item", "Measured per",
                           "Price, R", "Kind", "Source", "Price date"], 1):
        head(rl, 1, c, h)
    for i, row in enumerate(lib, 2):
        for c, v in enumerate(row, 1):
            x = rl.cell(i, c, v)
            x.font, x.border = F(9), BOX
        rl.cell(i, 5).number_format = RAND
    own0 = len(lib) + 3
    last = own0 + _OWN_RATES - 1
    rl.cell(own0 - 1, 1, "YOUR OWN RATES — add a line in the yellow rows and it appears in every BOQ dropdown").font = F(10, True, A)
    for i in range(own0, last + 1):
        x = rl.cell(i, 1, f'=IF(C{i}="","",B{i}&": "&C{i}&" ["&D{i}&"] — own rate")')
        x.font, x.fill, x.border = F(9), FILL_AUTO, BOX
        for c in range(2, 9):
            inp(rl, rl.cell(i, c).coordinate)
        rl.cell(i, 5).number_format = RAND
    d_sec = DataValidation(type="list", formula1=f"={sections}", allow_blank=True)
    d_kind = DataValidation(type="list", formula1=f"=Lists!$B$2:$B${len(KINDS) + 1}", allow_blank=True)
    rl.add_data_validation(d_sec)
    rl.add_data_validation(d_kind)
    d_sec.add(f"B{own0}:B{last}")
    d_kind.add(f"F{own0}:F{last}")
    for col, w in zip("ABCDEFGH", (70, 28, 58, 11, 13, 14, 30, 12)):
        rl.column_dimensions[col].width = w
    rl.freeze_panes = "A2"
    rl.auto_filter.ref = f"A1:H{max(2, len(lib) + 1)}"
    page(rl, True, "1:1")
    wb.defined_names["RateKeys"] = DefinedName("RateKeys", attr_text=f"'Rate Library'!$A$2:$A${last}")

    # ── BOQ ──
    bq = wb.create_sheet("BOQ", 1)
    put_logo(bq, "A1", 62)
    bq["D1"] = "BILL OF QUANTITIES"
    bq["D1"].font = F(16, True, A)
    for j, (label, f) in enumerate([("Project", '=IF(Cover!C21="","",Cover!C21)'), ("Client", '=IF(Cover!C13="","",Cover!C13)'),
                                    ("Site", '=IF(Cover!C23="","",Cover!C23)'), ("BOQ ref / revision", '=TRIM(Cover!C31&"  "&Cover!C32)'),
                                    ("Date", '=IF(Cover!C33="","",Cover!C33)')], 2):
        bq.cell(j, 4, label).font = F(9, True)
        x = bq.cell(j, 5, f)
        x.font, x.fill = F(9), FILL_SOFT
        bq.merge_cells(start_row=j, start_column=5, end_row=j, end_column=8)
        if label == "Date":
            x.number_format, x.alignment = "d mmm yyyy", Alignment(horizontal="left")
    HR, FR = 8, 9
    LR = FR + _LINES - 1
    bq["K2"] = "Bill total"
    bq["K2"].font = F(9, True)
    bq["L2"] = f"=SUM(J{FR}:J{LR})"
    bq["L2"].font, bq["L2"].number_format = F(12, True, A), RAND
    for c, h in enumerate(["Section", "Item No", "Pick from rate library", "Description", "Unit", "Qty",
                           "Library rate (R)", "Your rate (R)", "Rate", "Amount", "Type", "Rate source",
                           "Area / Room", "Notes"], 1):
        head(bq, HR, c, h)

    def look(col, r):
        return f"INDEX('Rate Library'!${col}$2:${col}${last},MATCH($C{r},RateKeys,0))"
    for r in range(FR, LR + 1):
        formulas = {2: f'=IF(D{r}="","",COUNTIFS($A${FR}:A{r},A{r},$D${FR}:D{r},"?*"))',
                    4: f'=IF($C{r}="","",IFERROR({look("C", r)},""))', 5: f'=IF($C{r}="","",IFERROR({look("D", r)},""))',
                    7: f'=IF($C{r}="","",IFERROR({look("E", r)},""))', 9: f'=IF(H{r}<>"",H{r},IF(G{r}="","",G{r}))',
                    10: f'=IF(OR(F{r}="",I{r}=""),"",ROUND(F{r}*I{r},2))',
                    11: f'=IF($C{r}="",IF(H{r}<>"","Own rate",""),IFERROR({look("F", r)},""))',
                    12: f'=IF($C{r}="","",IFERROR({look("G", r)}&" · "&{look("H", r)},""))'}
        for c in range(1, 15):
            x = bq.cell(r, c)
            x.border, x.font = BOX, F(9)
            if c in formulas:
                x.value, x.fill = formulas[c], FILL_AUTO
            else:
                x.fill, x.font = FILL_IN, F(9, color="0000FF")
                x.protection = UNLOCKED
            if c in (4, 5):       # description and unit: typed when no library item is picked
                x.protection = UNLOCKED
        for c in (7, 8, 9, 10):
            bq.cell(r, c).number_format = RAND
        bq.cell(r, 6).number_format = "#,##0.00"
        bq.cell(r, 4).alignment = Alignment(wrap_text=True, vertical="top")
        bq.cell(r, 15, f'=IF(D{r}="","",IFERROR(MATCH(A{r},{sections},0),99)*1000+ROW())')   # print order
    bq.column_dimensions["O"].hidden = True
    bq.cell(LR + 1, 9, "Bill total").font = F(10, True)
    x = bq.cell(LR + 1, 10, f"=SUM(J{FR}:J{LR})")
    x.font, x.number_format = F(10, True), RAND
    example = next((r for r in lib if r[1] == "Tiling"), lib[0] if lib else None)
    if example:
        bq.cell(FR, 1, example[1])
        bq.cell(FR, 3, example[0])
        bq.cell(FR, 6, 1)
        bq.cell(FR, 14, "Example — overwrite or delete this line")
    dv_s = DataValidation(type="list", formula1=f"={sections}", allow_blank=True)
    dv_i = DataValidation(type="list", formula1="=RateKeys", allow_blank=True, showErrorMessage=False)
    dv_q = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                          error="Quantity must be a number of 0 or more", showErrorMessage=True)
    dv_u = DataValidation(type="list", formula1=f"=Lists!$D$2:$D${len(UNITS) + 1}", allow_blank=True,
                          errorStyle="warning", error="Not one of the standard units — keep it anyway?",
                          showErrorMessage=True)
    dv_a = DataValidation(type="list", formula1="=AreaNames", allow_blank=True, showErrorMessage=False)
    for dv, ref in ((dv_s, "A"), (dv_i, "C"), (dv_q, "F"), (dv_u, "E"), (dv_a, "M")):
        bq.add_data_validation(dv)
        dv.add(f"{ref}{FR}:{ref}{LR}")
    bq.conditional_formatting.add(f"A{FR}:N{LR}", FormulaRule(formula=[f'AND($A{FR}<>"",$A{FR}<>$A{FR - 1})'],
                                                              border=Border(top=Side(style="medium", color=A))))
    bq["C8"].comment = Comment("Click the arrow or start typing. 'Bill rate' items are all-in rates from priced "
                               "bills; supplier prices are materials only; labour rates are labour only.", b["display_name"])
    bq["H8"].comment = Comment("Optional. A rate here replaces the library rate on this line only.", b["display_name"])
    for col, w in zip("ABCDEFGHIJKLMN", (30, 6, 52, 44, 8, 9, 13, 12, 13, 15, 13, 32, 18, 26)):
        bq.column_dimensions[col].width = w
    bq.row_dimensions[HR].height = 30
    bq.freeze_panes = f"D{FR}"
    page(bq, True, f"{HR}:{HR}")
    bq_print = (f"OFFSET(BOQ!$A$1,0,0,MAX({FR},SUMPRODUCT(MAX((BOQ!$D${FR}:$D${LR}<>\"\")"
                f"*ROW(BOQ!$D${FR}:$D${LR})))),14)")

    def titles(ws, title):
        put_logo(ws, "A1", 62)
        ws["C1"] = title
        ws["C1"].font = F(16, True, A)
        for ref, f, font in (("C2", '=IF(Cover!C21="","",Cover!C21)', F(11, True)),
                             ("C3", '=IF(Cover!C13="","","Client: "&Cover!C13)', F(9)),
                             ("C4", '=TRIM(Cover!C31&"  "&Cover!C32)', F(9)),
                             ("C5", '="Basis of measurement: "&Cover!C37', F_GREY)):
            ws[ref] = f
            ws[ref].font = font

    # ── Priced Bill: used lines only, in trade order, ending on the bill total ──
    pb = wb.create_sheet("Priced Bill", 2)
    titles(pb, "PRICED BILL OF QUANTITIES")
    PH = 6
    PF, PL = PH + 1, PH + _LINES
    for c, h in enumerate(["Section", "Item", "Description", "Unit", "Qty", "Rate (R)", "Amount (R)"], 1):
        head(pb, PH, c, h)
    for i in range(_LINES):
        r = PF + i
        pb.cell(r, 9, f'=IFERROR(SMALL(BOQ!$O${FR}:$O${LR},{i + 1}),"")')       # hidden: order key
        pb.cell(r, 10, f'=IF(I{r}="","",MOD(I{r},1000))')                          # hidden: BOQ row

        def g(col):
            return f'IF($J{r}="","",INDEX(BOQ!${col}:${col},$J{r}))'
        pb.cell(r, 1, (f'=IF($J{r}="","",IF(INDEX(BOQ!$A:$A,$J{r})=IF(J{r - 1}="","~",INDEX(BOQ!$A:$A,J{r - 1})),'
                       f'"",INDEX(BOQ!$A:$A,$J{r})))') if i else f'=IF($J{r}="","",INDEX(BOQ!$A:$A,$J{r}))')
        pb.cell(r, 2, "=" + g("B"))
        pb.cell(r, 3, f'=IF($J{r}="","",INDEX(BOQ!$D:$D,$J{r})&IF(INDEX(BOQ!$M:$M,$J{r})="",""," — "&INDEX(BOQ!$M:$M,$J{r})))')
        pb.cell(r, 4, "=" + g("E"))
        pb.cell(r, 5, "=" + g("F"))
        end = f'AND($J{r}="",$J{r - 1}<>"")' if i else "FALSE"
        pb.cell(r, 6, f'=IF({end},"Bill total",' + g("I") + ")")
        pb.cell(r, 7, f'=IF({end},SUM(G${PF}:G{r - 1}),' + g("J") + ")" if i else "=" + g("J"))
        for c in range(1, 8):
            x = pb.cell(r, c)
            x.font = F(9, c == 1, A if c == 1 else INK)
            x.alignment = Alignment(wrap_text=c == 3, vertical="top", horizontal="center" if c in (2, 4) else None)
        pb.cell(r, 5).number_format = "#,##0.00"
        pb.cell(r, 6).number_format = pb.cell(r, 7).number_format = RAND
    pb.conditional_formatting.add(f"A{PF}:G{PL}", FormulaRule(formula=[f'$A{PF}<>""'], fill=FILL_SOFT,
                                                              border=Border(top=Side(style="thin", color=A))))
    pb.conditional_formatting.add(f"A{PF}:G{PL}", FormulaRule(formula=[f'AND($J{PF}<>"",MOD(ROW(),2)=0)'], fill=FILL_BAND))
    pb.conditional_formatting.add(f"F{PF}:G{PL}", FormulaRule(
        formula=[f'$F{PF}="Bill total"'], font=Font(name="Arial", bold=True, color=A),
        border=Border(top=Side(style="medium", color=A), bottom=Side(style="double", color=A))))
    pb.column_dimensions["I"].hidden = pb.column_dimensions["J"].hidden = True
    pb["F2"] = "Bill total"
    pb["F2"].font = F(9, True)
    pb["G2"] = "=BOQ!L2"
    pb["G2"].font, pb["G2"].number_format = F(12, True, A), RAND
    for col, w in zip("ABCDEFG", (30, 6, 62, 7, 9, 13, 15)):
        pb.column_dimensions[col].width = w
    page(pb, False, f"{PH}:{PH}")
    pb_print = f"OFFSET('Priced Bill'!$A$1,0,0,{PH}+COUNT('Priced Bill'!$J${PF}:$J${PL})+1,7)"

    # ── Summary ──
    sm = wb.create_sheet("Summary", 3)
    titles(sm, "SUMMARY")
    R0 = 7
    head(sm, R0, 1, "Bill section")
    head(sm, R0, 2, "Amount, R")

    def total_row(j, label, formula, bold=False):
        sm.cell(j, 1, label).font = F(10, bold)
        sm.cell(j, 1).border = BOX
        x = sm.cell(j, 2, formula)
        x.number_format, x.fill, x.border = RAND, FILL_AUTO, BOX
    for j, s in enumerate(TRADES, R0 + 1):
        total_row(j, s, f"=SUMIFS(BOQ!$J${FR}:$J${LR},BOQ!$A${FR}:$A${LR},A{j})")
    o = R0 + 1 + NS
    total_row(o, "Lines with no section", f'=SUMIFS(BOQ!$J${FR}:$J${LR},BOQ!$A${FR}:$A${LR},"")')
    s1 = o + 1
    sm.cell(s1, 1, "Subtotal").font = F(10, True)
    x = sm.cell(s1, 2, f"=SUM(B{R0 + 1}:B{o})")
    x.number_format, x.font, x.border = RAND, F(10, True), Border(top=Side(style="medium", color=A))
    inputs = [("Contingency (%)", 0.05, "0.0%", "Assumption: 5% for unforeseen work. Change to suit the job."),
              ("Site management (R per week)", 0, RAND, "e.g. R12,000 per week. Enter 0 if not charged."),
              ("Site management (weeks)", '=IF(Cover!C26="",0,Cover!C26)', "0",
               "Taken from the Cover (Duration); type over to change."),
              ("Builder's fee on cost (%)", 0, "0.0%", "Optional. Enter 0 if the fee is charged as site management."),
              ("VAT registered?", "Yes" if b["vat_registered"] else "No", "@",
               "VAT at the rate below is added when 'Yes'."),
              ("VAT rate (%)", 0.15, "0.0%", "South African VAT: 15%.")]
    at = {}
    for j, (label, val, fmt, note) in enumerate(inputs, s1 + 2):
        sm.cell(j, 1, label).font = F()
        inp(sm, f"B{j}", val, fmt)
        sm.cell(j, 3, note).font = F_GREY
        at[label] = j
    yes = DataValidation(type="list", formula1="=Lists!$C$2:$C$3", allow_blank=False)
    sm.add_data_validation(yes)
    yes.add(f"B{at['VAT registered?']}")
    k = at["VAT rate (%)"] + 2
    fee = at["Builder's fee on cost (%)"]
    calc = [("Contingency", f"=ROUND(B{s1}*B{at['Contingency (%)']},2)"),
            ("Site management", f"=B{at['Site management (R per week)']}*B{at['Site management (weeks)']}"),
            ("Builder's fee", f"=ROUND(B{s1}*B{fee},2)"),
            ("Total excl. VAT", f"=B{s1}+B{k}+B{k + 1}+B{k + 2}"),
            ("VAT", f'=IF(B{at["VAT registered?"]}="Yes",ROUND(B{k + 3}*B{at["VAT rate (%)"]},2),0)'),
            ("TOTAL", f"=B{k + 3}+B{k + 4}")]
    for j, (label, f) in enumerate(calc, k):
        big = label == "TOTAL"
        sm.cell(j, 1, label).font = F(12 if big else 10, label in ("Total excl. VAT", "TOTAL"), A if big else INK)
        x = sm.cell(j, 2, f)
        x.number_format, x.fill, x.font = RAND, FILL_SOFT if big else FILL_AUTO, F(12 if big else 10, big, A if big else INK)
        if big:
            x.border = Border(top=Side(style="medium", color=A), bottom=Side(style="double", color=A))
    t0 = k + len(calc) + 2
    head(sm, t0, 1, "By type")
    head(sm, t0, 2, "Amount, R")
    for j, t in enumerate(KINDS[1:5] + ["Own rate"], t0 + 1):
        total_row(j, t, f"=SUMIFS(BOQ!$J${FR}:$J${LR},BOQ!$K${FR}:$K${LR},A{j})")
    a0 = t0 + 8
    head(sm, a0, 1, "By area / room")
    head(sm, a0, 2, "Amount, R")
    sm.cell(a0, 3, "Type your rooms or areas in the yellow cells; they appear in the BOQ's Area / Room dropdown.").font = F_GREY
    for j in range(a0 + 1, a0 + 1 + _AREAS):
        inp(sm, f"A{j}")
        x = sm.cell(j, 2, f'=IF(A{j}="","",SUMIFS(BOQ!$J${FR}:$J${LR},BOQ!$M${FR}:$M${LR},A{j}))')
        x.number_format, x.fill, x.border = RAND, FILL_AUTO, BOX
    total_row(a0 + 1 + _AREAS, "Lines with no area", f'=SUMIFS(BOQ!$J${FR}:$J${LR},BOQ!$M${FR}:$M${LR},"")')
    wb.defined_names["AreaNames"] = DefinedName("AreaNames", attr_text=f"Summary!$A${a0 + 1}:$A${a0 + _AREAS}")
    sm.column_dimensions["A"].width, sm.column_dimensions["B"].width, sm.column_dimensions["C"].width = 34, 18, 70
    page(sm, False)
    sm.page_setup.fitToHeight = 1

    # ── Read me ──
    rm = wb.create_sheet("Read me")
    put_logo(rm, "A1", 50)
    rm["B2"] = "How to use this Bill of Quantities"
    rm["B2"].font = F(14, True, A)
    text = [
        ("1  Cover", "Fill in the client, project and bill details (yellow). They appear on every page."),
        ("2  BOQ", "Pick a Section, then pick an item from 'Pick from rate library' — description, unit and rate fill in. "
                   "Enter the Qty; the Amount is worked out."),
        ("", "Different price for this job? Type it in 'Your rate (R)'. Item not in the library? Leave the pick empty "
             "and type the Description, Unit and 'Your rate (R)'."),
        ("", "Area / Room (optional): type the room names on the Summary, then pick one per line. The bill stays "
             "grouped by trade; the Summary also totals by room."),
        ("3  Priced Bill", "The client copy: only the lines you've used, in trade order, ready to print or save as PDF."),
        ("4  Summary", "Totals by trade, by type and by room, plus contingency, site management, fee and VAT (yellow cells)."),
        ("5  Rate Library", f"{len(lib)} prices from your own documents, each with its source and date. Add your own in "
                            "the yellow rows at the bottom — they appear in every dropdown."),
        ("", ""),
        ("Yellow", "You fill these in."),
        ("Grey", "Worked out automatically (protected so they can't be overwritten by accident)."),
        ("Blue text", "Your inputs."),
        ("", ""),
        ("Prices", "Your rate list first, then what your invoices, quotes and priced bills show over the last 18 "
                   "months: the middle of what was paid, or of what was quoted when nothing was paid yet. "
                   "'Bill rate' items are all-in rates from priced bills; supplier prices are materials only; "
                   "labour rates are labour only."),
        ("Sections", "Trades in the order of the ASAQS model bills (Standard System of Measuring Building Work)."),
        ("Units", "The standard units: m, m², m³, No, kg, t, Item, Sum, Prov Sum, PC Sum. Supplier prices keep "
                  "their pack size (bag, tin, sheet): they are materials prices, not measured items."),
        ("Basis", "Choose on the Cover how the quantities were measured. Only choose ASAQS if the bill was "
                  "measured to the Standard System — the Cover then says so."),
        ("Updated", f"{today.day} {today:%b %Y}. Ask Vula for a fresh template to get the latest prices."),
        ("", ""),
        ("Send to Vula", f"Send the finished workbook to {b['display_name']}'s Vula WhatsApp with the project name as "
                         "the caption — Vula reads every BOQ line and sets the project budget from it."),
        ("Printing", "Each sheet is set up for A4 with your logo, page numbers and company details. The BOQ and "
                     "Priced Bill print only as far as their last line in Excel."),
        ("Protection", "Sheets are protected without a password so formulas stay intact. Review › Unprotect sheet "
                       "if you need to change the layout."),
    ]
    for i, (a, t) in enumerate(text, 4):
        rm.cell(i, 1, a).font = F(10, True)
        rm.cell(i, 2, t).font = F(10)
        rm.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if a == "Yellow":
            rm.cell(i, 1).fill = FILL_IN
        elif a == "Grey":
            rm.cell(i, 1).fill = FILL_AUTO
        elif a == "Blue text":
            rm.cell(i, 1).font = F(10, True, "0000FF")
    rm.column_dimensions["A"].width, rm.column_dimensions["B"].width = 18, 110
    page(rm, False)

    # ── protection, order, print areas ──
    for ws in (cv, bq, pb, sm, rl):
        ws.protection.sheet = True
        ws.protection.formatColumns = ws.protection.formatRows = False
        ws.protection.autoFilter = ws.protection.sort = False
    wb.move_sheet("Read me", offset=-(len(wb.sheetnames) - 2))
    wb.move_sheet("Lists", offset=len(wb.sheetnames))
    wb.active = 0
    for name, formula in (("BOQ", bq_print), ("Priced Bill", pb_print)):
        wb[name].defined_names["_xlnm.Print_Area"] = DefinedName("_xlnm.Print_Area", attr_text=formula)
    wb.calculation.fullCalcOnLoad = True          # no cached values: Excel / Sheets work them out on open
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def link(tenant_id: str) -> Tuple[Optional[str], Dict[str, Any]]:
    """Build, store privately (documents bucket) and return a signed link that lasts a week."""
    from vula.api.whatsapp import _upload_to_storage
    from vula.storage_links import SHARE_TTL, signed
    data, info = build(tenant_id)
    url = _upload_to_storage("documents", f"{tenant_id}/templates/BOQ_Template_{date.today().isoformat()}.xlsx",
                             data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    return (signed(url, SHARE_TTL) if url else None), info

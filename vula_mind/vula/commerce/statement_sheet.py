"""
vula/commerce/statement_sheet.py — import a bank statement the owner (or their bookkeeper) has
already categorised in a spreadsheet, keeping every allocation.

2026-09-28 (DIGG): Ian sent "DIGG_63214254607_Statement_Breakdown.xlsx" — 340 lines, 10 Jul –
12 Sep, each already given a Category ("HPC001 project", "Owner drawings", "Materials (card)"…)
and Sub-category (the trade: "Labour, wages & subcontractors", "Glazing & tinting"…). Vula's
statement import was PDF-only through an LLM, which would have thrown that work away. Here the
sheet's own columns are read as they are:

- a Category naming a project ("HPC001 project", "Sporty.TV project", "22 Porterfield") → the
  line's project (mapped onto the tenant's project name — preview() suggests the mapping, the
  owner confirms it), its Sub-category → the trade;
- an income Sub-category naming one ("HPC001 payment certificate") → that project's receipts;
- every other Category → its chart-of-accounts code (drawings, travel, meals, software, bank
  charges, insurance…) — the business's overheads.

Rows are saved categorized_by='owner' (reconcile never re-categorises them), deduped on the
statement key, and each one teaches the allocation rules (vula/commerce/allocation.py) so the
next PDF statement is allocated the same way. Money is integer cents.
"""
from __future__ import annotations

import csv
import logging
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

log = logging.getLogger(__name__)

_COLS = {
    "date": re.compile(r"^(transaction )?date$|^txn date$|^date of transaction$"),
    "description": re.compile(r"^description$|^details$|^narrative$"),
    "payee": re.compile(r"^payee$|^counterparty$|^merchant$"),
    "in": re.compile(r"^(money in|credit|credits|deposit|deposits|in|amount in)$"),
    "out": re.compile(r"^(money out|debit|debits|withdrawal|withdrawals|out|amount out)$"),
    "amount": re.compile(r"^amount$"),
    "category": re.compile(r"^category$|^account$"),
    "sub": re.compile(r"^sub[\s-]?category$|^trade$|^detail$"),
}

_PROJECT_CAT = re.compile(r"\bproject\b", re.IGNORECASE)
_ACCOUNT_FOR = [
    (re.compile(r"drawing|household", re.I), "owner_drawings"),
    (re.compile(r"personal", re.I), "owner_drawings"),
    (re.compile(r"travel|vehicle|fuel|parking", re.I), "fuel"),
    (re.compile(r"bank (charge|fee)|interest", re.I), "bank_charges"),
    (re.compile(r"insurance", re.I), "insurance"),
    (re.compile(r"software|it\b|internet|connectivity|domain|hosting", re.I), "utilities"),
    (re.compile(r"professional|accounting|legal", re.I), "professional_fees"),
    (re.compile(r"material|hardware|merchant", re.I), "cost_of_sales"),
    (re.compile(r"labou?r|wage|salar|subcontract", re.I), "casual_labour"),
]
_LABOUR = re.compile(r"labou?r|wage|salar|subcontract", re.I)
_OVERHEAD_CODES = {"owner_drawings", "fuel", "bank_charges", "insurance", "utilities",
                   "professional_fees", "other_expense", "rent", "marketing", "wages"}


def _clean(v: Any) -> str:
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def _cents(v: Any) -> Optional[int]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(round(abs(float(v)) * 100))
    s = re.sub(r"[R\s,]", "", str(v))
    try:
        return int(round(abs(float(s)) * 100))
    except ValueError:
        return None


def _date(v: Any) -> Optional[str]:
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = _clean(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d", "%d-%m-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(s[:len(datetime.now().strftime(fmt)) + 4].strip(), fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _header(row: List[Any]) -> Optional[Dict[str, int]]:
    cols: Dict[str, int] = {}
    for i, cell in enumerate(row):
        label = _clean(cell).lower()
        for key, pat in _COLS.items():
            if key not in cols and pat.match(label):
                cols[key] = i
                break
    ok = "date" in cols and "description" in cols and ("amount" in cols or ("in" in cols and "out" in cols))
    return cols if ok else None


def _rows(path: Path) -> List[List[Any]]:
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
            return [list(r) for r in csv.reader(fh)]
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        # The transactions sheet is the one with a recognisable header; a Summary sheet first
        # (as in DIGG's workbook) is skipped.
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            if any(_header(r) for r in rows[:15]):
                return rows
    finally:
        wb.close()
    return []


def parse(path: Path) -> List[Dict[str, Any]]:
    """Statement lines: date, description, payee, direction, amount_cents, category, sub."""
    rows = _rows(Path(path))
    cols = None
    out: List[Dict[str, Any]] = []
    for row in rows:
        if cols is None:
            cols = _header(row)
            continue

        def cell(k):
            i = cols.get(k)
            return row[i] if i is not None and i < len(row) else None
        d = _date(cell("date"))
        desc = _clean(cell("description"))
        if not d or not desc:
            continue
        cin, cout = _cents(cell("in")), _cents(cell("out"))
        if "amount" in cols and cin is None and cout is None:
            raw = cell("amount")
            try:
                val = float(re.sub(r"[R\s,]", "", str(raw)))
            except (TypeError, ValueError):
                continue
            cin, cout = (int(round(val * 100)), None) if val > 0 else (None, int(round(-val * 100)))
        if cin:
            direction, amount = "in", cin
        elif cout:
            direction, amount = "out", cout
        else:
            continue
        out.append({"date": d, "description": desc, "payee": _clean(cell("payee")) or None,
                    "direction": direction, "amount_cents": amount,
                    "category": _clean(cell("category")) or None, "sub": _clean(cell("sub")) or None})
    return out


def project_label(line: Dict[str, Any]) -> Optional[str]:
    """The project a line names, as written in the sheet: "HPC001 project" → "HPC001";
    an income line "HPC001 payment certificate" → "HPC001"; "22 Porterfield" (a category that
    names a property with a number) → "22 Porterfield"."""
    cat, sub = line.get("category") or "", line.get("sub") or ""
    if _PROJECT_CAT.search(cat):
        return _PROJECT_CAT.sub("", cat).strip(" -–—") or None
    if line["direction"] == "in":
        m = re.match(r"^([A-Za-z]+[\w.]*\d+[\w.]*|[A-Z][\w.]+\.TV)\b.*\b(certificate|claim|payment)\b", sub)
        if m:
            return m.group(1)
    if re.match(r"^\d+\s+[A-Za-z]", cat):
        return cat
    return None


def _alpha_token(label: str) -> str:
    from vula.commerce.service import project_key
    for tok in project_key(label).split():
        tok = re.sub(r"\d+", "", tok)
        if len(tok) >= 3:
            return tok
    return ""


def suggest_projects(tenant_id: str, labels: List[str]) -> Dict[str, Optional[str]]:
    """Each sheet label → the tenant's project it most likely is: an exact name/number match
    (canonical_project), else the known project sharing its first word ("HPC001" ↔ "HPC Bokaap",
    "Sporty.TV" ↔ "Sporty – Phase 2", "22 Porterfield" ↔ "Porterfield"). None = new project."""
    from vula.commerce import service
    known: Dict[str, int] = {}
    try:
        db = service._client()
        for r in (db.table("vula_projects").select("name").eq("tenant_id", tenant_id)
                  .limit(500).execute().data or []):
            if r.get("name"):
                known[r["name"]] = known.get(r["name"], 0) + 1000
        for r in (db.table("vula_filed_documents").select("project").eq("tenant_id", tenant_id)
                  .limit(5000).execute().data or []):
            if r.get("project"):
                known[r["project"]] = known.get(r["project"], 0) + 1
    except Exception as exc:
        log.debug("project suggestions skipped: %s", exc)
    out: Dict[str, Optional[str]] = {}
    for label in labels:
        exact = service.canonical_project(tenant_id, label)
        if exact and exact != label:
            out[label] = exact
            continue
        tok = _alpha_token(label)
        cands = [(n, w) for n, w in known.items() if tok and _alpha_token(n) == tok]
        out[label] = max(cands, key=lambda c: c[1])[0] if cands else None
    return out


def account_for(line: Dict[str, Any], project: Optional[str]) -> str:
    if line["direction"] == "in":
        text = f"{line.get('sub') or ''} {line.get('category') or ''}"
        if re.search(r"revers|return|refund|transfer|loan", text, re.I):
            return "other_income"       # money back or moved in — not a sale
        if project or re.search(r"client|certificate|invoice|fee|deposit|payment", text, re.I):
            return "sales"
        return "other_income"
    text = f"{line.get('category') or ''} {line.get('sub') or ''}"
    if project:
        return "casual_labour" if _LABOUR.search(line.get("sub") or "") else "cost_of_sales"
    for pat, code in _ACCOUNT_FOR:
        if pat.search(text):
            return code
    return "other_expense"


def _existing_in_period(tenant_id: str, first: Optional[str], last: Optional[str]) -> List[Dict[str, Any]]:
    """Bank lines Vula already holds for the sheet's dates that came from somewhere else (weekly
    PDF statements), not yet matched to anything — what the sheet would double-count."""
    if not first or not last:
        return []
    from vula.commerce.ledger import _all_pages
    from vula.commerce.service import _client

    def make():
        return (_client().table("commerce_bank_transactions")
                .select("id,source_file,match_status,matched_invoice_id,matched_expense_id,matched_order_id")
                .eq("tenant_id", tenant_id).gte("txn_date", first).lte("txn_date", last))
    try:
        rows = _all_pages(make)
    except Exception as exc:
        log.debug("existing-lines read skipped: %s", exc)
        return []
    return [r for r in rows
            if not str(r.get("source_file") or "").lower().endswith((".xlsx", ".xlsm", ".csv"))
            and r.get("match_status") not in ("matched", "ignored")
            and not (r.get("matched_invoice_id") or r.get("matched_expense_id") or r.get("matched_order_id"))]


def preview(tenant_id: str, path: Path) -> Dict[str, Any]:
    lines = parse(path)
    labels = sorted({p for p in (project_label(li) for li in lines) if p})
    sums: Dict[str, Dict[str, int]] = {}
    for li in lines:
        p = project_label(li)
        if p:
            s = sums.setdefault(p, {"in": 0, "out": 0, "lines": 0})
            s[li["direction"]] += li["amount_cents"]
            s["lines"] += 1
    first = min((li["date"] for li in lines), default=None)
    last = max((li["date"] for li in lines), default=None)
    return {"lines": len(lines), "first": first, "last": last,
            # Lines Vula already has for these dates (from PDF statements) — importing on top
            # would count that money twice; import with replace_existing to set them aside.
            "existing_lines_in_period": len(_existing_in_period(tenant_id, first, last)),
            "money_in_cents": sum(li["amount_cents"] for li in lines if li["direction"] == "in"),
            "money_out_cents": sum(li["amount_cents"] for li in lines if li["direction"] == "out"),
            "projects": [{"label": p, "suggested": s, **sums[p]}
                         for p, s in suggest_projects(tenant_id, labels).items()]}


def import_sheet(tenant_id: str, path: Path, project_map: Optional[Dict[str, Optional[str]]] = None,
                 source_file: str = "statement.xlsx", replace_existing: bool = False) -> Dict[str, Any]:
    """Save the sheet's lines with their allocations. `project_map` maps a sheet label to the
    tenant's project name (from preview); an unmapped label is used as written."""
    from vula.commerce import accounting, allocation
    from vula.commerce.service import _client, canonical_project
    lines = parse(path)
    project_map = project_map or {}
    db = _client()
    set_aside = 0
    if replace_existing and lines:
        # The sheet is the complete, categorised record for its dates: the unmatched PDF-read
        # lines for the same dates are set aside (match_status 'ignored' — reversible, nothing
        # deleted) so no rand is counted twice.
        first, last = min(li["date"] for li in lines), max(li["date"] for li in lines)
        for r in _existing_in_period(tenant_id, first, last):
            try:
                (db.table("commerce_bank_transactions").update({"match_status": "ignored"})
                 .eq("tenant_id", tenant_id).eq("id", r["id"]).execute())
                set_aside += 1
            except Exception as exc:
                log.warning("could not set aside bank line %s: %s", r.get("id"), exc)
    acc_map = {a["code"]: a for a in accounting.ensure_chart(tenant_id)}
    vat_reg = accounting.is_vat_registered(tenant_id)
    saved, failed, projects = 0, 0, {}
    canon: Dict[str, Optional[str]] = {}
    # The statement key is (date, amount, description): two real R500 "PAY" lines on one day
    # would collapse into one. The nth repeat within the sheet gets " (n)" — deterministic, so
    # importing the same sheet again still dedupes.
    seen: Dict[tuple, int] = {}
    for li in lines:
        k = (li["date"], li["amount_cents"], li["description"])
        seen[k] = seen.get(k, 0) + 1
        if seen[k] > 1:
            li["description"] = f"{li['description']} ({seen[k]})"
    for li in lines:
        label = project_label(li)
        project = None
        if label:
            if label not in canon:
                canon[label] = project_map.get(label) or canonical_project(tenant_id, label)
            project = canon[label]
        code = account_for(li, project)
        if code not in acc_map:
            code = "other_expense" if li["direction"] == "out" else "other_income"
        trade = li.get("sub") if project and li["direction"] == "out" else None
        row = {
            "tenant_id": tenant_id, "txn_date": li["date"], "description": li["description"],
            "amount_cents": li["amount_cents"], "direction": li["direction"],
            "payee": li.get("payee"), "match_status": "unmatched", "source_file": source_file,
            "account_code": code, "categorized_by": "owner",
            "vat_cents": accounting.vat_for(acc_map.get(code), li["amount_cents"], vat_reg),
            "vat_treatment": (acc_map.get(code) or {}).get("vat_treatment"),
            "project": project, "trade": trade,
        }
        try:
            db.table("commerce_bank_transactions").upsert(
                row, on_conflict="tenant_id,txn_date,amount_cents,description").execute()
            saved += 1
        except Exception as exc:
            failed += 1
            log.warning("statement sheet row skipped: %s", exc)
            continue
        if project:
            projects[project] = projects.get(project, 0) + 1
            allocation.learn(tenant_id, li["description"], project, trade, payee=li.get("payee"))
    # What the sheet taught now allocates the other lines Vula holds (earlier/later statements).
    also_allocated = allocation.apply_to_existing(tenant_id) if projects else 0
    return {"parsed": len(lines), "saved": saved, "failed": failed, "projects": projects,
            "set_aside": set_aside, "also_allocated": also_allocated}

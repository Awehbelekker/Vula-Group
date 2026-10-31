"""
vula/commerce/sheet_builder.py — a spreadsheet the owner describes in their own words, built from the books.

31 Oct 2026 (Ian: "plan the sheet builder"). Vula only had fixed exports (supplier invoices,
BOQ template, expense sheet, project bill). "Every Tile Co invoice this year by project, VAT split,
totalled per month" got text or the nearest fixed export.

How it works — the model never supplies a figure:
  1. The model turns the owner's words into a SPEC: a data source, filters, columns from that
     source's fixed catalogue, a grouping and the totals. Code validates it (unknown columns are
     refused with the valid list).
  2. design: the rows are read from the books (tenant-scoped, paged), the filters resolved by the
     same helpers the rest of Vula uses (period, supplier names and aliases, project spellings), and
     the owner gets back what was understood, the first rows and the totals. Kept as a draft.
  3. The owner corrects it in words; each correction is a new spec and a new preview.
  4. build: the file is written from the draft — group subtotals and the grand total as SUM
     formulas — then re-opened and evaluated (xlsx_check) and must equal the books' totals to the
     cent; and the totals must still be the preview's, or the owner is told the books changed.
  5. save / use: the agreed layout is kept as the business's own template ("the Tile Co monthly
     sheet") and re-run with a new period.

Sources (v1): supplier invoices, supplier invoice lines, bank lines, expenses and slips, customer
invoices and orders. Money is integer cents until written; refunds and credit notes are negative,
money out of the bank is negative, so every SUM equals the books.
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, Dict, List, Optional, Tuple

log = logging.getLogger(__name__)

TABLE = "vula_sheet_specs"
MAX_ROWS = 5000
PREVIEW_ROWS = 5
DRAFT_DAYS = 7
MAX_GROUPS = 300
PERIODS = ("this_month", "last_month", "this_year", "last_year", "tax_year", "last_tax_year",
           "last_12_months", "all")
GROUPS = ("none", "month", "project", "supplier", "customer", "account", "status")


@dataclass(frozen=True)
class Field:
    id: str
    label: str
    kind: str = "text"            # text | date | money | number

    @property
    def money(self) -> bool:
        return self.kind == "money"


def _f(*specs: Tuple[str, str, str]) -> Dict[str, Field]:
    return {i: Field(i, label, kind) for i, label, kind in specs}


@dataclass(frozen=True)
class Source:
    id: str
    label: str
    scope: str                                    # the dashboard access it needs
    fields: Dict[str, Field]
    default_columns: Tuple[str, ...]
    filters: Tuple[str, ...]                      # which filters apply
    loader: str                                   # function name below
    party: str = "supplier"                       # the field a supplier/customer filter reads


SOURCES: Dict[str, Source] = {
    "supplier_invoices": Source(
        "supplier_invoices", "Supplier invoices", "finances",
        _f(("date", "Date", "date"), ("number", "Invoice no.", "text"), ("supplier", "Supplier", "text"),
           ("project", "Project", "text"), ("excl_vat", "Excl VAT", "money"), ("vat", "VAT", "money"),
           ("total", "Total", "money"), ("category", "Type", "text"), ("status", "Status", "text")),
        ("date", "number", "supplier", "project", "excl_vat", "vat", "total"),
        ("period", "supplier", "project", "status"), "_supplier_invoices"),
    "supplier_lines": Source(
        "supplier_lines", "Supplier invoice lines", "finances",
        _f(("date", "Date", "date"), ("number", "Invoice no.", "text"), ("supplier", "Supplier", "text"),
           ("project", "Project", "text"), ("description", "Item", "text"), ("quantity", "Qty", "number"),
           ("unit_price", "Unit price", "money"), ("line_total", "Line total", "money")),
        ("date", "number", "supplier", "project", "description", "quantity", "unit_price", "line_total"),
        ("period", "supplier", "project"), "_supplier_lines"),
    "bank": Source(
        "bank", "Bank lines", "finances",
        _f(("date", "Date", "date"), ("description", "Description", "text"), ("payee", "Payee", "text"),
           ("reference", "Reference", "text"), ("money_in", "Money in", "money"),
           ("money_out", "Money out", "money"), ("amount", "Amount (+in / −out)", "money"),
           ("project", "Project", "text"), ("account", "Category", "text"), ("status", "Matched", "text")),
        ("date", "description", "money_in", "money_out", "project", "account"),
        ("period", "supplier", "project", "account", "status"), "_bank", party="payee"),
    "expenses": Source(
        "expenses", "Expenses and slips", "finances",
        _f(("date", "Date", "date"), ("supplier", "Supplier", "text"), ("description", "Description", "text"),
           ("amount", "Amount", "money"), ("vat", "VAT", "money"), ("account", "Category", "text"),
           ("project", "Project", "text"), ("paid_with", "Paid with", "text"), ("paid_by", "Paid by", "text"),
           ("reimbursable", "Owed back", "text"), ("status", "Status", "text")),
        ("date", "supplier", "description", "amount", "vat", "account", "project", "paid_with"),
        ("period", "supplier", "project", "account", "status"), "_expenses"),
    "customer_invoices": Source(
        "customer_invoices", "Customer invoices", "finances",
        _f(("number", "Invoice no.", "text"), ("type", "Type", "text"), ("customer", "Customer", "text"),
           ("date", "Date", "date"), ("due", "Due", "date"), ("project", "Project", "text"),
           ("excl_vat", "Excl VAT", "money"), ("vat", "VAT", "money"), ("total", "Total", "money"),
           ("paid", "Paid", "money"), ("owed", "Still owed", "money"), ("status", "Status", "text")),
        ("number", "customer", "date", "project", "total", "paid", "owed", "status"),
        ("period", "customer", "project", "status", "type"), "_customer_invoices", party="customer"),
    "orders": Source(
        "orders", "Orders", "orders",
        _f(("number", "Order", "text"), ("date", "Date", "date"), ("customer", "Customer", "text"),
           ("channel", "Channel", "text"), ("status", "Status", "text"), ("subtotal", "Subtotal", "money"),
           ("delivery", "Delivery", "money"), ("total", "Total", "money")),
        ("number", "date", "customer", "status", "total"),
        ("period", "customer", "status"), "_orders", party="customer"),
}

_GROUP_FIELD = {"month": "date", "project": "project", "supplier": "supplier", "customer": "customer",
                "account": "account", "status": "status"}


class SpecError(ValueError):
    """The spec can't be built as given; the message says what to change, in plain words."""


# ── the spec ───────────────────────────────────────────────────────────────

def validate(spec: Dict[str, Any]) -> Dict[str, Any]:
    """A clean spec, or SpecError naming what's wrong and what is allowed."""
    if not isinstance(spec, dict):
        raise SpecError("I need to know which data the sheet is about.")
    src = SOURCES.get(str(spec.get("source") or ""))
    if not src:
        raise SpecError("Which data should the sheet hold? " + ", ".join(s.label.lower() for s in SOURCES.values()) + ".")
    cols = [str(c) for c in (spec.get("columns") or [])] or list(src.default_columns)
    bad = [c for c in cols if c not in src.fields]
    if bad:
        raise SpecError(f"{src.label} can't have the column(s) {', '.join(bad)}. Columns it can have: "
                        + ", ".join(f"{f.id} ({f.label})" for f in src.fields.values()) + ".")
    cols = list(dict.fromkeys(cols))
    group = str(spec.get("group_by") or "none")
    if group not in GROUPS:
        raise SpecError("It can be grouped by " + ", ".join(GROUPS[1:]) + ", or not at all.")
    if group != "none" and _GROUP_FIELD[group] not in src.fields:
        raise SpecError(f"{src.label} can't be grouped by {group}.")
    money = [c for c in cols if src.fields[c].money]
    totals = [str(t) for t in (spec.get("totals") or [])] or [c for c in money if c != "unit_price"]
    bad = [t for t in totals if t not in money]
    if bad:
        raise SpecError(f"Only money columns in the sheet can be totalled: {', '.join(money) or 'none'}.")
    raw = spec.get("filters") or {}
    filters: Dict[str, Any] = {}
    for k, v in raw.items():
        if v in (None, "", []):
            continue
        if k in ("since", "until"):
            try:
                filters[k] = date.fromisoformat(str(v)[:10]).isoformat()
            except ValueError:
                raise SpecError(f"I couldn't read the date '{v}' — use e.g. 2026-10-01.")
            continue
        if k == "period":
            if v not in PERIODS:
                raise SpecError("The period can be " + ", ".join(PERIODS) + ", or exact dates.")
            filters[k] = v
            continue
        if k not in src.filters:
            raise SpecError(f"{src.label} can't be filtered by {k}. It can be filtered by "
                            + ", ".join(src.filters) + ".")
        filters[k] = str(v).strip()[:120]
    title = re.sub(r"\s+", " ", str(spec.get("title") or "")).strip()[:90] or src.label
    return {"source": src.id, "columns": cols, "group_by": group, "totals": totals,
            "filters": filters, "title": title}


def period_range(filters: Dict[str, Any], today: Optional[date] = None) -> Tuple[Optional[str], Optional[str]]:
    """(since, until) from exact dates or a named period, worked out here — never by a model."""
    from vula.commerce import reports
    t = today or date.today()
    if filters.get("since") or filters.get("until"):
        return filters.get("since"), filters.get("until")
    p = filters.get("period")
    if not p or p == "all":
        return None, None
    if p == "last_year":
        return date(t.year - 1, 1, 1).isoformat(), date(t.year - 1, 12, 31).isoformat()
    return reports.period_dates(p, today=t)


# ── reading the books ──────────────────────────────────────────────────────

def _db():
    from vula.commerce import service
    return service._client()


def _pages(make) -> List[Dict[str, Any]]:
    from vula.commerce.ledger import _all_pages
    out = _all_pages(make)
    if len(out) > MAX_ROWS * 4:
        raise SpecError("That's more than the sheet can hold — narrow it down (a period, a supplier or a project).")
    return out


def _labels(tenant_id: str) -> Dict[str, str]:
    try:
        from vula.commerce.accounting import ensure_chart
        return {a["code"]: a.get("name") or a["code"] for a in ensure_chart(tenant_id)}
    except Exception:  # noqa: BLE001 — a label is never why a sheet fails
        return {}


def _filed(tenant_id: str, since: Optional[str], until: Optional[str]) -> List[Dict[str, Any]]:
    from vula.commerce import service

    def make():
        return (_db().table("vula_filed_documents")
                .select("id,filename,category,summary,fields,created_at,project,status,doc_status")
                .eq("tenant_id", tenant_id).eq("category", "Invoice").order("created_at"))
    out = []
    for r in _pages(make):
        if r.get("doc_status") == "superseded":
            continue
        e = service._export_row(r)
        e["project"] = r.get("project") or ""
        e["status"] = r.get("status") or ""
        if _in_range(e["date"], since, until):
            out.append(e)
    return out


def _in_range(d: str, since: Optional[str], until: Optional[str]) -> bool:
    d = (d or "")[:10]
    if since and (not d or d < since):
        return False
    if until and (not d or d > until):
        return False
    return True


def _supplier_invoices(tenant_id, since, until, labels):
    out = []
    for e in _filed(tenant_id, since, until):
        total, vat = e["total_cents"], e["vat_cents"]
        out.append({"date": e["date"], "number": e["ref"], "supplier": e["party"], "project": e["project"],
                    "excl_vat": total - vat if total is not None and vat is not None else total,
                    "vat": vat, "total": total, "category": "Refund / credit" if e["is_refund"] else e["category"],
                    "status": e["status"]})
    return out


def _supplier_lines(tenant_id, since, until, labels):
    out = []
    for e in _filed(tenant_id, since, until):
        for li in e["lines"]:
            out.append({"date": e["date"], "number": e["ref"], "supplier": e["party"], "project": e["project"],
                        "description": li["description"], "quantity": li["quantity"],
                        "unit_price": li["unit_cents"], "line_total": li["total_cents"]})
    return out


def _bank(tenant_id, since, until, labels):
    def make():
        q = (_db().table("commerce_bank_transactions")
             .select("id,txn_date,description,payee,reference,amount_cents,direction,account_code,project,"
                     "match_status")
             .eq("tenant_id", tenant_id))
        if since:
            q = q.gte("txn_date", since)
        if until:
            q = q.lte("txn_date", until)
        return q.order("txn_date")
    out = []
    for r in _pages(make):
        if r.get("match_status") == "ignored":
            continue
        cents = abs(int(r.get("amount_cents") or 0))
        incoming = (r.get("direction") or "").lower() == "in"
        code = r.get("account_code") or ""
        out.append({"date": str(r.get("txn_date") or "")[:10], "description": r.get("description") or "",
                    "payee": r.get("payee") or "", "reference": r.get("reference") or "",
                    "money_in": cents if incoming else 0, "money_out": 0 if incoming else cents,
                    "amount": cents if incoming else -cents, "project": r.get("project") or "",
                    "account": labels.get(code, code), "account_code": code,
                    "status": r.get("match_status") or ""})
    return out


def _expenses(tenant_id, since, until, labels):
    def make():
        q = (_db().table("commerce_expenses")
             .select("id,date,supplier,description,amount_cents,vat_cents,account_code,category,project,"
                     "paid_with,paid_by_name,reimbursable,status")
             .eq("tenant_id", tenant_id))
        if since:
            q = q.gte("date", since)
        if until:
            q = q.lte("date", until)
        return q.order("date")
    out = []
    for r in _pages(make):
        code = r.get("account_code") or ""
        out.append({"date": str(r.get("date") or "")[:10], "supplier": r.get("supplier") or "",
                    "description": r.get("description") or "", "amount": int(r.get("amount_cents") or 0),
                    "vat": int(r["vat_cents"]) if r.get("vat_cents") is not None else None,
                    "account": labels.get(code, code or r.get("category") or ""), "account_code": code,
                    "project": r.get("project") or "", "paid_with": r.get("paid_with") or "",
                    "paid_by": r.get("paid_by_name") or "",
                    "reimbursable": "yes" if r.get("reimbursable") else "", "status": r.get("status") or ""})
    return out


_DOC_TYPE_LABEL = {"invoice": "Invoice", "credit_note": "Credit note", "quote": "Quote",
                   "proforma": "Pro forma", "delivery_note": "Delivery note"}


def _customer_invoices(tenant_id, since, until, labels):
    def make():
        q = (_db().table("commerce_invoices")
             .select("id,invoice_number,doc_type,direction,customer_name,issue_date,due_date,project,"
                     "subtotal_cents,vat_cents,total_cents,total_paid_cents,status")
             .eq("tenant_id", tenant_id))
        if since:
            q = q.gte("issue_date", since)
        if until:
            q = q.lte("issue_date", until)
        return q.order("issue_date")
    out = []
    for r in _pages(make):
        if (r.get("direction") or "outbound") == "inbound":
            continue                                   # supplier bills are the supplier sources
        kind = r.get("doc_type") or "invoice"
        sign = -1 if kind == "credit_note" else 1
        total = sign * abs(int(r.get("total_cents") or 0))
        vat = sign * abs(int(r.get("vat_cents") or 0))
        paid = sign * abs(int(r.get("total_paid_cents") or 0))
        out.append({"number": r.get("invoice_number") or "", "type": _DOC_TYPE_LABEL.get(kind, kind),
                    "doc_type": kind, "customer": r.get("customer_name") or "",
                    "date": str(r.get("issue_date") or "")[:10], "due": str(r.get("due_date") or "")[:10],
                    "project": r.get("project") or "",
                    "excl_vat": sign * abs(int(r.get("subtotal_cents") or 0)) if r.get("subtotal_cents") is not None else total - vat,
                    "vat": vat, "total": total, "paid": paid, "owed": total - paid, "status": r.get("status") or ""})
    return out


def _orders(tenant_id, since, until, labels):
    def make():
        q = (_db().table("commerce_orders")
             .select("id,display_id,created_at,customer_name,channel,status,subtotal_cents,delivery_cents,total_cents")
             .eq("tenant_id", tenant_id))
        if since:
            q = q.gte("created_at", since)
        if until:
            q = q.lte("created_at", until + "T23:59:59")
        return q.order("created_at")
    return [{"number": r.get("display_id") or "", "date": str(r.get("created_at") or "")[:10],
             "customer": r.get("customer_name") or "", "channel": r.get("channel") or "",
             "status": r.get("status") or "", "subtotal": int(r.get("subtotal_cents") or 0),
             "delivery": int(r.get("delivery_cents") or 0), "total": int(r.get("total_cents") or 0)}
            for r in _pages(make)]


_LOADERS: Dict[str, Callable] = {"_supplier_invoices": _supplier_invoices, "_supplier_lines": _supplier_lines,
                                 "_bank": _bank, "_expenses": _expenses,
                                 "_customer_invoices": _customer_invoices, "_orders": _orders}
# Left out unless the owner asks for that status: they're not real money.
_NOT_COUNTED = {"expenses": {"cancelled", "void", "deleted"},
                "customer_invoices": {"cancelled", "void", "draft"},
                "orders": {"cancelled", "abandoned", "pending_payment"}}


# ── filters ────────────────────────────────────────────────────────────────

def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (s or "").lower()).strip()


def _party_match(row_value: str, names: List[str], query: str) -> bool:
    """A row's supplier/customer/payee is the one asked for: one of the resolved names/aliases
    (whole words), or — with no known supplier — the words asked for, all present."""
    v = f" {_norm(row_value)} "
    if names:
        from vula.commerce.service import _canonical_party
        canon = {_norm(_canonical_party(n)) for n in names} | {_norm(n) for n in names}
        return any(c and f" {c} " in v for c in canon)
    words = [w for w in _norm(query).split() if len(w) > 1]
    return bool(words) and all(f" {w} " in v or f" {w}" in v for w in words)


def _status_ok(source: str, row: Dict[str, Any], wanted: Optional[str]) -> bool:
    st = (row.get("status") or "").lower()
    if wanted:
        return _norm(wanted) in _norm(st)
    return st not in _NOT_COUNTED.get(source, set())


async def resolve(tenant_id: str, spec: Dict[str, Any]) -> Dict[str, Any]:
    """The filters as Vula understood them: supplier names and aliases, the project's spelling."""
    from vula.commerce import service
    f = spec["filters"]
    out: Dict[str, Any] = {"names": []}
    q = f.get("supplier") or f.get("customer")
    if f.get("supplier"):
        out["names"] = await service._resolve_supplier_names(tenant_id, f["supplier"])
    if f.get("project"):
        out["project"] = service.canonical_project(tenant_id, f["project"]) or f["project"]
    out["query"] = q or ""
    return out


def rows_for(tenant_id: str, spec: Dict[str, Any], resolved: Dict[str, Any],
             today: Optional[date] = None) -> List[Dict[str, Any]]:
    """The sheet's rows, from the books, filtered as the spec says."""
    from vula.commerce.service import project_key
    src = SOURCES[spec["source"]]
    f = spec["filters"]
    since, until = period_range(f, today)
    rows = _LOADERS[src.loader](tenant_id, since, until, _labels(tenant_id))
    out = []
    pk = project_key(resolved.get("project")) if f.get("project") else None
    for r in rows:
        party = (f"{r.get('payee') or ''} {r.get('description') or ''}" if src.id == "bank"
                 else r.get(src.party) or "")
        if (f.get("supplier") or f.get("customer")) and not _party_match(
                party, resolved.get("names") or [], resolved.get("query") or ""):
            continue
        if pk is not None and project_key(r.get("project")) != pk:
            continue
        if f.get("account") and _norm(f["account"]) not in _norm(f"{r.get('account')} {r.get('account_code')}"):
            continue
        if f.get("type") and _norm(f["type"]) not in _norm(f"{r.get('type')} {r.get('doc_type')}"):
            continue
        if not _status_ok(src.id, r, f.get("status")):
            continue
        out.append(r)
    if len(out) > MAX_ROWS:
        raise SpecError(f"That's {len(out):,} rows — more than a sheet should hold. Narrow it down "
                        "(a period, a supplier or a project).")
    return out


def _group_key(group: str, row: Dict[str, Any]) -> str:
    if group == "none":
        return ""
    v = row.get(_GROUP_FIELD[group]) or ""
    if group == "month":
        return str(v)[:7] or "No date"
    return str(v) or {"project": "No project", "account": "Uncategorised"}.get(group, "—")


def arrange(spec: Dict[str, Any], rows: List[Dict[str, Any]]) -> List[Tuple[str, List[Dict[str, Any]]]]:
    """Rows sorted by date within groups; groups in order (months by date, others A–Z)."""
    rows = sorted(rows, key=lambda r: (str(r.get("date") or ""), str(r.get("number") or "")))
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for r in rows:
        groups.setdefault(_group_key(spec["group_by"], r), []).append(r)
    if len(groups) > MAX_GROUPS:
        raise SpecError(f"That makes {len(groups)} groups — group it another way, or not at all.")
    return sorted(groups.items(), key=lambda kv: kv[0].lower())


def totals(spec: Dict[str, Any], rows: List[Dict[str, Any]]) -> Dict[str, int]:
    return {c: sum(int(r.get(c) or 0) for r in rows) for c in spec["totals"]}


# ── what the owner sees first ──────────────────────────────────────────────

def _r(cents: Optional[int]) -> str:
    if cents is None:
        return "—"
    return f"R{cents / 100:,.2f}" if cents >= 0 else f"−R{-cents / 100:,.2f}"


def _when(d: Optional[str]) -> str:
    try:
        return date.fromisoformat(d).strftime("%-d %b %Y") if d else ""
    except ValueError:
        return d or ""


def _cell_text(field: Field, v: Any) -> str:
    if field.money:
        return _r(v)
    if v is None or v == "":
        return "—"
    return str(v)


def understood(spec: Dict[str, Any], resolved: Dict[str, Any], today: Optional[date] = None) -> List[str]:
    src = SOURCES[spec["source"]]
    f = spec["filters"]
    since, until = period_range(f, today)
    out = [f"• Data: {src.label}"]
    if f.get("supplier"):
        names = resolved.get("names") or []
        out.append(f"• Supplier: {names[0]}" + (f" (also as {', '.join(names[1:4])})" if len(names) > 1 else "")
                   if names else f"• Supplier: names containing “{f['supplier']}”")
    if f.get("customer"):
        out.append(f"• Customer: names containing “{f['customer']}”")
    if f.get("project"):
        out.append(f"• Project: {resolved.get('project') or f['project']}")
    if f.get("account"):
        out.append(f"• Category: {f['account']}")
    if f.get("type"):
        out.append(f"• Type: {f['type']}")
    if f.get("status"):
        out.append(f"• Status: {f['status']}")
    elif spec["source"] in _NOT_COUNTED:
        out.append("• Left out: " + ", ".join(sorted(_NOT_COUNTED[spec["source"]])))
    out.append(f"• Period: {_when(since) or 'the start'} – {_when(until) or 'today'}" if since or until
               else "• Period: everything on record")
    out.append(f"• Grouped by: {spec['group_by']}" if spec["group_by"] != "none" else "• Not grouped")
    return out


def preview_text(spec: Dict[str, Any], resolved: Dict[str, Any], rows: List[Dict[str, Any]],
                 today: Optional[date] = None) -> str:
    src = SOURCES[spec["source"]]
    tot = totals(spec, rows)
    lines = [f"📊 *{spec['title']}* — here's what I understood:", *understood(spec, resolved, today),
             "Columns: " + " | ".join(src.fields[c].label for c in spec["columns"])]
    if not rows:
        lines.append("\nNothing in the books matches that. Tell me what to change (a different period, "
                     "supplier or project).")
        return "\n".join(lines)
    lines.append(f"\n{len(rows):,} row(s). The first {min(PREVIEW_ROWS, len(rows))}:")
    for r in sorted(rows, key=lambda r: str(r.get("date") or ""))[:PREVIEW_ROWS]:
        lines.append("– " + " · ".join(_cell_text(src.fields[c], r.get(c)) for c in spec["columns"][:6]))
    if tot:
        lines.append("\nTotals: " + " · ".join(f"{src.fields[c].label} {_r(v)}" for c, v in tot.items()))
    if spec["group_by"] != "none" and spec["totals"]:
        first = spec["totals"][-1]
        groups = arrange(spec, rows)
        shown = " · ".join(f"{k or '—'} {_r(sum(int(r.get(first) or 0) for r in g))}" for k, g in groups[:8])
        lines.append(f"{src.fields[first].label} per {spec['group_by']}: {shown}" + (" …" if len(groups) > 8 else ""))
    lines.append("\nReply *build it*, or tell me what to change (e.g. “only one project”, “drop VAT”, "
                 "“per project instead”).")
    return "\n".join(lines)


# ── the file ───────────────────────────────────────────────────────────────

def workbook(business: str, spec: Dict[str, Any], resolved: Dict[str, Any], rows: List[Dict[str, Any]],
             accent: Optional[str] = None, today: Optional[date] = None) -> Tuple[bytes, Dict[str, int]]:
    """The xlsx and the cells holding each grand total (ref -> cents) for the read-back."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    from vula.commerce import xlsx as X
    from vula.commerce.xlsx_check import safe_text
    src = SOURCES[spec["source"]]
    cols = spec["columns"]
    fill = X._fill_for(accent)
    token = X._fill_var.set(fill)
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = re.sub(r"[\[\]:*?/\\]", " ", spec["title"])[:31] or "Sheet"
        safe_text(ws["A1"], business)
        ws["A1"].font = Font(bold=True, size=12)
        safe_text(ws["A2"], spec["title"])
        ws["A2"].font = Font(bold=True, size=11)
        safe_text(ws["A3"], "  ".join(s.lstrip("• ") for s in understood(spec, resolved, today)))
        ws["A3"].font = X._MUTED
        ws["A4"] = f"From the books, as at {(today or date.today()).isoformat()}. Totals are formulas."
        ws["A4"].font = X._MUTED
        head = 6
        widths = [max(10, min(48, len(src.fields[c].label) + 4 if not src.fields[c].money else 15)) for c in cols]
        for i, c in enumerate(cols):
            if c in ("description", "supplier", "customer", "project", "account"):
                widths[i] = 32 if c == "description" else 24
        X._table(ws, head, [src.fields[c].label for c in cols], widths)
        money_cols = [i + 1 for i, c in enumerate(cols) if c in spec["totals"]]
        r = head + 1
        sub_cells: Dict[int, List[str]] = {col: [] for col in money_cols}
        groups = arrange(spec, rows)
        for key, grows in groups:
            if spec["group_by"] != "none":
                safe_text(ws.cell(row=r, column=1), key or "—")
                ws.cell(row=r, column=1).font = X._BOLD
                r += 1
            first = r
            for row in grows:
                for i, c in enumerate(cols, start=1):
                    fld, v = src.fields[c], row.get(c)
                    cell = ws.cell(row=r, column=i)
                    if fld.money:
                        X._money(cell, v)
                    elif fld.kind == "date":
                        cell.value = X._as_date(v) if v else None
                        cell.number_format = X._DATE_FMT
                    elif fld.kind == "number":
                        cell.value = v
                    else:
                        safe_text(cell, v if v not in (None, "") else None)
                r += 1
            last = r - 1
            if spec["group_by"] != "none" and money_cols:
                safe_text(ws.cell(row=r, column=1), f"Subtotal {key or '—'}")
                ws.cell(row=r, column=1).font = X._BOLD
                for col in money_cols:
                    L = get_column_letter(col)
                    cell = ws.cell(row=r, column=col, value=f"=SUM({L}{first}:{L}{last})")
                    cell.number_format, cell.font = X._MONEY_FMT, X._BOLD
                    sub_cells[col].append(f"{L}{r}")
                r += 2
        want: Dict[str, int] = {}
        if money_cols and rows:
            ws.cell(row=r, column=1, value="TOTAL").font = X._BOLD
            tot = totals(spec, rows)
            for col in money_cols:
                L = get_column_letter(col)
                if spec["group_by"] != "none":
                    formula = "=" + "+".join(sub_cells[col]) if sub_cells[col] else 0
                else:
                    formula = f"=SUM({L}{head + 1}:{L}{r - 1})"
                cell = ws.cell(row=r, column=col, value=formula)
                cell.number_format, cell.font = X._MONEY_FMT, X._BOLD
                want[f"{L}{r}"] = tot[cols[col - 1]]
        X._finish_table(ws, head, max(r - 1, head), len(cols))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), want
    finally:
        X._fill_var.reset(token)


# ── drafts and templates (migration 238) ───────────────────────────────────

def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def save_draft(tenant_id: str, spec: Dict[str, Any], request: str, by: str, preview: Dict[str, int],
               rows: int, template_id: Optional[str] = None) -> str:
    row = {"tenant_id": tenant_id, "status": "draft", "spec": {**spec, "_preview_totals": preview, "_rows": rows},
           "request": (request or "")[:1000], "created_by": by or None, "template_id": template_id,
           "created_at": _now(), "updated_at": _now()}
    res = _db().table(TABLE).insert(row).execute()
    return ((res.data or [{}])[0]).get("id") or ""


def latest_draft(tenant_id: str, by: str, built_too: bool = False) -> Optional[Dict[str, Any]]:
    """This person's most recent sheet draft (the last 7 days); with built_too, also one already
    built — "save it as the Tile Co sheet" comes after "build it"."""
    since = (datetime.now(timezone.utc) - timedelta(days=DRAFT_DAYS)).isoformat()
    q = (_db().table(TABLE).select("*").eq("tenant_id", tenant_id)
         .in_("status", ["draft", "built"] if built_too else ["draft"]).gte("created_at", since))
    q = q.eq("created_by", by) if by else q.is_("created_by", "null")
    rows = q.order("created_at", desc=True).limit(1).execute().data or []
    return rows[0] if rows else None


def templates(tenant_id: str) -> List[Dict[str, Any]]:
    try:
        return (_db().table(TABLE).select("id,name,spec,last_built_at,uses").eq("tenant_id", tenant_id)
                .eq("status", "template").order("name").execute().data or [])
    except Exception as exc:  # noqa: BLE001 — before migration 238
        log.debug("sheet templates unreadable: %s", exc)
        return []


def find_template(tenant_id: str, name: str) -> Optional[Dict[str, Any]]:
    want = _norm(name)
    ts = templates(tenant_id)
    exact = [t for t in ts if _norm(t.get("name")) == want]
    if exact:
        return exact[0]
    near = [t for t in ts if want and want in _norm(t.get("name"))]
    return near[0] if len(near) == 1 else None


def save_template(tenant_id: str, draft: Dict[str, Any], name: str, by: str) -> Dict[str, Any]:
    """Keep a draft's layout under a name for this business; a same-named template is replaced.
    Read back."""
    name = re.sub(r"\s+", " ", name or "").strip()[:80]
    if not name:
        raise SpecError("What should I call this sheet? e.g. “the Tile Co monthly sheet”.")
    spec = {k: v for k, v in (draft.get("spec") or {}).items() if not k.startswith("_")}
    old = find_template(tenant_id, name)
    if old and _norm(old.get("name")) == _norm(name):
        _db().table(TABLE).update({"spec": spec, "updated_at": _now(), "created_by": by or None}) \
            .eq("tenant_id", tenant_id).eq("id", old["id"]).execute()
        tid, replaced = old["id"], True
    else:
        res = _db().table(TABLE).insert({"tenant_id": tenant_id, "name": name, "status": "template", "spec": spec,
                                         "request": draft.get("request"), "created_by": by or None,
                                         "created_at": _now(), "updated_at": _now()}).execute()
        tid, replaced = ((res.data or [{}])[0]).get("id"), False
    back = (_db().table(TABLE).select("id,name,spec").eq("tenant_id", tenant_id).eq("id", tid)
            .limit(1).execute().data or []) if tid else []
    if not back or back[0].get("spec", {}).get("source") != spec.get("source"):
        raise RuntimeError("the sheet was not saved")
    return {**back[0], "replaced": replaced}


def mark_built(tenant_id: str, draft: Dict[str, Any]) -> None:
    try:
        _db().table(TABLE).update({"status": "built", "last_built_at": _now(), "updated_at": _now()}) \
            .eq("tenant_id", tenant_id).eq("id", draft["id"]).execute()
        if draft.get("template_id"):
            t = (_db().table(TABLE).select("uses").eq("tenant_id", tenant_id).eq("id", draft["template_id"])
                 .limit(1).execute().data or [{}])[0]
            _db().table(TABLE).update({"last_built_at": _now(), "uses": int(t.get("uses") or 0) + 1}) \
                .eq("tenant_id", tenant_id).eq("id", draft["template_id"]).execute()
    except Exception as exc:  # noqa: BLE001 — bookkeeping only
        log.debug("sheet build not recorded: %s", exc)

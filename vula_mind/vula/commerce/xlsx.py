"""vula/commerce/xlsx.py

Renders a supplier spend/materials history (a find_filed_document result, the same data
format_supplier_history_reply turns into WhatsApp text) as a real .xlsx workbook, so an
explicit "in excel"/"spreadsheet" request can be answered with an actual file instead of the
text-only reply. Mirrors vula/commerce/pdf.py::render_invoice_pdf's shape (a pure function
returning bytes) and vula/takeoff/boq_export.py's styling conventions.
"""
from __future__ import annotations

import io
import re
from typing import Any, Dict, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="FF2C5545")
_HEADER_FONT = Font(bold=True, color="FFFFFFFF")
_BOLD = Font(bold=True)
_MONEY_FMT = 'R #,##0.00'


def _header_row(ws: Any, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="left")


def _autosize(ws: Any, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def render_supplier_history_xlsx(result: Dict[str, Any], supplier: str = "") -> Optional[bytes]:
    """A supplier spend/materials find_filed_document result as .xlsx bytes — an "Invoices"
    sheet (date, document, amount, refund flag) plus a "Materials" sheet when the result
    carries a materials roll-up. None when there's nothing complete to export (mirrors
    format_supplier_history_reply's own None case), so the caller can fall back to text.
    """
    if result.get("status") != "found" or "total_amount_cents" not in result:
        return None
    matches = result.get("matches") or []
    total_n = int(result.get("total_matches") or len(matches))
    if not total_n:
        return None
    name = supplier or result.get("resolved_supplier") or next(
        (m.get("party") for m in matches if m.get("party")), None) or "Supplier"

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Invoices")
    ws.cell(row=1, column=1, value=name).font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=f"{total_n} document{'s' if total_n != 1 else ''}, "
                                    f"total spend {result.get('total_amount')}")
    _header_row(ws, 4, ["Date", "Document", "Amount", "Refund"])
    row = 5
    for m in matches:
        doc_name = re.sub(r"\.(pdf|jpe?g|png)$", "", m.get("filename") or "document", flags=re.I)
        ws.cell(row=row, column=1, value=(m.get("filed_at") or "")[:10])
        ws.cell(row=row, column=2, value=doc_name)
        amt_cell = ws.cell(row=row, column=3)
        if m.get("amount") is not None:
            amt_cell.value = float(m["amount"])
            amt_cell.number_format = _MONEY_FMT
        ws.cell(row=row, column=4, value="Yes" if m.get("is_refund") else "")
        row += 1
    if total_n > len(matches):
        ws.cell(row=row, column=2,
                value=f"…and {total_n - len(matches)} more (all included in the total)")
        row += 1
    total_cell = ws.cell(row=row + 1, column=2, value="Total")
    total_cell.font = _BOLD
    total_amt = ws.cell(row=row + 1, column=3, value=int(result["total_amount_cents"]) / 100)
    total_amt.number_format = _MONEY_FMT
    total_amt.font = _BOLD
    _autosize(ws, [12, 42, 14, 9])

    materials = result.get("materials") or []
    if materials:
        ms = wb.create_sheet("Materials")
        _header_row(ms, 1, ["Description", "Quantity", "Spend", "Check quantity"])
        for i, it in enumerate(materials, start=2):
            ms.cell(row=i, column=1, value=it.get("description"))
            ms.cell(row=i, column=2, value=it.get("quantity"))
            spend_cell = ms.cell(row=i, column=3)
            if it.get("spend_cents") is not None:
                spend_cell.value = int(it["spend_cents"]) / 100
                spend_cell.number_format = _MONEY_FMT
            ms.cell(row=i, column=4, value="⚠️" if it.get("unit_price_varies") else "")
        _autosize(ms, [42, 12, 14, 14])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
